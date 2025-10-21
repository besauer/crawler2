from __future__ import annotations

import base64
import io
import json
import math
import os
import platform
import re
import subprocess
import threading
import time
import traceback
import unicodedata
import uuid
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Set, Tuple
from urllib.parse import urljoin, urlparse, urlunparse

import requests
from flask import Flask, Response, jsonify, render_template, request
from flask.signals import got_request_exception
from flask.typing import ResponseReturnValue
from openpyxl import load_workbook
from bs4 import BeautifulSoup
from markupsafe import Markup, escape
from werkzeug.exceptions import HTTPException

from crawler import (
    CrawlProgress,
    CrawlResult,
    crawl_site,
    MAX_PAGES_DEFAULT,
)
from storage import storage

app = Flask(__name__)


@app.template_filter("break_every")
def template_break_every(value: object, interval: int = 30) -> Markup:
    text = "" if value is None else str(value)
    if not text:
        return Markup("")
    interval = max(1, int(interval or 1))
    chunks = [f"<span class='d-block'>{escape(text[i : i + interval])}</span>" for i in range(0, len(text), interval)]
    return Markup("".join(chunks))

saved_search_lock = threading.Lock()
search_definition_lock = threading.Lock()
search_run_lock = threading.Lock()
search_result_lock = threading.Lock()
search_log_lock = threading.Lock()
search_jobs_lock = threading.Lock()
search_categories_lock = threading.Lock()
DEFAULT_CONCURRENCY = 5
MAX_CONCURRENCY = 150
MAX_MAX_PAGES = 1000
DEFAULT_RESPECT_ROBOTS = True
stammdaten_lock = threading.Lock()
settings_lock = threading.Lock()
synonym_cache_lock = threading.Lock()
keyword_finder_cache_lock = threading.Lock()
OPENAI_CHAT_COMPLETIONS_URL = "https://api.openai.com/v1/chat/completions"
OPENAI_MODELS_URL = "https://api.openai.com/v1/models?limit=1"
OPENAI_MODEL_NAME = "gpt-4o-mini"
OPENAI_TIMEOUT = 20
DEFAULT_SYNONYM_PROMPT = (
    "Du bist ein Synonym-Experte für deutschsprachige Websuche und spezialisiert auf Synonyme, die sich auf Webseiten "
    "von Schulen in Baden Württemberg finden. Denke dabei vor allem an Keywords, die so banal sind, dass man sie schnell "
    "vergisst. Die aber häufig vorkommen. Gib ausschließlich eine JSON-Liste mit gebräuchlichen Synonymen, nahen Begriffen "
    "und typischen Kurzformen des folgenden Begriffs zurück. Keine Erklärungen, keine weiteren Wörter. Vermeide Homonyme "
    "und irrelevante Bedeutungen."
)
DEFAULT_KEYWORD_FINDER_PROMPT = (
    "Du kombinierst Vorschläge des Google Keyword Planner mit zusätzlicher Analyse. Du erhältst ein JSON-Objekt mit "
    "Keyword-Vorschlägen (einschließlich Suchvolumen, Wettbewerb und Relevanz). Gruppiere die Begriffe semantisch in "
    "3–8 Themencluster, entferne irrelevante oder doppelte Begriffe und gib pro Cluster die drei bis fünf relevantesten "
    "Begriffe zurück. Nutze kurze, aussagekräftige Cluster-Namen. Antworte ausschließlich im folgenden JSON-Format: "
    "{\"clusters\": [{\"label\": \"<Cluster>\", \"keywords\": [{\"term\": \"<Begriff>\", \"search_volume\": <Zahl|null>, "
    "\"competition\": \"<niedrig/mittel/hoch/null>\", \"relevance\": <Zahl|null>}]}]}}. Keine weiteren Texte."
)
DEFAULT_KEYWORD_FINDER_MAX_RESULTS = 150
DEFAULT_KEYWORD_FINDER_TEMPERATURE = 0.3
GOOGLE_KEYWORD_TIMEOUT = 20
GOOGLE_KEYWORD_API_URL = "https://keywordsearch.googleapis.com/v1beta1/suggestKeywords"
DEFAULT_KEYWORD_EVALUATION_PROMPT = (
    "Du bewertest deutsche Suchbegriffe für eine Websuche. "
    "Antworte ausschließlich mit einem JSON-Objekt mit den Schlüsseln "
    '"quality"' " (Werte: \"hoch\", \"mittel\" oder \"niedrig\") und "
    '"better_keywords"' " (Liste von bis zu fünf besseren oder spezifischeren deutschen Suchbegriffen). "
    "Falls es keine besseren Begriffe gibt, gib eine leere Liste zurück. Keine Erklärungen."
)
DEFAULT_SYNONYM_MAX_RESULTS = 50
DEFAULT_SYNONYM_TEMPERATURE = 0.2
AVAILABLE_OPENAI_MODELS = [
    "gpt-4o-mini",
    "gpt-4o",
    "gpt-4.1-mini",
    "gpt-4.1",
    "o4-mini",
    "gpt-3.5-turbo",
    "llama-3.1-70b",
]
data_quality_lock = threading.Lock()
DATA_QUALITY_SITE_PROMPT = (
    "Du prüfst, ob eine Webseite zur angegebenen Schule gehört. Du erhältst Schulname und Ort sowie komprimierte "
    "Textauszüge aus Startseite und Impressum. Bewerte ausschließlich, ob die Seite mit hoher Wahrscheinlichkeit "
    "die offizielle Webseite der Schule ist. Antworte als JSON-Objekt mit den Schlüsseln \"bewertung\" (Werte: OK, "
    "NEIN, UNSICHER), \"confidence\" (0.0–1.0), \"begruendung\" (kurze Begründung), \"gefundene_signale\" "
    "(Liste kurzer Stichworte), \"auftritt_typ\" (Werte: EXKLUSIV für eine eigenständige Schulwebseite oder GETEILT "
    "für einen geteilten Auftritt auf einer fremden Domain) und \"auftritt_begruendung\" (kurze Begründung für den "
    "Auftrittstyp)."
)
DATA_QUALITY_SEARCH_PROMPT = (
    "Du erhältst Schulname, Ort und eine Liste möglicher Webseiten aus einer Google-Suche. Wähle die URL aus, die am "
    "ehesten die offizielle Seite der Schule ist. Entferne irrelevante Treffer. Gib ein JSON-Objekt mit den "
    "Schlüsseln \"empfehlung\" (URL oder null), \"confidence\" (0.0–1.0), \"begruendung\" (kurze Begründung) und "
    "\"signale\" (Liste kurzer Stichworte) zurück."
)
DEFAULT_DATA_QUALITY_SETTINGS = {
    "model": OPENAI_MODEL_NAME,
    "temperature": 0.1,
    "confidence_threshold": 0.85,
    "max_search_results": 10,
    "batch_size": 5000,
    "parallel_checks": 5,
    "dry_run": False,
    "search_language": "de",
    "search_region": "de",
    "allowed_domains": [],
    "blocked_domains": [],
}
GOOGLE_SEARCH_API_URL = "https://www.googleapis.com/customsearch/v1"
GOOGLE_SEARCH_TIMEOUT = 20
DATA_QUALITY_IMPRESSUM_KEYWORDS = {
    "impressum",
    "kontakt",
    "datenschutz",
    "rechtliches",
    "legal",
    "anbieter",
}
MAX_DATA_QUALITY_HISTORY = 20

SITE_SCOPE_EXCLUSIVE = "exclusive"
SITE_SCOPE_SHARED = "shared"
SITE_SCOPE_UNKNOWN = "unknown"
SITE_SCOPE_LABELS = {
    SITE_SCOPE_EXCLUSIVE: "Exklusive Webseite",
    SITE_SCOPE_SHARED: "Geteilter Auftritt",
    SITE_SCOPE_UNKNOWN: "Auftritt unbekannt",
}
SITE_SCOPE_CLASSES = {
    SITE_SCOPE_EXCLUSIVE: "text-bg-success",
    SITE_SCOPE_SHARED: "text-bg-warning",
    SITE_SCOPE_UNKNOWN: "text-bg-secondary",
}

DEFAULT_SEARCH_EVALUATION_PROMPT = (
    "Bewertung schulischer KI-Integration Anweisung: Analysiere den folgenden Text einer Schulwebseite. "
    "Bewerte, welche der 9 Dimensionen schulischer KI-Integration zutreffen und wie stark sie ausgeprägt sind (1–5). "
    "Gib zusätzlich für jede Bewertung eine Sicherheit (0–1) an. Mehrere Dimensionen können gleichzeitig zutreffen. "
    "Formatiere das Ergebnis immer exakt im folgenden JSON-Schema: { \"Bewertung\": [ {\"Dimension\": \"Einsatz im Unterricht\", "
    "\"Wert\": 0, \"Sicherheit\": 0.0}, {\"Dimension\": \"Medienbildung & KI-Kompetenzen\", \"Wert\": 0, \"Sicherheit\": 0.0}, "
    "{\"Dimension\": \"Ethik, Datenschutz & Verantwortung\", \"Wert\": 0, \"Sicherheit\": 0.0}, {\"Dimension\": \"Partizipation & Mitgestaltung\", "
    "\"Wert\": 0, \"Sicherheit\": 0.0}, {\"Dimension\": \"Infrastruktur & Zugänglichkeit\", \"Wert\": 0, \"Sicherheit\": 0.0}, "
    "{\"Dimension\": \"Schulentwicklung & Steuerung\", \"Wert\": 0, \"Sicherheit\": 0.0}, {\"Dimension\": \"Fortbildung & Professionalisierung\", "
    "\"Wert\": 0, \"Sicherheit\": 0.0}, {\"Dimension\": \"Schulkultur & Kommunikation\", \"Wert\": 0, \"Sicherheit\": 0.0}, "
    "{\"Dimension\": \"Kooperation & Netzwerke\", \"Wert\": 0, \"Sicherheit\": 0.0} ] } Bewertungslogik Wert Bedeutung 1 Klarer Hinweis "
    "auf KI-Aktivität, jedoch nur sporadisch oder punktuell. 2 KI ist im Alltag erkennbar integriert, jedoch ohne formale Konzepte. "
    "3 Schulweite Nutzung oder erkennbare Konzepte und feste Strukturen vorhanden. 4 Klare, ausgearbeitete Konzepte und tiefe "
    "Verankerung im Schulalltag. 5 Außergewöhnlich stark ausgeprägt, innovativ, tief reflektiert und strategisch verankert. "
    "Text zur Analyse {{CONTENT}}"
)
DEFAULT_SEARCH_MODEL = OPENAI_MODEL_NAME
DEFAULT_SEARCH_TEMPERATURE = 0.2
DEFAULT_SEARCH_MAX_RESULTS = 200
DEFAULT_SEARCH_CONCURRENCY = 20
DEFAULT_SEARCH_CATEGORIES = ["KI"]
SEARCH_STATUS_NEW = "neu"
SEARCH_STATUS_RUNNING = "in_arbeit"
SEARCH_STATUS_FINISHED = "abgeschlossen"
SEARCH_DIMENSIONS = [
    "Einsatz im Unterricht",
    "Medienbildung & KI-Kompetenzen",
    "Ethik, Datenschutz & Verantwortung",
    "Partizipation & Mitgestaltung",
    "Infrastruktur & Zugänglichkeit",
    "Schulentwicklung & Steuerung",
    "Fortbildung & Professionalisierung",
    "Schulkultur & Kommunikation",
    "Kooperation & Netzwerke",
]


def utcnow_iso() -> str:
    return datetime.utcnow().isoformat() + "Z"


EXPORT_SCHEMA_VERSION = 1
EXPORT_SECTION_DEFINITIONS: Dict[str, Dict[str, str]] = {
    "stammdaten": {
        "label": "Stammdaten",
        "description": "Schulverzeichnis inkl. aller Zusatzfelder.",
    },
    "keywords": {
        "label": "Keywords & Synonyme",
        "description": "Gespeicherte Suchläufe, Synonym- und Keyword-Cache.",
    },
    "searches": {
        "label": "Suchergebnisse",
        "description": "Suchdefinitionen, Auswertungen und Bewertungsresultate.",
    },
    "settings": {
        "label": "Anwendungs- & Crawl-Einstellungen",
        "description": "Voreinstellungen für Synonyme, Crawl, Keyword-Finder und Datenqualität.",
    },
    "data_quality": {
        "label": "Datenqualitäts-Ergebnisse",
        "description": "Prüfstatus, Vorschläge und Historie pro Schule.",
    },
    "api_keys": {
        "label": "API-Schlüssel & Tokens",
        "description": "OpenAI-, Google- und weitere Zugangsdaten.",
    },
    "debug": {
        "label": "Debugging-Protokolle",
        "description": "Gesammelte Fehlermeldungen und Kontextinformationen.",
    },
}
EXPORT_SECTION_ORDER = [
    "stammdaten",
    "keywords",
    "searches",
    "settings",
    "data_quality",
    "debug",
    "api_keys",
]

IMPORT_SESSION_TTL_SECONDS = 600
IMPORT_SESSION_LIMIT = 8

audit_log_lock = threading.Lock()
error_log_lock = threading.Lock()


@app.context_processor
def inject_export_import_metadata() -> Dict[str, Any]:
    """Make export/import metadata available to all templates."""

    return {
        "export_sections": EXPORT_SECTION_ORDER,
        "export_definitions": EXPORT_SECTION_DEFINITIONS,
    }

_import_sessions: Dict[str, Dict[str, Any]] = {}
import_sessions_lock = threading.Lock()
QUALITY_STATUS_PENDING = "pending"
QUALITY_STATUS_OK = "ok"
QUALITY_STATUS_UNSURE = "unsure"
QUALITY_STATUS_INVALID = "invalid"

api_key_lock = threading.Lock()
_api_key_value: Optional[str] = None
google_api_key_lock = threading.Lock()
_google_api_key_value: Optional[str] = None

BACKUP_DIR = Path("backups")
BACKUP_FILE = BACKUP_DIR / "auto_backup.json"
BACKUP_INTERVAL_SECONDS = 300
_backup_thread_started = False

STAMMDATEN_PRIMARY_FIELDS = [
    {"key": "jahr", "label": "Jahr", "type": "numeric"},
    {"key": "schul_id", "label": "Schul ID"},
    {"key": "status", "label": "Status"},
    {"key": "anzahl_aussenstellen", "label": "Anzahl Außenstellen", "type": "numeric"},
    {"key": "rb", "label": "RB"},
    {"key": "kkz", "label": "KKZ"},
    {"key": "rkz", "label": "RKZ"},
    {"key": "schulname", "label": "Schulname"},
    {"key": "strasse", "label": "Straße"},
    {"key": "plz", "label": "PLZ"},
    {"key": "ort", "label": "Ort"},
    {"key": "telefon", "label": "Telefon"},
    {"key": "fax", "label": "Fax"},
    {"key": "homepage", "label": "URL"},
    {"key": "schueler_gesamt", "label": "Schüler/-innen insgesamt", "type": "numeric"},
    {"key": "klassen_gesamt", "label": "Klassen insgesamt", "type": "numeric"},
]

STAMMDATEN_BOOLEAN_FIELDS = [
    {"key": "angebot_grundschulen", "label": "12 – Grundschulen", "type": "bool"},
    {"key": "angebot_werkreal_hauptschulen", "label": "14 – Werkreal-/Hauptschulen", "type": "bool"},
    {"key": "angebot_realschulen", "label": "15 – Realschulen", "type": "bool"},
    {"key": "angebot_gymnasien", "label": "16 – Gymnasien", "type": "bool"},
    {"key": "angebot_schulen_besonderer_art", "label": "17 – Schulen besonderer Art", "type": "bool"},
    {"key": "angebot_freie_waldorfschulen", "label": "19 – Freie Waldorfschulen", "type": "bool"},
    {"key": "angebot_gemeinschaftsschule_sek1", "label": "21 – Gemeinschaftsschulen – Sekundarstufe I", "type": "bool"},
    {"key": "angebot_gemeinschaftsschule_sek2", "label": "21 – Gemeinschaftsschulen – Sekundarstufe II", "type": "bool"},
    {"key": "angebot_berufsschulen", "label": "31 – Berufsschulen", "type": "bool"},
    {"key": "angebot_berufsfachschulen", "label": "32 – Berufsfachschulen", "type": "bool"},
    {"key": "angebot_berufskollegs", "label": "33 – Berufskollegs", "type": "bool"},
    {"key": "angebot_berufsoberschulen", "label": "34 – Berufsoberschulen", "type": "bool"},
    {"key": "angebot_fachschulen", "label": "35 – Fachschulen", "type": "bool"},
    {"key": "angebot_gesundheitswesen", "label": "35 – Schulen für Berufe des Gesundheitswesens", "type": "bool"},
    {"key": "angebot_berufliche_gymnasien", "label": "36 – Berufliche Gymnasien", "type": "bool"},
    {"key": "angebot_sbbz", "label": "51 – Sonderpädagogische Bildungs- und Beratungszentren", "type": "bool"},
    {"key": "angebot_sonderberufsschulen", "label": "52 – Sonderberufsschulen", "type": "bool"},
    {"key": "angebot_zweiter_bildungsweg", "label": "Zweiter Bildungsweg", "type": "bool"},
]

STAMMDATEN_FIELDS = STAMMDATEN_PRIMARY_FIELDS + STAMMDATEN_BOOLEAN_FIELDS
STAMMDATEN_NUMERIC_KEYS = {field["key"] for field in STAMMDATEN_FIELDS if field.get("type") == "numeric"}
STAMMDATEN_BOOLEAN_KEYS = {field["key"] for field in STAMMDATEN_FIELDS if field.get("type") == "bool"}
LEGACY_STAMMDATEN_ALIASES = {
    "schulname": ["name"],
    "schueler_gesamt": ["anzahl_schueler"],
    "klassen_gesamt": ["anzahl_klassen"],
    "homepage": ["homepage", "url"],
    "status": ["traeger"],
    "ort": ["ort"],
}


class OpenAIIntegrationError(Exception):
    """Raised when the OpenAI integration cannot provide synonyms."""


class KeywordPlannerError(Exception):
    """Raised when keyword planner suggestions cannot be fetched."""


def _default_synonym_settings() -> Dict[str, object]:
    return {
        "prompt": DEFAULT_SYNONYM_PROMPT,
        "max_results": DEFAULT_SYNONYM_MAX_RESULTS,
        "temperature": DEFAULT_SYNONYM_TEMPERATURE,
        "model": OPENAI_MODEL_NAME,
    }


def _default_search_settings() -> Dict[str, object]:
    return {
        "prompt": DEFAULT_SEARCH_EVALUATION_PROMPT,
        "max_results": DEFAULT_SEARCH_MAX_RESULTS,
        "temperature": DEFAULT_SEARCH_TEMPERATURE,
        "model": DEFAULT_SEARCH_MODEL,
        "concurrency": DEFAULT_SEARCH_CONCURRENCY,
    }


def _read_settings_unlocked() -> Dict[str, object]:
    data = storage.get_json("settings", "data", {})
    if isinstance(data, dict):
        return data
    return {}


def load_settings_data() -> Dict[str, object]:
    with settings_lock:
        return dict(_read_settings_unlocked())


def save_settings_data(data: Dict[str, object]) -> None:
    with settings_lock:
        storage.set_json("settings", "data", data)


def is_developer_mode_enabled() -> bool:
    data = load_settings_data()
    return bool(data.get("developer_mode"))


def set_developer_mode_enabled(enabled: bool) -> bool:
    data = load_settings_data()
    data["developer_mode"] = bool(enabled)
    save_settings_data(data)
    return bool(data["developer_mode"])


def get_synonym_defaults() -> Dict[str, object]:
    data = load_settings_data()
    defaults = _default_synonym_settings()
    result = {}
    result.update(defaults)
    settings_synonyms = data.get("synonym_defaults") if isinstance(data, dict) else {}
    if isinstance(settings_synonyms, dict):
        prompt = str(settings_synonyms.get("prompt", "")).strip()
        if prompt:
            result["prompt"] = prompt
        max_results = settings_synonyms.get("max_results")
        if isinstance(max_results, int) and 1 <= max_results <= 150:
            result["max_results"] = max_results
        temperature = settings_synonyms.get("temperature")
        try:
            temp_value = float(temperature)
        except (TypeError, ValueError):
            temp_value = None
        if temp_value is not None and 0 <= temp_value <= 2:
            result["temperature"] = temp_value
        model = str(settings_synonyms.get("model", "")).strip()
        if model:
            result["model"] = model
    return result


def get_search_defaults() -> Dict[str, object]:
    data = load_settings_data()
    defaults = _default_search_settings()
    result = dict(defaults)
    settings_search = data.get("search_defaults") if isinstance(data, dict) else {}
    if isinstance(settings_search, dict):
        prompt = str(settings_search.get("prompt", "")).strip()
        if prompt:
            result["prompt"] = prompt
        max_results_value = settings_search.get("max_results")
        try:
            max_results = int(max_results_value)
        except (TypeError, ValueError):
            max_results = defaults["max_results"]
        max_results = max(1, min(max_results, 200))
        result["max_results"] = max_results
        temperature_value = settings_search.get("temperature")
        try:
            temperature = float(temperature_value)
        except (TypeError, ValueError):
            temperature = defaults["temperature"]
        temperature = max(0.0, min(temperature, 2.0))
        result["temperature"] = temperature
        model_value = str(settings_search.get("model", "")).strip()
        if model_value:
            result["model"] = model_value
        concurrency_value = settings_search.get("concurrency")
        try:
            concurrency = int(concurrency_value)
        except (TypeError, ValueError):
            concurrency = defaults["concurrency"]
        if concurrency < 1:
            concurrency = 1
        result["concurrency"] = concurrency
    return result


def update_synonym_defaults(values: Dict[str, object]) -> Dict[str, object]:
    current = load_settings_data()
    defaults = _default_synonym_settings()
    prompt = str(values.get("prompt", "")).strip() or defaults["prompt"]
    max_results_raw = values.get("max_results")
    try:
        max_results = int(max_results_raw)
    except (TypeError, ValueError):
        max_results = defaults["max_results"]
    if max_results < 1:
        max_results = 1
    elif max_results > 150:
        max_results = 150
    temperature_raw = values.get("temperature")
    try:
        temperature = float(temperature_raw)
    except (TypeError, ValueError):
        temperature = defaults["temperature"]
    if temperature < 0:
        temperature = 0.0
    elif temperature > 2:
        temperature = 2.0
    model = str(values.get("model", "")).strip() or defaults["model"]

    current.setdefault("synonym_defaults", {})
    current["synonym_defaults"] = {
        "prompt": prompt,
        "max_results": max_results,
        "temperature": temperature,
        "model": model,
    }
    save_settings_data(current)
    return current["synonym_defaults"]


def update_search_defaults(values: Dict[str, object]) -> Dict[str, object]:
    current = load_settings_data()
    defaults = _default_search_settings()
    prompt = str(values.get("prompt", "")).strip() or defaults["prompt"]
    try:
        max_results = int(values.get("max_results", defaults["max_results"]))
    except (TypeError, ValueError):
        max_results = defaults["max_results"]
    max_results = max(1, min(max_results, 200))
    try:
        temperature = float(values.get("temperature", defaults["temperature"]))
    except (TypeError, ValueError):
        temperature = defaults["temperature"]
    temperature = max(0.0, min(temperature, 2.0))
    model = str(values.get("model", "")).strip() or defaults["model"]
    try:
        concurrency = int(values.get("concurrency", defaults["concurrency"]))
    except (TypeError, ValueError):
        concurrency = defaults["concurrency"]
    if concurrency < 1:
        concurrency = 1

    current.setdefault("search_defaults", {})
    current["search_defaults"] = {
        "prompt": prompt,
        "max_results": max_results,
        "temperature": temperature,
        "model": model,
        "concurrency": concurrency,
    }
    save_settings_data(current)
    return current["search_defaults"]


def _default_keyword_finder_settings() -> Dict[str, object]:
    return {
        "prompt": DEFAULT_KEYWORD_FINDER_PROMPT,
        "max_results": DEFAULT_KEYWORD_FINDER_MAX_RESULTS,
        "temperature": DEFAULT_KEYWORD_FINDER_TEMPERATURE,
        "model": OPENAI_MODEL_NAME,
    }


def get_keyword_finder_defaults() -> Dict[str, object]:
    data = load_settings_data()
    defaults = _default_keyword_finder_settings()
    result = dict(defaults)
    if isinstance(data, dict):
        stored = data.get("keyword_finder_defaults")
        if isinstance(stored, dict):
            prompt = str(stored.get("prompt", "")).strip()
            if prompt:
                result["prompt"] = prompt
            max_results = stored.get("max_results")
            if isinstance(max_results, int):
                result["max_results"] = max(1, min(max_results, DEFAULT_KEYWORD_FINDER_MAX_RESULTS))
            temperature = stored.get("temperature")
            try:
                temp_value = float(temperature)
            except (TypeError, ValueError):
                temp_value = result["temperature"]
            else:
                temp_value = max(0.0, min(temp_value, 2.0))
            result["temperature"] = temp_value
            model_value = str(stored.get("model", "")).strip()
            if model_value:
                result["model"] = model_value
    return result


def update_keyword_finder_defaults(values: Dict[str, object]) -> Dict[str, object]:
    current = load_settings_data()
    defaults = _default_keyword_finder_settings()

    prompt = str(values.get("prompt", "")).strip() or defaults["prompt"]
    try:
        max_results = int(values.get("max_results", defaults["max_results"]))
    except (TypeError, ValueError):
        max_results = defaults["max_results"]
    max_results = max(1, min(max_results, DEFAULT_KEYWORD_FINDER_MAX_RESULTS))

    try:
        temperature = float(values.get("temperature", defaults["temperature"]))
    except (TypeError, ValueError):
        temperature = defaults["temperature"]
    temperature = max(0.0, min(temperature, 2.0))

    model = str(values.get("model", "")).strip() or defaults["model"]

    current.setdefault("keyword_finder_defaults", {})
    current["keyword_finder_defaults"] = {
        "prompt": prompt,
        "max_results": max_results,
        "temperature": temperature,
        "model": model,
    }
    save_settings_data(current)
    return current["keyword_finder_defaults"]


def get_google_search_credentials() -> Dict[str, str]:
    data = load_settings_data()
    if not isinstance(data, dict):
        return {"api_key": "", "cx": ""}
    return {
        "api_key": str(data.get("google_search_api_key", "") or "").strip(),
        "cx": str(data.get("google_search_cx", "") or "").strip(),
    }


def update_google_search_credentials(api_key: str, cx: str) -> Dict[str, str]:
    data = load_settings_data()
    if not isinstance(data, dict):
        data = {}
    data["google_search_api_key"] = api_key.strip()
    data["google_search_cx"] = cx.strip()
    save_settings_data(data)
    return get_google_search_credentials()


def clear_google_search_credentials() -> None:
    data = load_settings_data()
    if not isinstance(data, dict):
        data = {}
    data.pop("google_search_api_key", None)
    data.pop("google_search_cx", None)
    save_settings_data(data)


def _default_data_quality_settings() -> Dict[str, object]:
    return dict(DEFAULT_DATA_QUALITY_SETTINGS)


def get_data_quality_settings() -> Dict[str, object]:
    data = load_settings_data()
    defaults = _default_data_quality_settings()
    if not isinstance(data, dict):
        return defaults
    stored = data.get("data_quality_settings")
    if not isinstance(stored, dict):
        return defaults
    result = dict(defaults)
    model = str(stored.get("model", "") or "").strip()
    if model:
        result["model"] = model
    try:
        temperature = float(stored.get("temperature", defaults["temperature"]))
    except (TypeError, ValueError):
        temperature = defaults["temperature"]
    temperature = max(0.0, min(temperature, 2.0))
    result["temperature"] = temperature
    try:
        threshold = float(stored.get("confidence_threshold", defaults["confidence_threshold"]))
    except (TypeError, ValueError):
        threshold = defaults["confidence_threshold"]
    threshold = max(0.0, min(threshold, 1.0))
    result["confidence_threshold"] = threshold
    try:
        max_search = int(stored.get("max_search_results", defaults["max_search_results"]))
    except (TypeError, ValueError):
        max_search = defaults["max_search_results"]
    result["max_search_results"] = max(1, min(max_search, 100))
    try:
        batch_size = int(stored.get("batch_size", defaults["batch_size"]))
    except (TypeError, ValueError):
        batch_size = defaults["batch_size"]
    result["batch_size"] = max(1, batch_size)
    try:
        parallel_checks = int(stored.get("parallel_checks", defaults["parallel_checks"]))
    except (TypeError, ValueError):
        parallel_checks = defaults["parallel_checks"]
    result["parallel_checks"] = max(1, min(parallel_checks, 50))
    result["dry_run"] = bool(stored.get("dry_run", defaults["dry_run"]))
    language = str(stored.get("search_language", defaults["search_language"]) or "").strip() or defaults["search_language"]
    result["search_language"] = language
    region = str(stored.get("search_region", defaults["search_region"]) or "").strip() or defaults["search_region"]
    result["search_region"] = region
    allowed = stored.get("allowed_domains", defaults["allowed_domains"])
    if isinstance(allowed, str):
        allowed = [item.strip() for item in allowed.splitlines() if item.strip()]
    elif isinstance(allowed, list):
        allowed = [str(item).strip() for item in allowed if str(item).strip()]
    else:
        allowed = []
    result["allowed_domains"] = allowed
    blocked = stored.get("blocked_domains", defaults["blocked_domains"])
    if isinstance(blocked, str):
        blocked = [item.strip() for item in blocked.splitlines() if item.strip()]
    elif isinstance(blocked, list):
        blocked = [str(item).strip() for item in blocked if str(item).strip()]
    else:
        blocked = []
    result["blocked_domains"] = blocked
    return result


def update_data_quality_settings(values: Dict[str, object]) -> Dict[str, object]:
    data = load_settings_data()
    if not isinstance(data, dict):
        data = {}
    defaults = _default_data_quality_settings()
    model = str(values.get("model", "") or defaults["model"]).strip() or defaults["model"]
    try:
        temperature = float(values.get("temperature", defaults["temperature"]))
    except (TypeError, ValueError):
        temperature = defaults["temperature"]
    temperature = max(0.0, min(temperature, 2.0))
    try:
        threshold = float(values.get("confidence_threshold", defaults["confidence_threshold"]))
    except (TypeError, ValueError):
        threshold = defaults["confidence_threshold"]
    threshold = max(0.0, min(threshold, 1.0))
    try:
        max_search = int(values.get("max_search_results", defaults["max_search_results"]))
    except (TypeError, ValueError):
        max_search = defaults["max_search_results"]
    max_search = max(1, min(max_search, 100))
    try:
        batch_size = int(values.get("batch_size", defaults["batch_size"]))
    except (TypeError, ValueError):
        batch_size = defaults["batch_size"]
    batch_size = max(1, batch_size)
    try:
        parallel_checks = int(values.get("parallel_checks", defaults["parallel_checks"]))
    except (TypeError, ValueError):
        parallel_checks = defaults["parallel_checks"]
    parallel_checks = max(1, min(parallel_checks, 50))
    dry_run = bool(values.get("dry_run"))
    language = str(values.get("search_language", defaults["search_language"]) or "").strip() or defaults["search_language"]
    region = str(values.get("search_region", defaults["search_region"]) or "").strip() or defaults["search_region"]
    allowed_raw = values.get("allowed_domains", "")
    if isinstance(allowed_raw, str):
        allowed_domains = [item.strip() for item in allowed_raw.splitlines() if item.strip()]
    elif isinstance(allowed_raw, list):
        allowed_domains = [str(item).strip() for item in allowed_raw if str(item).strip()]
    else:
        allowed_domains = []
    blocked_raw = values.get("blocked_domains", "")
    if isinstance(blocked_raw, str):
        blocked_domains = [item.strip() for item in blocked_raw.splitlines() if item.strip()]
    elif isinstance(blocked_raw, list):
        blocked_domains = [str(item).strip() for item in blocked_raw if str(item).strip()]
    else:
        blocked_domains = []
    data["data_quality_settings"] = {
        "model": model,
        "temperature": temperature,
        "confidence_threshold": threshold,
        "max_search_results": max_search,
        "batch_size": batch_size,
        "parallel_checks": parallel_checks,
        "dry_run": dry_run,
        "search_language": language,
        "search_region": region,
        "allowed_domains": allowed_domains,
        "blocked_domains": blocked_domains,
    }
    save_settings_data(data)
    return get_data_quality_settings()


def _default_crawl_settings() -> Dict[str, int]:
    return {
        "max_pages": MAX_PAGES_DEFAULT,
        "concurrency": DEFAULT_CONCURRENCY,
    }


def get_crawl_defaults() -> Dict[str, int]:
    data = load_settings_data()
    defaults = _default_crawl_settings()
    result = dict(defaults)
    if isinstance(data, dict):
        stored = data.get("crawl_defaults")
        if isinstance(stored, dict):
            max_pages = stored.get("max_pages")
            if isinstance(max_pages, int):
                result["max_pages"] = max(1, min(max_pages, MAX_MAX_PAGES))
            concurrency = stored.get("concurrency")
            if isinstance(concurrency, int):
                result["concurrency"] = max(1, min(concurrency, MAX_CONCURRENCY))
    return result


def update_crawl_defaults(values: Dict[str, object]) -> Dict[str, int]:
    current = load_settings_data()
    defaults = _default_crawl_settings()

    try:
        max_pages = int(values.get("max_pages", defaults["max_pages"]))
    except (TypeError, ValueError):
        max_pages = defaults["max_pages"]
    max_pages = max(1, min(max_pages, MAX_MAX_PAGES))

    try:
        concurrency = int(values.get("concurrency", defaults["concurrency"]))
    except (TypeError, ValueError):
        concurrency = defaults["concurrency"]
    concurrency = max(1, min(concurrency, MAX_CONCURRENCY))

    current.setdefault("crawl_defaults", {})
    current["crawl_defaults"] = {
        "max_pages": max_pages,
        "concurrency": concurrency,
    }
    save_settings_data(current)
    return current["crawl_defaults"]


def _read_synonym_cache_unlocked() -> Dict[str, object]:
    data = storage.get_json("synonym_cache", "entries", {})
    if isinstance(data, dict):
        return data
    return {}


def _write_synonym_cache_unlocked(data: Dict[str, object]) -> None:
    storage.set_json("synonym_cache", "entries", data)


def load_synonym_cache() -> Dict[str, object]:
    with synonym_cache_lock:
        return dict(_read_synonym_cache_unlocked())


def update_synonym_cache_entry(key: str, synonyms: List[str]) -> None:
    with synonym_cache_lock:
        cache = _read_synonym_cache_unlocked()
        cache[key] = {
            "synonyms": list(synonyms),
            "updated_at": datetime.utcnow().isoformat() + "Z",
        }
        _write_synonym_cache_unlocked(cache)


def replace_synonym_cache(data: Dict[str, object]) -> None:
    with synonym_cache_lock:
        payload = data if isinstance(data, dict) else {}
        _write_synonym_cache_unlocked(payload)


def get_synonyms_from_cache(key: str) -> Optional[List[str]]:
    cache = load_synonym_cache()
    entry = cache.get(key) if isinstance(cache, dict) else None
    if isinstance(entry, dict):
        values = entry.get("synonyms")
        if isinstance(values, list):
            result: List[str] = []
            for item in values:
                text = str(item).strip()
                if text:
                    result.append(text)
            return result
    return None


def synonym_cache_key(keyword: str, prompt: str, model: str, temperature: float) -> str:
    payload = {
        "keyword": keyword.strip().lower(),
        "prompt": prompt.strip(),
        "model": model.strip(),
        "temperature": round(float(temperature), 3),
    }
    return json.dumps(payload, ensure_ascii=False, sort_keys=True)


def _read_keyword_finder_cache_unlocked() -> Dict[str, object]:
    data = storage.get_json("keyword_finder_cache", "entries", {})
    if isinstance(data, dict):
        return data
    return {}


def _write_keyword_finder_cache_unlocked(data: Dict[str, object]) -> None:
    storage.set_json("keyword_finder_cache", "entries", data)


def load_keyword_finder_cache() -> Dict[str, object]:
    with keyword_finder_cache_lock:
        return dict(_read_keyword_finder_cache_unlocked())


def update_keyword_finder_cache_entry(key: str, payload: Dict[str, object]) -> None:
    with keyword_finder_cache_lock:
        cache = _read_keyword_finder_cache_unlocked()
        cache[key] = {
            "data": payload,
            "updated_at": datetime.utcnow().isoformat() + "Z",
        }
        _write_keyword_finder_cache_unlocked(cache)


def replace_keyword_finder_cache(data: Dict[str, object]) -> None:
    with keyword_finder_cache_lock:
        payload = data if isinstance(data, dict) else {}
        _write_keyword_finder_cache_unlocked(payload)


def get_keyword_finder_cache_entry(key: str) -> Optional[Dict[str, object]]:
    cache = load_keyword_finder_cache()
    entry = cache.get(key) if isinstance(cache, dict) else None
    if isinstance(entry, dict):
        data = entry.get("data")
        if isinstance(data, dict):
            return data
    return None


def _read_audit_log_unlocked() -> List[Dict[str, Any]]:
    return [entry for entry in storage.read_audit_log() if isinstance(entry, dict)]


def append_audit_event(action: str, details: Dict[str, Any]) -> None:
    entry = {
        "timestamp": datetime.utcnow().isoformat() + "Z",
        "action": action,
        "details": details,
    }
    with audit_log_lock:
        storage.append_audit_log(entry)


def record_audit_event(action: str, details: Dict[str, Any]) -> None:
    sanitized = {}
    for key, value in details.items():
        if isinstance(value, (str, int, float, bool)) or value is None:
            sanitized[key] = value
        elif isinstance(value, (list, dict)):
            sanitized[key] = value
        else:
            sanitized[key] = str(value)
    append_audit_event(action, sanitized)


def _sanitize_log_value(value: Any, depth: int = 0) -> Any:
    if depth > 3:
        return str(value)
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    if isinstance(value, dict):
        result: Dict[str, Any] = {}
        for key, item in value.items():
            key_str = str(key)
            if key_str.lower() in {"authorization", "cookie", "set-cookie"}:
                result[key_str] = "<redacted>"
            else:
                result[key_str] = _sanitize_log_value(item, depth + 1)
        return result
    if isinstance(value, (list, tuple, set)):
        return [_sanitize_log_value(item, depth + 1) for item in list(value)[:20]]
    return str(value)


def log_error_event(category: str, message: str, *, details: Optional[Dict[str, Any]] = None, stack: Optional[str] = None) -> None:
    entry: Dict[str, Any] = {
        "timestamp": datetime.utcnow().isoformat() + "Z",
        "category": str(category or "error"),
        "message": str(message or "Unbekannter Fehler"),
    }
    if details:
        entry["details"] = _sanitize_log_value(details)
    if stack:
        entry["stack"] = stack if isinstance(stack, str) else str(stack)
    with error_log_lock:
        storage.append_error_log(entry)


def log_developer_event(category: str, message: str, *, details: Optional[Dict[str, Any]] = None) -> None:
    if not is_developer_mode_enabled():
        return
    entry: Dict[str, Any] = {
        "timestamp": datetime.utcnow().isoformat() + "Z",
        "category": f"developer:{category}",
        "message": str(message or ""),
    }
    if details is not None:
        entry["details"] = _sanitize_log_value(details)
    with error_log_lock:
        storage.append_error_log(entry)


def list_error_log_entries(limit: Optional[int] = 200) -> List[Dict[str, Any]]:
    with error_log_lock:
        return list(storage.list_error_log(limit))


def replace_error_log_entries(entries: Iterable[Dict[str, Any]]) -> None:
    with error_log_lock:
        storage.replace_error_log(entries)


def capture_request_exception(sender, exception, **extra) -> None:  # pragma: no cover - defensive logging
    if isinstance(exception, HTTPException) and getattr(exception, "code", 500) < 500:
        return
    details: Dict[str, Any] = {"type": type(exception).__name__}
    try:
        details.update(
            {
                "method": request.method,
                "path": request.path,
                "query_string": request.query_string.decode("utf-8", "ignore"),
                "remote_addr": request.remote_addr,
                "user_agent": request.headers.get("User-Agent"),
                "status_code": getattr(exception, "code", None),
                "content_length": request.content_length,
            }
        )
    except RuntimeError:
        details["request"] = "außerhalb des Request-Kontexts"
    stack = "".join(traceback.format_exception(type(exception), exception, exception.__traceback__))
    log_error_event("flask-request", str(exception), details=details, stack=stack)


got_request_exception.connect(capture_request_exception, app)


@app.before_request
def log_incoming_request() -> None:
    if not is_developer_mode_enabled():
        return
    path = request.path or ""
    if path.startswith("/static/"):
        return
    payload: Optional[Dict[str, Any]] = None
    if request.method in {"POST", "PUT", "PATCH"}:
        data = request.get_json(silent=True)
        if isinstance(data, dict):
            payload = data
        elif request.form:
            payload = {key: request.form.getlist(key) for key in request.form.keys()}
    log_developer_event(
        "request",
        f"{request.method} {path}",
        details={
            "path": path,
            "method": request.method,
            "query": request.args.to_dict(flat=False),
            "payload": payload,
            "remote_addr": request.remote_addr,
        },
    )
    try:
        request.environ["developer_mode_started_at"] = time.time()
    except Exception:
        request.environ["developer_mode_started_at"] = None


@app.after_request
def log_outgoing_response(response: Response) -> Response:
    if is_developer_mode_enabled():
        path = request.path or ""
        if not path.startswith("/static/"):
            started_at = request.environ.get("developer_mode_started_at")
            duration = None
            if isinstance(started_at, (int, float)):
                duration = max(0.0, time.time() - float(started_at))
            log_developer_event(
                "response",
                f"{request.method} {path}",
                details={
                    "status": response.status_code,
                    "duration": duration,
                    "content_length": response.calculate_content_length(),
                },
            )
    return response


def _cleanup_import_sessions_locked() -> None:
    now = time.time()
    expired = [
        key
        for key, payload in _import_sessions.items()
        if now - payload.get("created", 0.0) > IMPORT_SESSION_TTL_SECONDS
    ]
    for key in expired:
        _import_sessions.pop(key, None)
    if len(_import_sessions) > IMPORT_SESSION_LIMIT:
        sorted_items = sorted(
            _import_sessions.items(), key=lambda item: item[1].get("created", now)
        )
        for key, _ in sorted_items[:-IMPORT_SESSION_LIMIT]:
            _import_sessions.pop(key, None)


def create_import_session(payload: Dict[str, Any], summary: Dict[str, Any], contains_sensitive: bool) -> str:
    session_id = str(uuid.uuid4())
    with import_sessions_lock:
        _cleanup_import_sessions_locked()
        _import_sessions[session_id] = {
            "payload": payload,
            "summary": summary,
            "contains_sensitive": contains_sensitive,
            "created": time.time(),
        }
    return session_id


def get_import_session(session_id: str) -> Optional[Dict[str, Any]]:
    with import_sessions_lock:
        data = _import_sessions.get(session_id)
        if not data:
            return None
        return dict(data)


def consume_import_session(session_id: str) -> Optional[Dict[str, Any]]:
    with import_sessions_lock:
        data = _import_sessions.pop(session_id, None)
        return dict(data) if data else None


def _stammdaten_key(record: Dict[str, Any]) -> str:
    if not isinstance(record, dict):
        return ""
    jahr = record.get("jahr")
    if jahr is None:
        jahr = record.get("Jahr")
    schul_id = (
        record.get("schul_id")
        or record.get("Schul_ID")
        or record.get("schulId")
        or record.get("id")
    )
    if schul_id is None:
        return ""
    try:
        jahr_value = str(jahr).strip()
    except Exception:
        jahr_value = ""
    try:
        schul_value = str(schul_id).strip()
    except Exception:
        schul_value = ""
    if not schul_value:
        return ""
    return f"{jahr_value}|{schul_value}".strip("|")


def _saved_search_signature(entry: Dict[str, Any]) -> str:
    if not isinstance(entry, dict):
        return ""
    job_id = entry.get("job_id")
    if job_id:
        return f"job:{job_id}"
    saved_at = entry.get("saved_at")
    start_urls = entry.get("start_urls")
    if isinstance(start_urls, list):
        start_urls_value = ",".join(sorted(str(url) for url in start_urls))
    else:
        start_urls_value = ""
    return f"sig:{saved_at or ''}:{start_urls_value}"


def _git_metadata() -> Dict[str, str]:
    metadata: Dict[str, str] = {
        "python_version": platform.python_version(),
    }
    try:
        branch = (
            subprocess.check_output(
                ["git", "rev-parse", "--abbrev-ref", "HEAD"],
                stderr=subprocess.DEVNULL,
                text=True,
            )
            .strip()
        )
    except (subprocess.SubprocessError, OSError):
        branch = ""
    if branch:
        metadata["git_branch"] = branch
    try:
        commit = (
            subprocess.check_output(
                ["git", "rev-parse", "HEAD"], stderr=subprocess.DEVNULL, text=True
            )
            .strip()
        )
    except (subprocess.SubprocessError, OSError):
        commit = ""
    if commit:
        metadata["git_commit"] = commit
    return metadata


def summarize_section(section: str, data: Any) -> Dict[str, Any]:
    base = {
        "label": EXPORT_SECTION_DEFINITIONS.get(section, {}).get("label", section),
    }
    if section == "stammdaten":
        base["entries"] = len(data) if isinstance(data, list) else 0
    elif section == "keywords":
        saved = data.get("saved_searches") if isinstance(data, dict) else []
        syn_cache = data.get("synonym_cache") if isinstance(data, dict) else {}
        finder_cache = data.get("keyword_finder_cache") if isinstance(data, dict) else {}
        base.update(
            {
                "saved_searches": len(saved) if isinstance(saved, list) else 0,
                "synonym_cache_entries": len(syn_cache)
                if isinstance(syn_cache, dict)
                else 0,
                "keyword_finder_cache_entries": len(finder_cache)
                if isinstance(finder_cache, dict)
                else 0,
            }
        )
    elif section == "settings":
        base["keys"] = len(data) if isinstance(data, dict) else 0
    elif section == "searches":
        definitions = []
        runs = []
        results = {}
        if isinstance(data, dict):
            if isinstance(data.get("definitions"), list):
                definitions = data["definitions"]
            if isinstance(data.get("runs"), list):
                runs = data["runs"]
            if isinstance(data.get("results"), dict):
                results = data["results"]
        base.update(
            {
                "definitions": len(definitions),
                "runs": len(runs),
                "result_sets": len(results),
            }
        )
    elif section == "data_quality":
        records = None
        if isinstance(data, dict):
            if "records" in data and isinstance(data["records"], dict):
                records = data["records"]
            elif "data" in data and isinstance(data["data"], dict):
                records = data["data"]
            elif isinstance(data, dict):
                records = data
        if isinstance(records, dict):
            base["records"] = len(records)
        else:
            base["records"] = 0
    elif section == "debug":
        if isinstance(data, list):
            base["entries"] = len(data)
        elif isinstance(data, dict) and "logs" in data and isinstance(data["logs"], list):
            base["entries"] = len(data["logs"])
        else:
            base["entries"] = 0
    elif section == "api_keys":
        openai_key = ""
        keyword_key = ""
        google_search = {}
        if isinstance(data, dict):
            openai_key = str(data.get("openai") or "")
            keyword_key = str(data.get("google_keyword_planner") or "")
            google_search = data.get("google_search") if isinstance(data.get("google_search"), dict) else {}
        base.update(
            {
                "openai": bool(openai_key),
                "google_keyword_planner": bool(keyword_key),
                "google_search_api_key": bool(google_search.get("api_key"))
                if isinstance(google_search, dict)
                else False,
                "google_search_cx": bool(google_search.get("cx"))
                if isinstance(google_search, dict)
                else False,
            }
        )
    return base


def build_export_payload(selected_sections: Set[str]) -> Tuple[Dict[str, Any], Dict[str, Any], bool]:
    requested = {
        section
        for section in selected_sections
        if section in EXPORT_SECTION_DEFINITIONS
    }
    if not requested:
        requested = set(EXPORT_SECTION_ORDER)

    payload: Dict[str, Any] = {
        "schema_version": EXPORT_SCHEMA_VERSION,
        "generated_at": datetime.utcnow().isoformat() + "Z",
        "meta": _git_metadata(),
        "sections": {},
    }
    summary: Dict[str, Any] = {
        "schema_version": EXPORT_SCHEMA_VERSION,
        "generated_at": payload["generated_at"],
        "meta": payload["meta"],
        "sections": {},
    }
    contains_sensitive = False

    if "stammdaten" in requested:
        stammdaten_data = load_stammdaten()
        payload["sections"]["stammdaten"] = stammdaten_data
        summary["sections"]["stammdaten"] = summarize_section("stammdaten", stammdaten_data)

    if "keywords" in requested:
        keywords_payload = {
            "saved_searches": load_saved_searches(),
            "synonym_cache": load_synonym_cache(),
            "keyword_finder_cache": load_keyword_finder_cache(),
        }
        payload["sections"]["keywords"] = keywords_payload
        summary["sections"]["keywords"] = summarize_section("keywords", keywords_payload)

    if "searches" in requested:
        searches_payload = collect_all_search_data()
        payload["sections"]["searches"] = searches_payload
        summary["sections"]["searches"] = summarize_section("searches", searches_payload)

    if "settings" in requested:
        settings_payload = load_settings_data()
        payload["sections"]["settings"] = settings_payload
        summary["sections"]["settings"] = summarize_section("settings", settings_payload)

    if "data_quality" in requested:
        dq_payload = {"records": load_data_quality_results()}
        payload["sections"]["data_quality"] = dq_payload
        summary["sections"]["data_quality"] = summarize_section("data_quality", dq_payload)

    if "debug" in requested:
        debug_payload = list_error_log_entries(limit=None)
        payload["sections"]["debug"] = debug_payload
        summary["sections"]["debug"] = summarize_section("debug", debug_payload)

    if "api_keys" in requested:
        api_payload = {
            "openai": get_api_key() or "",
            "google_keyword_planner": get_keyword_planner_key() or "",
            "google_search": get_google_search_credentials(),
        }
        payload["sections"]["api_keys"] = api_payload
        api_summary = summarize_section("api_keys", api_payload)
        summary["sections"]["api_keys"] = api_summary
        contains_sensitive = any(
            bool(api_summary.get(flag))
            for flag in [
                "openai",
                "google_keyword_planner",
                "google_search_api_key",
                "google_search_cx",
            ]
        )

    summary["requested_sections"] = sorted(requested)
    return payload, summary, contains_sensitive


def analyze_import_payload(payload: Dict[str, Any]) -> Tuple[Dict[str, Any], bool]:
    if not isinstance(payload, dict):
        raise ValueError("Ungültiges Exportformat")
    schema_version = payload.get("schema_version")
    if schema_version != EXPORT_SCHEMA_VERSION:
        raise ValueError("Die Exportdatei verwendet eine inkompatible Version.")
    sections = payload.get("sections")
    if not isinstance(sections, dict):
        raise ValueError("Die Exportdatei enthält keine Bereiche.")

    summary: Dict[str, Any] = {
        "schema_version": schema_version,
        "generated_at": payload.get("generated_at"),
        "meta": payload.get("meta", {}),
        "sections": {},
    }
    contains_sensitive = False

    for key, definition in EXPORT_SECTION_DEFINITIONS.items():
        if key not in sections:
            continue
        section_summary = summarize_section(key, sections[key])
        summary["sections"][key] = section_summary
        if key == "api_keys":
            contains_sensitive = any(
                bool(section_summary.get(flag))
                for flag in [
                    "openai",
                    "google_keyword_planner",
                    "google_search_api_key",
                    "google_search_cx",
                ]
            )

    summary["available_sections"] = sorted(summary["sections"].keys())
    summary["contains_sensitive"] = contains_sensitive
    return summary, contains_sensitive


def apply_stammdaten_import(data: Any, mode: str, dry_run: bool) -> Dict[str, Any]:
    if not isinstance(data, list):
        return {"status": "skipped", "reason": "Ungültiges Format"}
    current = load_stammdaten()
    if mode == "replace":
        if not dry_run:
            save_stammdaten(data)
        return {
            "status": "replaced",
            "previous": len(current),
            "new_total": len(data),
        }

    index: Dict[str, Dict[str, Any]] = {}
    for record in current:
        key = _stammdaten_key(record)
        if key:
            index[key] = dict(record)

    added = 0
    updated = 0
    skipped = 0
    for record in data:
        key = _stammdaten_key(record)
        if not key:
            skipped += 1
            continue
        existing = index.get(key)
        if existing is None:
            index[key] = dict(record)
            added += 1
        elif existing != record:
            index[key] = dict(record)
            updated += 1
        else:
            skipped += 1

    merged: List[Dict[str, Any]] = []
    seen_keys: Set[str] = set()
    for record in current:
        key = _stammdaten_key(record)
        if key and key in index:
            merged.append(index[key])
            seen_keys.add(key)
    for key, record in index.items():
        if key not in seen_keys:
            merged.append(record)

    if not dry_run and (added or updated):
        save_stammdaten(merged)

    return {
        "status": "merged",
        "added": added,
        "updated": updated,
        "unchanged": skipped,
        "total": len(index),
    }


def apply_settings_import(data: Any, mode: str, dry_run: bool) -> Dict[str, Any]:
    if not isinstance(data, dict):
        return {"status": "skipped", "reason": "Ungültiges Format"}
    current = load_settings_data()
    if mode == "replace":
        if not dry_run:
            save_settings_data(data)
        return {
            "status": "replaced",
            "keys": len(data),
        }

    merged = dict(current)
    updated = 0
    for key, value in data.items():
        if merged.get(key) != value:
            merged[key] = value
            updated += 1
    if not dry_run and updated:
        save_settings_data(merged)
    return {
        "status": "merged",
        "updated_keys": updated,
        "total_keys": len(merged),
    }


def apply_saved_search_merge(existing: List[Dict[str, Any]], incoming: List[Dict[str, Any]]) -> Tuple[List[Dict[str, Any]], Dict[str, int]]:
    signatures = { _saved_search_signature(entry): entry for entry in existing }
    stats = {"added": 0, "skipped": 0}
    merged = list(existing)
    for entry in incoming:
        sig = _saved_search_signature(entry)
        if not sig:
            sig = f"raw:{json.dumps(entry, ensure_ascii=False, sort_keys=True)}"
        if sig in signatures:
            stats["skipped"] += 1
            continue
        signatures[sig] = entry
        merged.append(entry)
        stats["added"] += 1
    return merged, stats


def apply_keywords_import(data: Any, mode: str, dry_run: bool) -> Dict[str, Any]:
    if not isinstance(data, dict):
        return {"status": "skipped", "reason": "Ungültiges Format"}

    incoming_saved = data.get("saved_searches") if isinstance(data.get("saved_searches"), list) else []
    incoming_synonyms = data.get("synonym_cache") if isinstance(data.get("synonym_cache"), dict) else {}
    incoming_finder = data.get("keyword_finder_cache") if isinstance(data.get("keyword_finder_cache"), dict) else {}

    result: Dict[str, Any] = {}

    if mode == "replace":
        if not dry_run:
            replace_saved_searches(list(incoming_saved))
            replace_synonym_cache(dict(incoming_synonyms))
            replace_keyword_finder_cache(dict(incoming_finder))
        result.update(
            {
                "status": "replaced",
                "saved_searches": len(incoming_saved),
                "synonym_cache_entries": len(incoming_synonyms),
                "keyword_finder_cache_entries": len(incoming_finder),
            }
        )
        return result

    current_saved = load_saved_searches()
    current_synonyms = load_synonym_cache()
    current_finder = load_keyword_finder_cache()

    merged_saved, saved_stats = apply_saved_search_merge(current_saved, list(incoming_saved))
    synonym_updates = 0
    for key, value in incoming_synonyms.items():
        if current_synonyms.get(key) != value:
            current_synonyms[key] = value
            synonym_updates += 1
    finder_updates = 0
    for key, value in incoming_finder.items():
        if current_finder.get(key) != value:
            current_finder[key] = value
            finder_updates += 1

    if not dry_run and saved_stats["added"]:
        replace_saved_searches(merged_saved)
    if not dry_run and synonym_updates:
        replace_synonym_cache(current_synonyms)
    if not dry_run and finder_updates:
        replace_keyword_finder_cache(current_finder)

    result.update(
        {
            "status": "merged",
            "saved_searches_added": saved_stats["added"],
            "saved_searches_skipped": saved_stats["skipped"],
            "synonym_updates": synonym_updates,
            "keyword_finder_updates": finder_updates,
        }
    )
    return result


def apply_searches_import(data: Any, mode: str, dry_run: bool) -> Dict[str, Any]:
    if not isinstance(data, dict):
        return {"status": "skipped", "reason": "Ungültiges Format"}

    definitions = data.get("definitions") if isinstance(data.get("definitions"), list) else []
    runs = data.get("runs") if isinstance(data.get("runs"), list) else []
    raw_results = data.get("results") if isinstance(data.get("results"), dict) else {}
    raw_logs = data.get("logs") if isinstance(data.get("logs"), dict) else {}

    normalised_results: Dict[str, List[Any]] = {}
    for run_id, entries in raw_results.items():
        run_id_str = str(run_id or "").strip()
        if not run_id_str or not isinstance(entries, list):
            continue
        normalised_results[run_id_str] = entries

    normalised_logs: Dict[str, List[Any]] = {}
    for run_id, entries in raw_logs.items():
        run_id_str = str(run_id or "").strip()
        if not run_id_str or not isinstance(entries, list):
            continue
        normalised_logs[run_id_str] = entries

    def_count = len(definitions)
    run_count = len(runs)
    result_entries = sum(len(value) for value in normalised_results.values())
    log_entries = sum(len(value) for value in normalised_logs.values())

    if mode == "replace":
        if not dry_run:
            existing_ids = [str(item.get("id")) for item in load_search_definitions() if str(item.get("id"))]
            if existing_ids:
                delete_search_definitions(existing_ids)
            for definition in definitions:
                if isinstance(definition, dict) and definition.get("id"):
                    save_search_definition(definition)
            for run in runs:
                if isinstance(run, dict) and run.get("id") and run.get("definition_id"):
                    save_search_run(run)
            run_ids = set(normalised_results.keys()) | set(normalised_logs.keys())
            for run_id_str in run_ids:
                clear_search_results(run_id_str)
                cleaned_results = [entry for entry in normalised_results.get(run_id_str, []) if isinstance(entry, dict)]
                if cleaned_results:
                    append_search_results(run_id_str, cleaned_results)
                cleaned_logs = [entry for entry in normalised_logs.get(run_id_str, []) if isinstance(entry, dict)]
                if cleaned_logs:
                    append_search_logs(run_id_str, cleaned_logs)
        return {
            "status": "replaced",
            "definitions": def_count,
            "runs": run_count,
            "results": result_entries,
            "logs": log_entries,
        }

    existing_definitions = {str(item.get("id")): item for item in load_search_definitions() if str(item.get("id"))}
    added_defs = 0
    updated_defs = 0
    if not dry_run:
        for definition in definitions:
            if not isinstance(definition, dict):
                continue
            identifier = str(definition.get("id") or "").strip()
            if not identifier:
                continue
            if identifier not in existing_definitions:
                added_defs += 1
            elif existing_definitions[identifier] != definition:
                updated_defs += 1
            save_search_definition(definition)
            existing_definitions[identifier] = definition
    else:
        for definition in definitions:
            if not isinstance(definition, dict):
                continue
            identifier = str(definition.get("id") or "").strip()
            if not identifier:
                continue
            if identifier not in existing_definitions:
                added_defs += 1
            elif existing_definitions[identifier] != definition:
                updated_defs += 1
            existing_definitions[identifier] = definition

    existing_runs = {str(item.get("id")): item for item in list_search_runs() if str(item.get("id"))}
    added_runs = 0
    updated_runs = 0
    if not dry_run:
        for run in runs:
            if not isinstance(run, dict):
                continue
            run_id = str(run.get("id") or "").strip()
            definition_id = str(run.get("definition_id") or "").strip()
            if not run_id or not definition_id:
                continue
            if run_id not in existing_runs:
                added_runs += 1
            elif existing_runs[run_id] != run:
                updated_runs += 1
            save_search_run(run)
            existing_runs[run_id] = run
    else:
        for run in runs:
            if not isinstance(run, dict):
                continue
            run_id = str(run.get("id") or "").strip()
            definition_id = str(run.get("definition_id") or "").strip()
            if not run_id or not definition_id:
                continue
            if run_id not in existing_runs:
                added_runs += 1
            elif existing_runs[run_id] != run:
                updated_runs += 1
            existing_runs[run_id] = run

    imported_results = 0
    imported_logs = 0
    processed_for_logs: Set[str] = set()
    for run_id_str, entries in normalised_results.items():
        cleaned_entries = [entry for entry in entries if isinstance(entry, dict)]
        imported_results += len(cleaned_entries)
        if not dry_run:
            clear_search_results(run_id_str)
            if cleaned_entries:
                append_search_results(run_id_str, cleaned_entries)
            cleaned_logs = [entry for entry in normalised_logs.get(run_id_str, []) if isinstance(entry, dict)]
            if cleaned_logs:
                append_search_logs(run_id_str, cleaned_logs)
                imported_logs += len(cleaned_logs)
        else:
            imported_logs += len([entry for entry in normalised_logs.get(run_id_str, []) if isinstance(entry, dict)])
        processed_for_logs.add(run_id_str)

    for run_id_str, entries in normalised_logs.items():
        if run_id_str in processed_for_logs:
            continue
        cleaned_logs = [entry for entry in entries if isinstance(entry, dict)]
        if not cleaned_logs:
            continue
        if not dry_run:
            clear_search_logs(run_id_str)
            append_search_logs(run_id_str, cleaned_logs)
        imported_logs += len(cleaned_logs)

    return {
        "status": "merged",
        "definitions_added": added_defs,
        "definitions_updated": updated_defs,
        "runs_added": added_runs,
        "runs_updated": updated_runs,
        "results_updated": imported_results,
        "logs_updated": imported_logs,
    }


def apply_data_quality_import(data: Any, mode: str, dry_run: bool) -> Dict[str, Any]:
    records = None
    if isinstance(data, dict):
        if "records" in data and isinstance(data["records"], dict):
            records = data["records"]
        elif "data" in data and isinstance(data["data"], dict):
            records = data["data"]
        elif all(isinstance(key, str) for key in data.keys()):
            records = data
    if not isinstance(records, dict):
        return {"status": "skipped", "reason": "Ungültiges Format"}

    current = load_data_quality_results()
    if mode == "replace":
        if not dry_run:
            replace_data_quality_results(records)
        return {
            "status": "replaced",
            "records": len(records),
        }

    merged = dict(current)
    added = 0
    updated = 0
    for key, value in records.items():
        if key not in merged:
            merged[key] = value
            added += 1
        elif merged[key] != value:
            merged[key] = value
            updated += 1
    if not dry_run and (added or updated):
        replace_data_quality_results(merged)
    return {
        "status": "merged",
        "added": added,
        "updated": updated,
        "total": len(merged),
    }


def apply_api_keys_import(data: Any, mode: str, dry_run: bool) -> Dict[str, Any]:
    if not isinstance(data, dict):
        return {"status": "skipped", "reason": "Ungültiges Format"}
    incoming_openai = str(data.get("openai") or "").strip()
    incoming_planner = str(data.get("google_keyword_planner") or "").strip()
    incoming_search = data.get("google_search") if isinstance(data.get("google_search"), dict) else {}
    incoming_search_api = str(incoming_search.get("api_key") or "").strip()
    incoming_search_cx = str(incoming_search.get("cx") or "").strip()

    current_openai = get_api_key() or ""
    current_planner = get_keyword_planner_key() or ""
    current_search = get_google_search_credentials()

    updated = {"openai": False, "keyword_planner": False, "google_search": False}

    if mode == "replace":
        if not dry_run:
            set_api_key(incoming_openai)
            set_keyword_planner_key(incoming_planner)
            update_google_search_credentials(incoming_search_api, incoming_search_cx)
        updated = {
            "openai": incoming_openai != current_openai,
            "keyword_planner": incoming_planner != current_planner,
            "google_search": (
                incoming_search_api != current_search.get("api_key")
                or incoming_search_cx != current_search.get("cx")
            ),
        }
        return {
            "status": "replaced",
            "changes": updated,
        }

    if incoming_openai and not current_openai:
        if not dry_run:
            set_api_key(incoming_openai)
        updated["openai"] = True
    if incoming_planner and not current_planner:
        if not dry_run:
            set_keyword_planner_key(incoming_planner)
        updated["keyword_planner"] = True
    if (incoming_search_api or incoming_search_cx) and (
        not current_search.get("api_key") or not current_search.get("cx")
    ):
        if not dry_run:
            update_google_search_credentials(
                incoming_search_api or current_search.get("api_key", ""),
                incoming_search_cx or current_search.get("cx", ""),
            )
        updated["google_search"] = True

    return {
        "status": "merged",
        "changes": updated,
    }


def apply_debug_import(data: Any, mode: str, dry_run: bool) -> Dict[str, Any]:
    if isinstance(data, dict) and "logs" in data:
        payload = data.get("logs")
    else:
        payload = data
    if not isinstance(payload, list):
        return {"status": "skipped", "reason": "Ungültiges Format"}
    entries = [entry for entry in payload if isinstance(entry, dict)]
    if dry_run:
        return {"status": "preview", "entries": len(entries)}
    if mode == "replace":
        replace_error_log_entries(entries)
        return {"status": "replaced", "entries": len(entries)}
    if mode == "merge" and entries:
        with error_log_lock:
            for entry in entries:
                storage.append_error_log(entry)
        return {"status": "merged", "entries": len(entries)}
    return {"status": "skipped", "reason": "Keine verwertbaren Einträge"}


def apply_import(payload: Dict[str, Any], actions: Dict[str, str], dry_run: bool) -> Dict[str, Any]:
    sections = payload.get("sections")
    if not isinstance(sections, dict):
        raise ValueError("Exportpaket unvollständig")

    results: Dict[str, Any] = {}

    for section in EXPORT_SECTION_DEFINITIONS.keys():
        data = sections.get(section)
        if data is None:
            continue
        mode = actions.get(section, "ignore")
        if mode == "ignore":
            results[section] = {"status": "ignored"}
            continue
        if section == "stammdaten":
            results[section] = apply_stammdaten_import(data, mode, dry_run)
        elif section == "keywords":
            results[section] = apply_keywords_import(data, mode, dry_run)
        elif section == "searches":
            results[section] = apply_searches_import(data, mode, dry_run)
        elif section == "settings":
            results[section] = apply_settings_import(data, mode, dry_run)
        elif section == "data_quality":
            results[section] = apply_data_quality_import(data, mode, dry_run)
        elif section == "debug":
            results[section] = apply_debug_import(data, mode, dry_run)
        elif section == "api_keys":
            results[section] = apply_api_keys_import(data, mode, dry_run)
        else:
            results[section] = {"status": "skipped"}

    return {"results": results, "dry_run": dry_run}


def keyword_finder_cache_key(
    keyword: str,
    prompt: str,
    model: str,
    temperature: float,
    max_results: int,
) -> str:
    payload = {
        "keyword": keyword.strip().lower(),
        "prompt": prompt.strip(),
        "model": model.strip(),
        "temperature": round(float(temperature), 3),
        "max_results": int(max_results),
    }
    return json.dumps(payload, ensure_ascii=False, sort_keys=True)


def set_api_key(value: str) -> None:
    sanitized = value.strip()
    with api_key_lock:
        global _api_key_value
        _api_key_value = sanitized or None
        if sanitized:
            os.environ["OPENAI_API_KEY"] = sanitized
            storage.set_json("api_keys", "openai", {"value": sanitized})
        else:
            os.environ.pop("OPENAI_API_KEY", None)
            storage.delete("api_keys", ["openai"])


def clear_api_key() -> None:
    with api_key_lock:
        global _api_key_value
        _api_key_value = None
        os.environ.pop("OPENAI_API_KEY", None)
        storage.delete("api_keys", ["openai"])


def get_api_key() -> Optional[str]:
    with api_key_lock:
        env_value = os.environ.get("OPENAI_API_KEY", "").strip()
        if env_value:
            global _api_key_value
            _api_key_value = env_value
            return env_value
        if _api_key_value:
            return _api_key_value
        stored = storage.get_json("api_keys", "openai", {})
        if isinstance(stored, dict):
            value = str(stored.get("value", "")).strip()
            if value:
                _api_key_value = value
                return value
        return None


def has_api_key() -> bool:
    return get_api_key() is not None


def set_keyword_planner_key(value: str) -> None:
    sanitized = value.strip()
    with google_api_key_lock:
        global _google_api_key_value
        _google_api_key_value = sanitized or None
        if sanitized:
            os.environ["GOOGLE_KEYWORD_PLANNER_KEY"] = sanitized
            storage.set_json("api_keys", "google_keyword_planner", {"value": sanitized})
        else:
            os.environ.pop("GOOGLE_KEYWORD_PLANNER_KEY", None)
            storage.delete("api_keys", ["google_keyword_planner"])


def clear_keyword_planner_key() -> None:
    with google_api_key_lock:
        global _google_api_key_value
        _google_api_key_value = None
        os.environ.pop("GOOGLE_KEYWORD_PLANNER_KEY", None)
        storage.delete("api_keys", ["google_keyword_planner"])


def get_keyword_planner_key() -> Optional[str]:
    with google_api_key_lock:
        env_value = os.environ.get("GOOGLE_KEYWORD_PLANNER_KEY", "").strip()
        if env_value:
            global _google_api_key_value
            _google_api_key_value = env_value
            return env_value
        if _google_api_key_value:
            return _google_api_key_value
        stored = storage.get_json("api_keys", "google_keyword_planner", {})
        if isinstance(stored, dict):
            value = str(stored.get("value", "")).strip()
            if value:
                _google_api_key_value = value
                return value
        return None


def has_keyword_planner_key() -> bool:
    return get_keyword_planner_key() is not None


def fetch_synonyms_from_openai(
    keyword: str,
    api_key: str,
    *,
    prompt: Optional[str] = None,
    model: Optional[str] = None,
    temperature: Optional[float] = None,
    max_results: Optional[int] = None,
) -> List[str]:
    """Request synonyms for a single keyword from the OpenAI API."""

    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json",
    }
    prompt_text = (prompt or "").strip() or DEFAULT_SYNONYM_PROMPT
    model_name = (model or "").strip() or OPENAI_MODEL_NAME
    try:
        temperature_value = float(temperature) if temperature is not None else DEFAULT_SYNONYM_TEMPERATURE
    except (TypeError, ValueError):
        temperature_value = DEFAULT_SYNONYM_TEMPERATURE
    max_results_value: Optional[int]
    if max_results is None:
        max_results_value = None
    else:
        try:
            max_results_value = int(max_results)
        except (TypeError, ValueError):
            max_results_value = None
    if max_results_value is not None:
        if max_results_value < 1:
            max_results_value = 1
        elif max_results_value > 150:
            max_results_value = 150
    payload = {
        "model": model_name,
        "messages": [
            {"role": "system", "content": prompt_text},
            {"role": "user", "content": f'Begriff: "{keyword}"'},
        ],
        "temperature": temperature_value,
        "max_tokens": 200,
    }

    try:
        response = requests.post(
            OPENAI_CHAT_COMPLETIONS_URL,
            headers=headers,
            json=payload,
            timeout=OPENAI_TIMEOUT,
        )
    except requests.RequestException as exc:  # pragma: no cover - network failure
        raise OpenAIIntegrationError(
            f'Synonymerweiterung für "{keyword}" nicht möglich: Keine Verbindung zur OpenAI-API.'
        ) from exc

    if response.status_code >= 400:
        raise OpenAIIntegrationError(
            f'Synonymerweiterung für "{keyword}" nicht möglich: OpenAI-API meldet Status {response.status_code}.'
        )

    try:
        data = response.json()
    except ValueError as exc:  # pragma: no cover - invalid JSON
        raise OpenAIIntegrationError(
            f'Synonymerweiterung für "{keyword}" nicht möglich: Antwort konnte nicht gelesen werden.'
        ) from exc

    content = ""
    choices = data.get("choices") if isinstance(data, dict) else None
    if choices and isinstance(choices, list) and choices:
        message = choices[0].get("message") or {}
        if isinstance(message, dict):
            content = message.get("content", "")

    if not content:
        raise OpenAIIntegrationError(
            f'Synonymerweiterung für "{keyword}" nicht möglich: Antwort enthielt keine Vorschläge.'
        )

    content = content.strip()
    try:
        parsed = json.loads(content)
    except json.JSONDecodeError:
        start_idx = content.find("[")
        end_idx = content.rfind("]")
        if start_idx == -1 or end_idx == -1 or end_idx <= start_idx:
            raise OpenAIIntegrationError(
                f'Synonymerweiterung für "{keyword}" nicht möglich: Antwort war nicht im JSON-Format.'
            )
        try:
            parsed = json.loads(content[start_idx : end_idx + 1])
        except json.JSONDecodeError as exc:
            raise OpenAIIntegrationError(
                f'Synonymerweiterung für "{keyword}" nicht möglich: Antwort war nicht im JSON-Format.'
            ) from exc

    if not isinstance(parsed, list):
        raise OpenAIIntegrationError(
            f'Synonymerweiterung für "{keyword}" nicht möglich: Antwort enthielt keine Begriffsliste.'
        )

    synonyms: List[str] = []
    seen: Set[str] = set()
    for item in parsed:
        text = str(item).strip()
        if not text:
            continue
        lowered = text.lower()
        if lowered in seen:
            continue
        seen.add(lowered)
        synonyms.append(text)
        if max_results_value is not None and len(synonyms) >= max_results_value:
            break

    return synonyms


def fetch_keyword_planner_suggestions(
    keyword: str,
    api_key: str,
    *,
    max_results: int = DEFAULT_KEYWORD_FINDER_MAX_RESULTS,
) -> List[Dict[str, object]]:
    """Fetch keyword suggestions from the Google Keyword Planner."""

    if not api_key:
        raise KeywordPlannerError("Kein Keyword-Planner-Schlüssel hinterlegt. Bitte in den Einstellungen speichern.")

    try:
        max_value = int(max_results)
    except (TypeError, ValueError):
        max_value = DEFAULT_KEYWORD_FINDER_MAX_RESULTS
    max_value = max(1, min(max_value, DEFAULT_KEYWORD_FINDER_MAX_RESULTS))

    params = {"key": api_key}
    payload = {
        "query": keyword,
        "languageCode": "de-DE",
        "maxSuggestions": max_value,
        "regionCode": "DE",
    }

    try:
        response = requests.post(
            GOOGLE_KEYWORD_API_URL,
            params=params,
            json=payload,
            timeout=GOOGLE_KEYWORD_TIMEOUT,
        )
    except requests.RequestException as exc:
        raise KeywordPlannerError("Keyword-Planner-Anfrage fehlgeschlagen (Netzwerkfehler).") from exc

    if response.status_code >= 400:
        raise KeywordPlannerError(
            f"Keyword-Planner-Anfrage fehlgeschlagen (Status {response.status_code}). Bitte Schlüssel und Berechtigungen prüfen."
        )

    try:
        data = response.json()
    except ValueError as exc:
        raise KeywordPlannerError("Keyword-Planner-Antwort konnte nicht gelesen werden.") from exc

    suggestions_raw: List[Dict[str, object]] = []
    if isinstance(data, dict):
        for key in ("suggestions", "results", "keywordIdeas", "ideas", "items"):
            values = data.get(key)
            if isinstance(values, list):
                for item in values:
                    if isinstance(item, dict):
                        suggestions_raw.append(item)
        # Some APIs nest results deeper
        if not suggestions_raw:
            keyword_ideas = data.get("keyword_ideas")
            if isinstance(keyword_ideas, list):
                for item in keyword_ideas:
                    if isinstance(item, dict):
                        suggestions_raw.append(item)

    if not suggestions_raw:
        raise KeywordPlannerError("Keyword-Planner lieferte keine Vorschläge.")

    suggestions: List[Dict[str, object]] = []
    seen: Set[str] = set()
    for item in suggestions_raw:
        term = str(
            item.get("keyword")
            or item.get("text")
            or item.get("term")
            or item.get("phrase")
            or ""
        ).strip()
        if not term:
            continue
        lowered = term.lower()
        if lowered in seen:
            continue
        seen.add(lowered)
        search_volume = item.get("searchVolume") or item.get("search_volume") or item.get("avgMonthlySearches")
        competition = item.get("competition") or item.get("competitionIndex") or item.get("competition_index")
        relevance = item.get("relevance") or item.get("relevanceScore") or item.get("relevance_score")
        suggestions.append(
            {
                "term": term,
                "search_volume": search_volume,
                "competition": competition,
                "relevance": relevance,
            }
        )
        if len(suggestions) >= max_value:
            break

    if not suggestions:
        raise KeywordPlannerError("Keyword-Planner lieferte keine verwertbaren Begriffe.")

    return suggestions


def _normalise_competition(value: object) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        if value <= 0:
            return "niedrig"
        if value < 0.5:
            return "niedrig"
        if value < 0.8:
            return "mittel"
        return "hoch"
    text = str(value).strip().lower()
    if not text:
        return None
    if text in {"low", "niedrig"}:
        return "niedrig"
    if text in {"medium", "mittel"}:
        return "mittel"
    if text in {"high", "hoch"}:
        return "hoch"
    return None


def _build_fallback_clusters(suggestions: List[Dict[str, object]]) -> Dict[str, object]:
    cleaned = []
    for item in suggestions:
        term = str(item.get("term", "")).strip()
        if not term:
            continue
        cleaned.append(
            {
                "term": term,
                "search_volume": item.get("search_volume"),
                "competition": _normalise_competition(item.get("competition")),
                "relevance": item.get("relevance"),
            }
        )
    return {
        "clusters": [
            {
                "label": "Weitere Vorschläge",
                "keywords": cleaned,
            }
        ]
    }


def analyse_keyword_clusters_with_openai(
    keyword: str,
    suggestions: List[Dict[str, object]],
    api_key: str,
    *,
    prompt: Optional[str] = None,
    model: Optional[str] = None,
    temperature: Optional[float] = None,
) -> Dict[str, object]:
    if not suggestions:
        return {"clusters": []}

    prompt_text = (prompt or "").strip() or DEFAULT_KEYWORD_FINDER_PROMPT
    model_name = (model or "").strip() or OPENAI_MODEL_NAME
    try:
        temperature_value = float(temperature) if temperature is not None else DEFAULT_KEYWORD_FINDER_TEMPERATURE
    except (TypeError, ValueError):
        temperature_value = DEFAULT_KEYWORD_FINDER_TEMPERATURE

    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json",
    }
    user_payload = {
        "keyword": keyword,
        "suggestions": suggestions,
    }
    request_payload = {
        "model": model_name,
        "messages": [
            {"role": "system", "content": prompt_text},
            {"role": "user", "content": json.dumps(user_payload, ensure_ascii=False)},
        ],
        "temperature": temperature_value,
        "max_tokens": 600,
    }

    try:
        response = requests.post(
            OPENAI_CHAT_COMPLETIONS_URL,
            headers=headers,
            json=request_payload,
            timeout=OPENAI_TIMEOUT,
        )
    except requests.RequestException as exc:
        raise OpenAIIntegrationError("LLM-Auswertung der Keyword-Vorschläge fehlgeschlagen (Netzwerkfehler).") from exc

    if response.status_code >= 400:
        raise OpenAIIntegrationError(
            f"LLM-Auswertung der Keyword-Vorschläge fehlgeschlagen (Status {response.status_code})."
        )

    try:
        data = response.json()
    except ValueError as exc:
        raise OpenAIIntegrationError("LLM-Antwort konnte nicht gelesen werden.") from exc

    content = ""
    choices = data.get("choices") if isinstance(data, dict) else None
    if isinstance(choices, list) and choices:
        message = choices[0].get("message") if isinstance(choices[0], dict) else None
        if isinstance(message, dict):
            content = message.get("content", "")

    if not content:
        raise OpenAIIntegrationError("LLM lieferte keine verwertbaren Daten.")

    content = content.strip()
    parsed: Optional[Dict[str, object]] = None
    try:
        parsed_obj = json.loads(content)
    except json.JSONDecodeError:
        start_idx = content.find("{")
        end_idx = content.rfind("}")
        if start_idx != -1 and end_idx != -1 and end_idx > start_idx:
            try:
                parsed_obj = json.loads(content[start_idx : end_idx + 1])
            except json.JSONDecodeError as exc:
                raise OpenAIIntegrationError("LLM-Antwort war nicht im JSON-Format.") from exc
        else:
            raise OpenAIIntegrationError("LLM-Antwort war nicht im JSON-Format.")
    else:
        parsed = parsed_obj if isinstance(parsed_obj, dict) else None

    if parsed is None:
        raise OpenAIIntegrationError("LLM-Antwort enthielt keine strukturierten Daten.")

    clusters_raw = parsed.get("clusters") if isinstance(parsed, dict) else None
    clusters: List[Dict[str, object]] = []
    if isinstance(clusters_raw, list):
        for cluster in clusters_raw:
            if not isinstance(cluster, dict):
                continue
            label = str(cluster.get("label", "")).strip() or "Weitere Vorschläge"
            keywords_raw = cluster.get("keywords")
            keywords: List[Dict[str, object]] = []
            if isinstance(keywords_raw, list):
                for keyword_entry in keywords_raw:
                    if not isinstance(keyword_entry, dict):
                        continue
                    term = str(keyword_entry.get("term", "")).strip()
                    if not term:
                        continue
                    keywords.append(
                        {
                            "term": term,
                            "search_volume": keyword_entry.get("search_volume"),
                            "competition": _normalise_competition(keyword_entry.get("competition")),
                            "relevance": keyword_entry.get("relevance"),
                        }
                    )
            if keywords:
                clusters.append({"label": label, "keywords": keywords})

    if not clusters:
        return _build_fallback_clusters(suggestions)

    return {"clusters": clusters}


def expand_keywords_with_ai(
    keywords: List[str],
    *,
    enabled: bool,
    api_key: Optional[str],
) -> Tuple[List[str], List[Dict[str, List[str]]], List[Dict[str, str]], bool]:
    """Return expanded keywords, grouping information and user messages."""

    cleaned_keywords = [kw.strip() for kw in keywords if kw and kw.strip()]
    expanded: List[str] = []
    seen_lower: Set[str] = set()
    for kw in cleaned_keywords:
        lowered = kw.lower()
        if lowered not in seen_lower:
            expanded.append(kw)
            seen_lower.add(lowered)

    keyword_groups: List[Dict[str, List[str]]] = []
    messages: List[Dict[str, str]] = []
    defaults = get_synonym_defaults()

    if not cleaned_keywords:
        return expanded, keyword_groups, messages, False

    if not enabled:
        for kw in cleaned_keywords:
            keyword_groups.append({"original": kw, "additional": []})
        return expanded, keyword_groups, messages, False

    if not api_key:
        messages.append(
            {
                "category": "warning",
                "text": "OpenAI-Schlüssel erforderlich. Bitte unter Einstellungen → API-Schlüssel hinterlegen.",
            }
        )
        for kw in cleaned_keywords:
            keyword_groups.append({"original": kw, "additional": []})
        return expanded, keyword_groups, messages, False

    message_texts: Set[str] = set()
    synonyms_expanded = False

    for kw in cleaned_keywords:
        additional: List[str] = []
        try:
            synonyms = fetch_synonyms_from_openai(
                kw,
                api_key,
                prompt=defaults.get("prompt"),
                model=str(defaults.get("model") or OPENAI_MODEL_NAME),
                temperature=defaults.get("temperature"),
                max_results=defaults.get("max_results"),
            )
        except OpenAIIntegrationError as exc:
            text = str(exc)
            if text and text not in message_texts:
                messages.append({"category": "warning", "text": text})
                message_texts.add(text)
            keyword_groups.append({"original": kw, "additional": []})
            continue

        for synonym in synonyms:
            lowered = synonym.lower()
            if lowered in seen_lower:
                continue
            seen_lower.add(lowered)
            expanded.append(synonym)
            additional.append(synonym)

        if additional:
            synonyms_expanded = True

        keyword_groups.append({"original": kw, "additional": additional})

    if synonyms_expanded:
        total_added = sum(len(group["additional"]) for group in keyword_groups)
        info_text = (
            "Synonymerweiterung aktiv: "
            f"{total_added} zusätzliche Begriffe werden durchsucht."
        )
        messages.append({"category": "info", "text": info_text})

    return expanded, keyword_groups, messages, synonyms_expanded


def evaluate_keyword_with_openai(
    keyword: str,
    api_key: str,
    *,
    prompt: Optional[str] = None,
) -> Dict[str, object]:
    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json",
    }
    prompt_text = (prompt or "").strip() or DEFAULT_KEYWORD_EVALUATION_PROMPT
    payload = {
        "model": OPENAI_MODEL_NAME,
        "messages": [
            {"role": "system", "content": prompt_text},
            {"role": "user", "content": f'Begriff: "{keyword}"'},
        ],
        "temperature": 0.2,
        "max_tokens": 200,
    }

    try:
        response = requests.post(
            OPENAI_CHAT_COMPLETIONS_URL,
            headers=headers,
            json=payload,
            timeout=OPENAI_TIMEOUT,
        )
    except requests.RequestException as exc:  # pragma: no cover - network failure
        raise OpenAIIntegrationError(
            f'Bewertung für "{keyword}" nicht möglich: Keine Verbindung zur OpenAI-API.'
        ) from exc

    if response.status_code >= 400:
        raise OpenAIIntegrationError(
            f'Bewertung für "{keyword}" nicht möglich: OpenAI-API meldet Status {response.status_code}.'
        )

    try:
        response_payload = response.json()
    except ValueError as exc:  # pragma: no cover - invalid JSON
        raise OpenAIIntegrationError(
            f'Bewertung für "{keyword}" nicht möglich: Antwort konnte nicht gelesen werden.'
        ) from exc

    choices = response_payload.get("choices") if isinstance(response_payload, dict) else None
    content = ""
    if isinstance(choices, list) and choices:
        message = choices[0].get("message") if isinstance(choices[0], dict) else None
        if isinstance(message, dict):
            content = str(message.get("content", ""))

    if not content:
        raise OpenAIIntegrationError(
            f'Bewertung für "{keyword}" nicht möglich: Antwort enthielt keine Daten.'
        )

    content = content.strip()
    try:
        parsed = json.loads(content)
    except json.JSONDecodeError:
        start_idx = content.find("{")
        end_idx = content.rfind("}")
        if start_idx == -1 or end_idx == -1 or end_idx <= start_idx:
            raise OpenAIIntegrationError(
                f'Bewertung für "{keyword}" nicht möglich: Antwort war nicht im JSON-Format.'
            )
        try:
            parsed = json.loads(content[start_idx : end_idx + 1])
        except json.JSONDecodeError as exc:
            raise OpenAIIntegrationError(
                f'Bewertung für "{keyword}" nicht möglich: Antwort war nicht im JSON-Format.'
            ) from exc

    if not isinstance(parsed, dict):
        raise OpenAIIntegrationError(
            f'Bewertung für "{keyword}" nicht möglich: Antwort enthielt kein Objekt.'
        )

    quality_raw = str(parsed.get("quality", "")).strip()
    quality = quality_raw or "unbewertet"

    better_raw = (
        parsed.get("better_keywords")
        or parsed.get("betterTerms")
        or parsed.get("better_terms")
        or parsed.get("bessere_suchbegriffe")
        or parsed.get("alternativen")
    )
    better_keywords: List[str] = []
    if isinstance(better_raw, list):
        seen: Set[str] = set()
        for item in better_raw:
            text = str(item).strip()
            if not text:
                continue
            lowered = text.lower()
            if lowered in seen:
                continue
            seen.add(lowered)
            better_keywords.append(text)
            if len(better_keywords) >= 8:
                break

    return {
        "keyword": keyword,
        "quality": quality,
        "better_keywords": better_keywords,
    }


def _collapse_whitespace(text: str) -> str:
    return " ".join((text or "").split())


def _truncate_text(text: str, limit: int) -> str:
    if limit <= 0:
        return ""
    collapsed = _collapse_whitespace(text)
    if len(collapsed) <= limit:
        return collapsed
    return collapsed[: limit - 3] + "..."


def _strip_accents(text: str) -> str:
    decomposed = unicodedata.normalize("NFKD", text or "")
    return "".join(ch for ch in decomposed if not unicodedata.combining(ch))


def _normalise_for_match(text: str) -> str:
    return re.sub(r"\s+", " ", _strip_accents(text).casefold()).strip()


def _extract_city_names(text: str) -> List[str]:
    candidates: List[str] = []
    if not text:
        return candidates
    for match in re.finditer(r"\b\d{5}\s+([A-Za-zÄÖÜäöüß\- ]{2,})", text):
        city = match.group(1).strip()
        if city:
            candidates.append(city)
    return candidates


def _check_impressum_location(snapshot: Dict[str, object], ort: str) -> Tuple[bool, Optional[str], Optional[str]]:
    """Return (match?, context, mismatch_hint)."""

    location_norm = _normalise_for_match(ort)
    if not location_norm:
        return True, None, None

    impressum_text = str(snapshot.get("impressum_excerpt", "") or "")
    structured = snapshot.get("structured_signals")
    if not isinstance(structured, dict):
        structured = {}

    addresses = structured.get("addresses") if isinstance(structured.get("addresses"), list) else []
    mismatch_hint: Optional[str] = None

    def _check_text(source_text: str) -> Optional[str]:
        normalised = _normalise_for_match(source_text)
        if location_norm and location_norm in normalised:
            return source_text.strip()
        return None

    # Check structured address lines first
    for address in addresses:
        context = _check_text(str(address))
        if context:
            return True, context, None
        for city in _extract_city_names(str(address)):
            city_norm = _normalise_for_match(city)
            if city_norm and city_norm != location_norm:
                mismatch_hint = city.strip()

    # Check the impressum excerpt itself
    if impressum_text:
        context = _check_text(impressum_text)
        if context:
            return True, context, None
        for city in _extract_city_names(impressum_text):
            city_norm = _normalise_for_match(city)
            if city_norm and city_norm != location_norm:
                mismatch_hint = city.strip()

    return False, None, mismatch_hint


def fetch_url_text(url: str) -> Optional[str]:
    headers = {
        "User-Agent": "Mozilla/5.0 (compatible; SchulCrawler/1.0)",
        "Accept": "text/html,application/xhtml+xml",
    }
    try:
        response = requests.get(url, headers=headers, timeout=15)
    except requests.RequestException:
        return None
    if response.status_code >= 400:
        return None
    content_type = response.headers.get("Content-Type", "").lower()
    if "text" not in content_type and "html" not in content_type:
        return None
    response.encoding = response.apparent_encoding or response.encoding
    return response.text


def extract_visible_text(html: str) -> str:
    if not html:
        return ""
    soup = BeautifulSoup(html, "html.parser")
    for element in soup(["script", "style", "noscript"]):
        element.decompose()
    texts = [segment.strip() for segment in soup.stripped_strings]
    return " ".join(texts)


def collect_site_snapshot(url: str) -> Dict[str, object]:
    snapshot = {
        "url": url,
        "domain": urlparse(url).netloc,
        "title": "",
        "description": "",
        "main_excerpt": "",
        "impressum_excerpt": "",
        "structured_signals": {},
        "fetched_urls": [],
        "errors": [],
    }
    html = fetch_url_text(url)
    if not html:
        snapshot["errors"].append("Startseite konnte nicht geladen werden.")
        return snapshot

    soup = BeautifulSoup(html, "html.parser")
    snapshot["title"] = _collapse_whitespace(soup.title.string if soup.title else "")
    description_tag = soup.find("meta", attrs={"name": re.compile("description", re.I)})
    if description_tag:
        snapshot["description"] = _collapse_whitespace(description_tag.get("content", ""))
    text = extract_visible_text(html)
    snapshot["main_excerpt"] = _truncate_text(text, 1500)
    snapshot["fetched_urls"].append(url)

    structured: Dict[str, object] = {}
    emails = sorted({match.group(0) for match in re.finditer(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+", text)})
    if emails:
        structured["emails"] = emails[:5]
    postal_lines: List[str] = []
    for line in text.split(". "):
        if re.search(r"\b\d{5}\b", line):
            postal_lines.append(_collapse_whitespace(line))
            if len(postal_lines) >= 5:
                break
    if postal_lines:
        structured["addresses"] = postal_lines
    snapshot["structured_signals"] = structured

    base_url = url
    impressum_links: List[str] = []
    auxiliary_links: List[str] = []
    for link in soup.find_all("a", href=True):
        link_text = link.get_text(" ", strip=True)
        link_text_lower = link_text.lower()
        href_lower = link["href"].lower()
        if any(keyword in link_text_lower for keyword in DATA_QUALITY_IMPRESSUM_KEYWORDS) or any(
            keyword in href_lower for keyword in DATA_QUALITY_IMPRESSUM_KEYWORDS
        ):
            absolute = urljoin(base_url, link["href"])
            if "impressum" in link_text_lower or "impressum" in href_lower:
                impressum_links.append(absolute)
            else:
                auxiliary_links.append(absolute)
    nav_links: List[str] = []
    seen_nav: Set[str] = set()
    for candidate_list in (impressum_links, auxiliary_links):
        for link in candidate_list:
            if link in seen_nav:
                continue
            seen_nav.add(link)
            nav_links.append(link)
            if len(nav_links) >= 3:
                break
        if len(nav_links) >= 3:
            break

    for nav_url in nav_links:
        html_nav = fetch_url_text(nav_url)
        if not html_nav:
            continue
        nav_text = extract_visible_text(html_nav)
        if not nav_text:
            continue
        if "impressum" in nav_url.lower() or "impressum" in nav_text.lower():
            snapshot["impressum_excerpt"] = _truncate_text(nav_text, 1500)
            snapshot["impressum_url"] = nav_url
        else:
            if not snapshot["impressum_excerpt"]:
                snapshot["impressum_excerpt"] = _truncate_text(nav_text, 800)
        snapshot["fetched_urls"].append(nav_url)
        if snapshot["impressum_excerpt"]:
            break

    return snapshot


def assess_school_website_with_llm(
    schulname: str,
    ort: str,
    url: str,
    snapshot: Dict[str, object],
    *,
    api_key: str,
    settings: Dict[str, object],
) -> Optional[Dict[str, object]]:
    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json",
    }
    model_name = str(settings.get("model") or OPENAI_MODEL_NAME)
    temperature = float(settings.get("temperature", DEFAULT_DATA_QUALITY_SETTINGS["temperature"]))
    payload = {
        "schule": schulname,
        "ort": ort,
        "url": url,
        "domain": snapshot.get("domain"),
        "seitentitel": snapshot.get("title"),
        "beschreibung": snapshot.get("description"),
        "startseite_text": snapshot.get("main_excerpt"),
        "impressum_text": snapshot.get("impressum_excerpt"),
        "signale": snapshot.get("structured_signals", {}),
    }
    request_payload = {
        "model": model_name,
        "messages": [
            {"role": "system", "content": DATA_QUALITY_SITE_PROMPT},
            {"role": "user", "content": json.dumps(payload, ensure_ascii=False)},
        ],
        "temperature": temperature,
        "max_tokens": 600,
    }

    try:
        response = requests.post(
            OPENAI_CHAT_COMPLETIONS_URL,
            headers=headers,
            json=request_payload,
            timeout=OPENAI_TIMEOUT,
        )
    except requests.RequestException:
        return None

    if response.status_code >= 400:
        return None

    try:
        data = response.json()
    except ValueError:
        return None

    choices = data.get("choices") if isinstance(data, dict) else None
    content = ""
    if isinstance(choices, list) and choices:
        message = choices[0].get("message") if isinstance(choices[0], dict) else None
        if isinstance(message, dict):
            content = str(message.get("content", ""))
    if not content:
        return None

    content = content.strip()
    try:
        parsed = json.loads(content)
    except json.JSONDecodeError:
        start_idx = content.find("{")
        end_idx = content.rfind("}")
        if start_idx == -1 or end_idx == -1 or end_idx <= start_idx:
            return None
        try:
            parsed = json.loads(content[start_idx : end_idx + 1])
        except json.JSONDecodeError:
            return None
    if not isinstance(parsed, dict):
        return None
    scope_value = normalise_site_scope(parsed.get("site_scope") or parsed.get("auftritt_typ"))
    parsed["site_scope"] = scope_value
    reason_value = parsed.get("site_scope_reason") or parsed.get("auftritt_begruendung")
    parsed["site_scope_reason"] = str(reason_value or "").strip()
    return parsed


def perform_google_search(
    query: str,
    api_key: str,
    cx: str,
    *,
    max_results: int,
    language: str = "de",
    region: str = "de",
) -> List[Dict[str, object]]:
    max_results = max(1, min(int(max_results), 200))
    collected: List[Dict[str, object]] = []
    seen_links: Set[str] = set()
    start_index = 1
    log_developer_event(
        "google-search",
        "Google-Suche gestartet",
        details={
            "query": query,
            "max_results": max_results,
            "language": language,
            "region": region,
        },
    )
    while len(collected) < max_results:
        remaining = max_results - len(collected)
        batch_size = max(1, min(remaining, 10))
        params = {
            "key": api_key,
            "cx": cx,
            "q": query,
            "num": batch_size,
            "start": start_index,
            "hl": (language or "de")[:5],
            "gl": (region or "de")[:5],
            "safe": "active",
        }
        try:
            response = requests.get(GOOGLE_SEARCH_API_URL, params=params, timeout=GOOGLE_SEARCH_TIMEOUT)
        except requests.RequestException:
            log_developer_event(
                "google-search",
                "Google-Suche fehlgeschlagen",
                details={"query": query, "reason": "Netzwerkfehler", "start": start_index},
            )
            break
        if response.status_code >= 400:
            log_developer_event(
                "google-search",
                "Google-Suche mit Fehlerstatus",
                details={
                    "query": query,
                    "status": response.status_code,
                    "start": start_index,
                },
            )
            break
        try:
            payload = response.json()
        except ValueError:
            log_developer_event(
                "google-search",
                "Antwort konnte nicht gelesen werden",
                details={"query": query, "start": start_index},
            )
            break
        items = payload.get("items") if isinstance(payload, dict) else None
        if not isinstance(items, list) or not items:
            log_developer_event(
                "google-search",
                "Keine weiteren Treffer",
                details={"query": query, "start": start_index},
            )
            break
        added = 0
        for item in items:
            if not isinstance(item, dict):
                continue
            link_raw = item.get("link")
            if not link_raw:
                continue
            link = str(link_raw)
            if link in seen_links:
                continue
            seen_links.add(link)
            collected.append(
                {
                    "title": _collapse_whitespace(str(item.get("title", ""))),
                    "snippet": _collapse_whitespace(str(item.get("snippet", ""))),
                    "displayLink": str(item.get("displayLink", "")),
                    "link": link,
                }
            )
            added += 1
            if len(collected) >= max_results:
                break
        if added == 0:
            break
        start_index += max(1, len(items))
        if start_index > 100 and len(collected) >= 100:
            break
    log_developer_event(
        "google-search",
        "Google-Suche abgeschlossen",
        details={
            "query": query,
            "retrieved": len(collected),
            "domains": [
                _extract_domain_from_url(item.get("link", ""))
                for item in collected[: min(len(collected), 5)]
            ],
        },
    )
    return collected


def _extract_domain_from_url(url: str) -> str:
    try:
        parsed = urlparse(url)
    except ValueError:
        return ""
    domain = parsed.netloc or parsed.path
    if domain.startswith("www."):
        domain = domain[4:]
    return domain.lower()


def fetch_page_text(url: str, *, timeout: int = 15) -> str:
    headers = {"User-Agent": "Mozilla/5.0 (compatible; SchulCrawler/1.0)"}
    try:
        response = requests.get(url, headers=headers, timeout=timeout)
    except requests.RequestException:
        return ""
    if response.status_code >= 400:
        return ""
    content_type = str(response.headers.get("Content-Type", "")).lower()
    if "text" not in content_type and "html" not in content_type:
        return ""
    try:
        soup = BeautifulSoup(response.text, "html.parser")
    except Exception:
        return ""
    for element in soup(["script", "style", "noscript", "template"]):
        element.decompose()
    text = soup.get_text(separator=" ", strip=True)
    collapsed = _collapse_whitespace(text)
    return collapsed[:20000]


def evaluate_page_with_llm(
    content: str,
    prompt: str,
    api_key: str,
    *,
    model: str,
    temperature: float,
) -> Optional[List[Dict[str, Any]]]:
    cleaned = (content or "").strip()
    if not cleaned:
        return None
    template = (prompt or "").strip() or DEFAULT_SEARCH_EVALUATION_PROMPT
    if "{{CONTENT}}" in template:
        user_prompt = template.replace("{{CONTENT}}", cleaned)
    else:
        user_prompt = f"{template}\n\nText zur Analyse:\n{cleaned}"
    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json",
    }
    payload = {
        "model": model or DEFAULT_SEARCH_MODEL,
        "messages": [
            {
                "role": "system",
                "content": "Du analysierst Schulwebseiten nach dem bereitgestellten Bewertungsprompt und gibst JSON aus.",
            },
            {"role": "user", "content": user_prompt},
        ],
        "temperature": max(0.0, min(float(temperature), 2.0)),
        "max_tokens": 900,
    }
    log_developer_event(
        "openai-search",
        "Bewertung angefordert",
        details={
            "model": payload["model"],
            "temperature": payload["temperature"],
            "prompt_chars": len(user_prompt),
            "content_chars": len(cleaned),
        },
    )
    try:
        response = requests.post(
            OPENAI_CHAT_COMPLETIONS_URL,
            headers=headers,
            json=payload,
            timeout=OPENAI_TIMEOUT,
        )
    except requests.RequestException:
        log_developer_event(
            "openai-search",
            "Bewertung fehlgeschlagen",
            details={"reason": "Netzwerkfehler", "model": payload["model"]},
        )
        return None
    if response.status_code >= 400:
        log_developer_event(
            "openai-search",
            "Bewertung mit Fehlerstatus",
            details={"status": response.status_code, "model": payload["model"]},
        )
        return None
    try:
        data = response.json()
    except ValueError:
        log_developer_event(
            "openai-search",
            "Bewertung lieferte keine JSON-Antwort",
            details={"model": payload["model"], "status": response.status_code},
        )
        return None
    choices = data.get("choices") if isinstance(data, dict) else None
    content_text = ""
    if isinstance(choices, list) and choices:
        first = choices[0]
        if isinstance(first, dict):
            message = first.get("message")
            if isinstance(message, dict):
                content_text = str(message.get("content", ""))
    content_text = content_text.strip()
    if not content_text:
        log_developer_event(
            "openai-search",
            "Bewertung ohne Inhalt",
            details={"model": payload["model"], "status": response.status_code},
        )
        return None
    try:
        parsed = json.loads(content_text)
    except json.JSONDecodeError:
        start = content_text.find("{")
        end = content_text.rfind("}")
        if start == -1 or end == -1 or end <= start:
            return None
        try:
            parsed = json.loads(content_text[start : end + 1])
        except json.JSONDecodeError:
            log_developer_event(
                "openai-search",
                "Bewertung konnte nicht geparst werden",
                details={"model": payload["model"], "status": response.status_code},
            )
            return None
    if not isinstance(parsed, dict):
        return None
    evaluations = parsed.get("Bewertung")
    if not isinstance(evaluations, list):
        return None
    cleaned_eval: List[Dict[str, Any]] = []
    for entry in evaluations:
        if not isinstance(entry, dict):
            continue
        dimension = str(entry.get("Dimension", "")).strip()
        try:
            value = int(entry.get("Wert", 0))
        except (TypeError, ValueError):
            try:
                value = int(float(entry.get("Wert", 0)))
            except (TypeError, ValueError):
                value = 0
        try:
            confidence = float(entry.get("Sicherheit", 0.0))
        except (TypeError, ValueError):
            confidence = 0.0
        cleaned_eval.append(
            {
                "dimension": dimension,
                "value": max(0, min(value, 5)),
                "confidence": max(0.0, min(confidence, 1.0)),
            }
        )
    log_developer_event(
        "openai-search",
        "Bewertung abgeschlossen",
        details={
            "model": payload["model"],
            "status": response.status_code,
            "dimensions": len(cleaned_eval),
            "usage": _sanitize_log_value(data.get("usage")) if isinstance(data, dict) else None,
        },
    )
    return cleaned_eval or None


def _normalise_domain(domain: str) -> str:
    text = (domain or "").strip().lower()
    if text.startswith("www."):
        text = text[4:]
    return text


def _domain_matches(domain: str, candidates: Set[str]) -> bool:
    if not domain:
        return False
    base = _normalise_domain(domain)
    for candidate in candidates:
        cand = _normalise_domain(candidate)
        if not cand:
            continue
        if base == cand or base.endswith(f".{cand}"):
            return True
    return False


def _domain_allowed(url: str, allowed: Set[str], blocked: Set[str]) -> bool:
    try:
        parsed = urlparse(url)
    except ValueError:
        return False
    domain = _normalise_domain(parsed.netloc)
    if not domain:
        domain = _normalise_domain(parsed.path)
    if _domain_matches(domain, blocked):
        return False
    if allowed and not _domain_matches(domain, allowed):
        return False
    return True


def _normalise_candidate_url(url: str) -> str:
    text = (url or "").strip()
    if not text:
        return ""
    parsed = urlparse(text, scheme="https")
    scheme = parsed.scheme.lower() if parsed.scheme else "https"
    if scheme not in {"http", "https"}:
        scheme = "https"
    netloc = parsed.netloc
    path = parsed.path
    if not netloc and parsed.path:
        netloc = parsed.path
        path = ""
    if not netloc:
        return ""
    normalised_path = path or "/"
    return urlunparse((scheme, netloc, normalised_path, "", parsed.query or "", ""))


def _urls_equivalent(first: str, second: str) -> bool:
    def _prep(value: str) -> str:
        normalised = _normalise_candidate_url(value)
        if not normalised:
            return ""
        parsed = urlparse(normalised)
        path = parsed.path.rstrip("/") or "/"
        return urlunparse((parsed.scheme, parsed.netloc.lower(), path, "", parsed.query, ""))

    return _prep(first) == _prep(second)


def analyse_search_results_with_llm(
    schulname: str,
    ort: str,
    original_url: str,
    results: List[Dict[str, object]],
    *,
    api_key: str,
    settings: Dict[str, object],
) -> Optional[Dict[str, object]]:
    if not results:
        return None
    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json",
    }
    model_name = str(settings.get("model") or OPENAI_MODEL_NAME)
    temperature = float(settings.get("temperature", DEFAULT_DATA_QUALITY_SETTINGS["temperature"]))
    payload = {
        "schule": schulname,
        "ort": ort,
        "original_url": original_url,
        "treffer": results,
    }
    request_payload = {
        "model": model_name,
        "messages": [
            {"role": "system", "content": DATA_QUALITY_SEARCH_PROMPT},
            {"role": "user", "content": json.dumps(payload, ensure_ascii=False)},
        ],
        "temperature": temperature,
        "max_tokens": 600,
    }
    try:
        response = requests.post(
            OPENAI_CHAT_COMPLETIONS_URL,
            headers=headers,
            json=request_payload,
            timeout=OPENAI_TIMEOUT,
        )
    except requests.RequestException:
        return None
    if response.status_code >= 400:
        return None
    try:
        data = response.json()
    except ValueError:
        return None
    choices = data.get("choices") if isinstance(data, dict) else None
    content = ""
    if isinstance(choices, list) and choices:
        message = choices[0].get("message") if isinstance(choices[0], dict) else None
        if isinstance(message, dict):
            content = str(message.get("content", ""))
    if not content:
        return None
    content = content.strip()
    try:
        parsed = json.loads(content)
    except json.JSONDecodeError:
        start_idx = content.find("{")
        end_idx = content.rfind("}")
        if start_idx == -1 or end_idx == -1 or end_idx <= start_idx:
            return None
        try:
            parsed = json.loads(content[start_idx : end_idx + 1])
        except json.JSONDecodeError:
            return None
    if not isinstance(parsed, dict):
        return None
    return parsed


def normalise_quality_decision(value: str) -> str:
    text = (value or "").strip().lower()
    if text in {"ok", "korrekt", "ja", "true"}:
        return QUALITY_STATUS_OK
    if text in {"nein", "falsch", "false"}:
        return QUALITY_STATUS_INVALID
    return QUALITY_STATUS_UNSURE


def test_openai_api_key(api_key: str) -> None:
    headers = {"Authorization": f"Bearer {api_key}"}
    try:
        response = requests.get(OPENAI_MODELS_URL, headers=headers, timeout=OPENAI_TIMEOUT)
    except requests.RequestException as exc:  # pragma: no cover - network failure
        raise OpenAIIntegrationError("Die Verbindung zur OpenAI-API konnte nicht hergestellt werden.") from exc

    if response.status_code >= 400:
        raise OpenAIIntegrationError(
            f"OpenAI-API meldet einen Fehler (Status {response.status_code})."
        )


@app.post("/synonyms")
def generate_synonyms() -> ResponseReturnValue:
    api_key = get_api_key()
    if not api_key:
        return jsonify({"error": "OpenAI-Schlüssel erforderlich."}), 400

    payload = request.get_json(silent=True) or {}
    if not isinstance(payload, dict):
        return jsonify({"error": "Ungültige Anfrage."}), 400

    keyword = str(payload.get("keyword", "")).strip()
    prompt = payload.get("prompt")
    model = payload.get("model")
    max_results = payload.get("max_results")
    temperature = payload.get("temperature")
    defaults = get_synonym_defaults()

    if not keyword:
        return jsonify({"error": "Bitte geben Sie einen Suchbegriff an."}), 400

    prompt_text = (prompt or "").strip() or str(defaults.get("prompt", DEFAULT_SYNONYM_PROMPT))
    model_name = str(model or defaults.get("model") or OPENAI_MODEL_NAME).strip() or OPENAI_MODEL_NAME
    try:
        temperature_value = float(temperature) if temperature is not None else float(defaults.get("temperature", DEFAULT_SYNONYM_TEMPERATURE))
    except (TypeError, ValueError):
        temperature_value = float(defaults.get("temperature", DEFAULT_SYNONYM_TEMPERATURE))
    if temperature_value < 0:
        temperature_value = 0.0
    elif temperature_value > 2:
        temperature_value = 2.0
    try:
        max_results_value = int(max_results) if max_results is not None else int(defaults.get("max_results", DEFAULT_SYNONYM_MAX_RESULTS))
    except (TypeError, ValueError):
        max_results_value = int(defaults.get("max_results", DEFAULT_SYNONYM_MAX_RESULTS))
    if max_results_value < 1:
        max_results_value = 1
    elif max_results_value > 150:
        max_results_value = 150

    cache_key = synonym_cache_key(keyword, prompt_text, model_name, temperature_value)
    cached = get_synonyms_from_cache(cache_key)
    if cached is not None:
        return jsonify(
            {
                "keyword": keyword,
                "synonyms": cached[:max_results_value],
                "cached": True,
            }
        )

    try:
        synonyms = fetch_synonyms_from_openai(
            keyword,
            api_key,
            prompt=prompt_text,
            model=model_name,
            temperature=temperature_value,
            max_results=max_results_value,
        )
    except OpenAIIntegrationError as exc:
        return jsonify({"error": str(exc)}), 502

    update_synonym_cache_entry(cache_key, synonyms)

    return jsonify(
        {
            "keyword": keyword,
            "synonyms": synonyms,
            "cached": False,
        }
    )


@app.post("/related-keywords")
def generate_related_keywords() -> ResponseReturnValue:
    payload = request.get_json(silent=True) or {}
    keyword = str(payload.get("keyword", "")).strip()
    if not keyword:
        return jsonify({"error": "Bitte geben Sie ein Suchwort an."}), 400

    defaults = get_keyword_finder_defaults()
    prompt = str(payload.get("prompt", "")).strip() or defaults.get("prompt", DEFAULT_KEYWORD_FINDER_PROMPT)
    model = str(payload.get("model", "")).strip() or defaults.get("model", OPENAI_MODEL_NAME)

    try:
        temperature = float(payload.get("temperature", defaults.get("temperature", DEFAULT_KEYWORD_FINDER_TEMPERATURE)))
    except (TypeError, ValueError):
        temperature = defaults.get("temperature", DEFAULT_KEYWORD_FINDER_TEMPERATURE)
    temperature = max(0.0, min(temperature, 2.0))

    try:
        max_results = int(payload.get("max_results", defaults.get("max_results", DEFAULT_KEYWORD_FINDER_MAX_RESULTS)))
    except (TypeError, ValueError):
        max_results = defaults.get("max_results", DEFAULT_KEYWORD_FINDER_MAX_RESULTS)
    max_results = max(1, min(max_results, DEFAULT_KEYWORD_FINDER_MAX_RESULTS))

    google_key = get_keyword_planner_key()
    if not google_key:
        return jsonify({"error": "Kein Keyword-Planner-Schlüssel hinterlegt. Bitte in den Einstellungen speichern."}), 400

    try:
        suggestions = fetch_keyword_planner_suggestions(keyword, google_key, max_results=max_results)
    except KeywordPlannerError as exc:
        return jsonify({"error": str(exc)}), 502

    cache_key = keyword_finder_cache_key(keyword, prompt, model, temperature, max_results)
    cached_entry = get_keyword_finder_cache_entry(cache_key)
    current_terms = [item.get("term", "").strip().lower() for item in suggestions]
    messages: List[str] = []

    if cached_entry and isinstance(cached_entry, dict):
        cached_suggestions = cached_entry.get("suggestions")
        if isinstance(cached_suggestions, list):
            cached_terms = [
                str(item.get("term", "")).strip().lower()
                for item in cached_suggestions
                if isinstance(item, dict)
            ]
            if cached_terms != current_terms:
                cached_entry = None
        else:
            cached_entry = None

    if cached_entry:
        clusters = cached_entry.get("clusters") if isinstance(cached_entry.get("clusters"), list) else []
        analysis = cached_entry.get("analysis") if isinstance(cached_entry.get("analysis"), str) else "cache"
        return jsonify(
            {
                "keyword": keyword,
                "clusters": clusters,
                "suggestions": suggestions,
                "analysis": analysis,
                "cached": True,
                "messages": messages,
                "defaults": {
                    "prompt": prompt,
                    "model": model,
                    "temperature": temperature,
                    "max_results": max_results,
                },
            }
        )

    openai_key = get_api_key()
    analysis_mode = "fallback"
    if not openai_key:
        messages.append("OpenAI-Schlüssel fehlt – Vorschläge werden ohne KI-Gruppierung angezeigt.")
        clusters_payload = _build_fallback_clusters(suggestions)
    else:
        try:
            clusters_payload = analyse_keyword_clusters_with_openai(
                keyword,
                suggestions,
                openai_key,
                prompt=prompt,
                model=model,
                temperature=temperature,
            )
        except OpenAIIntegrationError as exc:
            messages.append(str(exc))
            clusters_payload = _build_fallback_clusters(suggestions)
        else:
            analysis_mode = "llm"

    clusters = clusters_payload.get("clusters", []) if isinstance(clusters_payload, dict) else []

    cache_payload = {
        "clusters": clusters,
        "suggestions": suggestions,
        "analysis": analysis_mode,
    }
    update_keyword_finder_cache_entry(cache_key, cache_payload)

    return jsonify(
        {
            "keyword": keyword,
            "clusters": clusters,
            "suggestions": suggestions,
            "analysis": analysis_mode,
            "cached": False,
            "messages": messages,
            "defaults": {
                "prompt": prompt,
                "model": model,
                "temperature": temperature,
                "max_results": max_results,
            },
        }
    )


@app.post("/evaluate-keywords")
def evaluate_keywords() -> ResponseReturnValue:
    api_key = get_api_key()
    if not api_key:
        return jsonify({"error": "OpenAI-Schlüssel erforderlich."}), 400

    payload = request.get_json(silent=True) or {}
    if not isinstance(payload, dict):
        return jsonify({"error": "Ungültige Anfrage."}), 400

    keywords_raw = payload.get("keywords")
    prompt = payload.get("prompt")

    if not isinstance(keywords_raw, list):
        return jsonify({"error": "Bitte geben Sie eine Liste von Suchbegriffen an."}), 400

    cleaned: List[str] = []
    seen_lower: Set[str] = set()
    for item in keywords_raw:
        text = str(item).strip()
        if not text:
            continue
        lowered = text.lower()
        if lowered in seen_lower:
            continue
        seen_lower.add(lowered)
        cleaned.append(text)

    if not cleaned:
        return jsonify({"evaluations": [], "messages": []})

    evaluations: List[Dict[str, object]] = []
    messages: List[str] = []

    for keyword in cleaned:
        try:
            result = evaluate_keyword_with_openai(keyword, api_key, prompt=prompt)
        except OpenAIIntegrationError as exc:
            text = str(exc)
            if text:
                messages.append(text)
        else:
            evaluations.append(result)

    status = 200 if evaluations or not messages else 502
    return jsonify({"evaluations": evaluations, "messages": messages}), status





def _read_saved_searches_unlocked() -> List[Dict[str, object]]:
    return list(storage.list_saved_searches())


def load_saved_searches() -> List[Dict[str, object]]:
    with saved_search_lock:
        return list(_read_saved_searches_unlocked())


def append_saved_search(entry: Dict[str, object]) -> None:
    with saved_search_lock:
        storage.append_saved_search(entry)


def replace_saved_searches(entries: List[Dict[str, object]]) -> None:
    with saved_search_lock:
        storage.replace_saved_searches(entries)


def load_search_definitions() -> List[Dict[str, Any]]:
    with search_definition_lock:
        return list(storage.list_search_definitions())


def get_search_definition(definition_id: str) -> Optional[Dict[str, Any]]:
    with search_definition_lock:
        return storage.get_search_definition(definition_id)


def save_search_definition(definition: Dict[str, Any]) -> None:
    with search_definition_lock:
        storage.upsert_search_definition(definition)


def delete_search_definitions(identifiers: Iterable[str]) -> None:
    with search_definition_lock:
        storage.delete_search_definitions(list(identifiers))


def list_search_runs(definition_id: Optional[str] = None) -> List[Dict[str, Any]]:
    with search_run_lock:
        return list(storage.list_search_runs(definition_id))


def get_search_run(run_id: str) -> Optional[Dict[str, Any]]:
    with search_run_lock:
        return storage.get_search_run(run_id)


def save_search_run(run_payload: Dict[str, Any]) -> None:
    run_id = str(run_payload.get("id") or "").strip()
    definition_id = str(run_payload.get("definition_id") or "").strip()
    if not run_id or not definition_id:
        raise ValueError("Search run erfordert eine ID und eine Definition-ID")
    with search_run_lock:
        storage.upsert_search_run(run_id, definition_id, run_payload)


def append_search_results(run_id: str, entries: Iterable[Dict[str, Any]]) -> None:
    with search_result_lock:
        storage.append_search_results(run_id, list(entries))


def load_search_results(run_id: str) -> List[Dict[str, Any]]:
    with search_result_lock:
        return list(storage.load_search_results(run_id))


def clear_search_results(run_id: str) -> None:
    with search_result_lock:
        storage.clear_search_results(run_id)
    with search_log_lock:
        storage.clear_search_logs(run_id)


def append_search_logs(run_id: str, entries: Iterable[Dict[str, Any]]) -> None:
    with search_log_lock:
        storage.append_search_logs(run_id, list(entries))


def load_search_logs(run_id: str) -> List[Dict[str, Any]]:
    with search_log_lock:
        return list(storage.load_search_logs(run_id))


def clear_search_logs(run_id: str) -> None:
    with search_log_lock:
        storage.clear_search_logs(run_id)


def get_search_result_entry(result_id: int) -> Optional[Dict[str, Any]]:
    with search_result_lock:
        return storage.get_search_result(result_id)


def collect_all_search_data() -> Dict[str, Any]:
    definitions = load_search_definitions()
    runs = list_search_runs()
    results: Dict[str, List[Dict[str, Any]]] = {}
    logs: Dict[str, List[Dict[str, Any]]] = {}
    for run in runs:
        run_id = str(run.get("id") or "").strip()
        if not run_id or run_id in results:
            continue
        results[run_id] = load_search_results(run_id)
        logs[run_id] = load_search_logs(run_id)
    return {
        "definitions": definitions,
        "runs": runs,
        "results": results,
        "logs": logs,
    }


def _normalise_category_list(values: Iterable[object]) -> List[str]:
    seen: Set[str] = set()
    categories: List[str] = []
    for value in values:
        text = str(value or "").strip()
        if not text:
            continue
        lowered = text.lower()
        if lowered in seen:
            continue
        seen.add(lowered)
        categories.append(text)
    for default in DEFAULT_SEARCH_CATEGORIES:
        text = str(default or "").strip()
        if not text:
            continue
        lowered = text.lower()
        if lowered not in seen:
            categories.insert(0, text)
            seen.add(lowered)
    return categories


def get_search_categories() -> List[str]:
    with search_categories_lock:
        existing = storage.list_search_categories()
        categories = _normalise_category_list(existing)
        if categories != existing:
            storage.save_search_categories(categories)
        return list(categories)


def add_search_category(name: str) -> List[str]:
    cleaned = str(name or "").strip()
    if not cleaned:
        raise ValueError("Der Kategoriename darf nicht leer sein.")
    with search_categories_lock:
        existing = storage.list_search_categories()
        existing.append(cleaned)
        categories = _normalise_category_list(existing)
        storage.save_search_categories(categories)
        return list(categories)


def delete_search_category(name: str) -> List[str]:
    cleaned = str(name or "").strip()
    if not cleaned:
        raise ValueError("Ungültiger Kategoriename.")
    if cleaned.lower() in {default.lower() for default in DEFAULT_SEARCH_CATEGORIES}:
        raise ValueError("Die Standardkategorie kann nicht gelöscht werden.")
    with search_categories_lock:
        existing = [
            item
            for item in storage.list_search_categories()
            if str(item or "").strip().lower() != cleaned.lower()
        ]
        categories = _normalise_category_list(existing)
        storage.save_search_categories(categories)
        return list(categories)


def normalise_keywords(values: Iterable[object]) -> List[str]:
    cleaned: List[str] = []
    seen: Set[str] = set()
    for value in values:
        text = str(value or "").strip()
        if not text:
            continue
        lowered = text.lower()
        if lowered in seen:
            continue
        seen.add(lowered)
        cleaned.append(text)
    return cleaned


def normalise_school_ids(values: Iterable[object]) -> List[str]:
    cleaned: List[str] = []
    seen: Set[str] = set()
    for value in values:
        text = str(value or "").strip()
        if not text:
            continue
        if text in seen:
            continue
        seen.add(text)
        cleaned.append(text)
    return cleaned


def update_search_definition_meta(definition_id: str, **updates: Any) -> Optional[Dict[str, Any]]:
    if not definition_id:
        return None
    definition = get_search_definition(definition_id)
    if not definition:
        return None
    for key, value in updates.items():
        if value is not None:
            definition[key] = value
    definition["updated_at"] = utcnow_iso()
    save_search_definition(definition)
    return definition


def store_search_definition(payload: Dict[str, Any]) -> Dict[str, Any]:
    defaults = get_search_defaults()
    identifier = str(payload.get("id") or uuid.uuid4())
    existing = get_search_definition(identifier)
    definition: Dict[str, Any] = dict(existing) if isinstance(existing, dict) else {}
    created_at = definition.get("created_at") or utcnow_iso()

    name = str(payload.get("name") or definition.get("name") or "").strip()
    if not name:
        raise ValueError("Bitte vergeben Sie einen Namen für die Suche.")

    description = str(payload.get("description") or definition.get("description") or "").strip()

    keywords_raw = payload.get("keywords")
    if isinstance(keywords_raw, str):
        lines = [item.strip() for item in keywords_raw.splitlines() if item.strip()]
        keywords = normalise_keywords(lines)
    elif isinstance(keywords_raw, list):
        keywords = normalise_keywords(keywords_raw)
    else:
        keywords = normalise_keywords(definition.get("keywords", []))
    if not keywords:
        raise ValueError("Bitte geben Sie mindestens ein Suchstichwort an.")

    school_ids_raw = payload.get("school_ids")
    if isinstance(school_ids_raw, list):
        school_ids = normalise_school_ids(school_ids_raw)
    else:
        school_ids = normalise_school_ids(definition.get("school_ids", []))

    prompt_text = str(payload.get("prompt") or definition.get("prompt") or defaults["prompt"]).strip()
    model_value = str(payload.get("model") or definition.get("model") or defaults["model"]).strip()

    temperature_raw = payload.get("temperature", definition.get("temperature", defaults["temperature"]))
    try:
        temperature_value = float(temperature_raw)
    except (TypeError, ValueError):
        temperature_value = defaults["temperature"]
    temperature_value = max(0.0, min(temperature_value, 2.0))

    max_results_raw = payload.get("max_results", definition.get("max_results", defaults["max_results"]))
    try:
        max_results_value = int(max_results_raw)
    except (TypeError, ValueError):
        max_results_value = defaults["max_results"]
    max_results_value = max(1, min(max_results_value, 200))

    concurrency_raw = payload.get("concurrency", definition.get("concurrency", defaults["concurrency"]))
    try:
        concurrency_value = int(concurrency_raw)
    except (TypeError, ValueError):
        concurrency_value = defaults["concurrency"]
    if concurrency_value < 1:
        concurrency_value = 1

    category_value = str(payload.get("category") or definition.get("category") or "").strip()
    categories = get_search_categories()
    if not category_value and categories:
        category_value = categories[0]
    status_candidate = str(payload.get("status") or definition.get("status") or SEARCH_STATUS_NEW)
    if status_candidate not in {SEARCH_STATUS_NEW, SEARCH_STATUS_RUNNING, SEARCH_STATUS_FINISHED}:
        status_candidate = SEARCH_STATUS_NEW

    definition.update(
        {
            "id": identifier,
            "name": name,
            "description": description,
            "keywords": keywords,
            "school_ids": school_ids,
            "prompt": prompt_text,
            "model": model_value,
            "temperature": temperature_value,
            "max_results": max_results_value,
            "concurrency": concurrency_value,
            "category": category_value or (categories[0] if categories else ""),
            "status": status_candidate,
            "created_at": created_at,
            "updated_at": utcnow_iso(),
        }
    )

    save_search_definition(definition)
    log_developer_event(
        "search-definition",
        "Suche gespeichert",
        details={
            "id": identifier,
        "name": name,
        "keywords": len(keywords),
        "schools": len(school_ids) if school_ids else None,
        "max_results": max_results_value,
        "model": model_value,
        "concurrency": concurrency_value,
    },
    )
    return definition


def _read_data_quality_results_unlocked() -> Dict[str, Dict[str, object]]:
    data = storage.get_json("data_quality", "results", {})
    if isinstance(data, dict):
        cleaned: Dict[str, Dict[str, object]] = {}
        for key, value in data.items():
            if isinstance(key, str) and isinstance(value, dict):
                cleaned[key] = value
        return cleaned
    return {}


def load_data_quality_results() -> Dict[str, Dict[str, object]]:
    with data_quality_lock:
        return dict(_read_data_quality_results_unlocked())


def save_data_quality_results(payload: Dict[str, Dict[str, object]]) -> None:
    with data_quality_lock:
        data = {key: value for key, value in payload.items() if isinstance(key, str) and isinstance(value, dict)}
        storage.set_json("data_quality", "results", data)


def replace_data_quality_results(payload: Dict[str, Dict[str, object]]) -> None:
    with data_quality_lock:
        data = payload if isinstance(payload, dict) else {}
        storage.set_json("data_quality", "results", data)


def update_data_quality_record(school_id: str, updates: Dict[str, object]) -> Dict[str, object]:
    if not school_id:
        return {}
    with data_quality_lock:
        data = _read_data_quality_results_unlocked()
        record = data.get(school_id, {}) if isinstance(data, dict) else {}
        if not isinstance(record, dict):
            record = {}
        updates_copy = dict(updates)
        force_scope = bool(updates_copy.pop("_force_site_scope", False))
        manual_existing = bool(record.get("site_scope_manual"))
        manual_update = updates_copy.pop("site_scope_manual", None)
        if manual_update is True:
            record["site_scope_manual"] = True
        elif manual_update is False:
            record["site_scope_manual"] = False
        elif "site_scope_manual" not in record:
            record["site_scope_manual"] = False
        if "site_scope" in updates_copy:
            updates_copy["site_scope"] = normalise_site_scope(updates_copy.get("site_scope"))
        if manual_existing and manual_update is None and not force_scope:
            updates_copy.pop("site_scope", None)
            updates_copy.pop("site_scope_reason", None)
            updates_copy.pop("site_scope_source", None)
            updates_copy.pop("site_scope_updated_at", None)
        record.update(updates_copy)
        data[school_id] = record
        storage.set_json("data_quality", "results", data)
        return dict(record)


def append_data_quality_history(school_id: str, entry: Dict[str, object]) -> None:
    if not school_id:
        return
    if not isinstance(entry, dict):
        return
    with data_quality_lock:
        data = _read_data_quality_results_unlocked()
        record = data.get(school_id)
        if not isinstance(record, dict):
            record = {}
        history = record.get("history") if isinstance(record.get("history"), list) else []
        history = list(history)
        entry_with_ts = dict(entry)
        entry_with_ts.setdefault("timestamp", datetime.utcnow().isoformat() + "Z")
        history.append(entry_with_ts)
        if len(history) > MAX_DATA_QUALITY_HISTORY:
            history = history[-MAX_DATA_QUALITY_HISTORY:]
        record["history"] = history
        data[school_id] = record
        storage.set_json("data_quality", "results", data)


def _apply_data_quality_url(
    school_id: str,
    new_url: str,
    *,
    dry_run: bool,
    manual_note_suffix: Optional[str] = None,
    history_stage: str = "manual",
    history_status: str = "accepted",
    history_detail: Optional[str] = "URL-Vorschlag übernommen",
    history_extra: Optional[Dict[str, object]] = None,
) -> bool:
    target_url = str(new_url or "").strip()
    if not school_id or not target_url:
        return False

    updated = True
    if not dry_run:
        updated = update_stammdaten_url(school_id, target_url)
    if not updated:
        return False

    timestamp = datetime.utcnow().isoformat() + "Z"
    note = "Richtiger Wert übernommen"
    if manual_note_suffix:
        note += f" ({manual_note_suffix})"
    if dry_run:
        note += " (Dry-Run)"

    update_data_quality_record(
        school_id,
        {
            "status": QUALITY_STATUS_OK,
            "original_url": target_url,
            "last_checked": timestamp,
            "manual_note": note,
            "suggested_url": None,
            "suggested_confidence": None,
            "suggested_reason": "",
            "suggested_signals": [],
            "correction_applied": True,
            "correction_applied_at": timestamp,
        },
    )

    history_entry: Dict[str, object] = {
        "stage": history_stage,
        "status": history_status,
    }
    if history_detail:
        detail_text = str(history_detail)
        if dry_run:
            detail_text += " (Dry-Run)"
        history_entry["detail"] = detail_text
    elif dry_run:
        history_entry["detail"] = "Dry-Run"
    if history_extra:
        history_entry.update(history_extra)

    append_data_quality_history(school_id, history_entry)
    return True


def _normalize_header(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value or "")
    normalized = "".join(ch for ch in normalized if not unicodedata.combining(ch))
    normalized = normalized.lower().replace("-", " ").replace("_", " ")
    return " ".join(normalized.split())


def normalise_site_scope(value: object) -> str:
    text = str(value or "").strip().lower()
    if not text:
        return SITE_SCOPE_UNKNOWN
    if text in {"exclusive", "exklusiv", "eigenständig", "nur schule", "schule"}:
        return SITE_SCOPE_EXCLUSIVE
    if text in {"shared", "geteilt", "unterseite", "gemeinsam", "portal", "fremd"}:
        return SITE_SCOPE_SHARED
    if "exklus" in text:
        return SITE_SCOPE_EXCLUSIVE
    if "teil" in text or "unterseite" in text or "gemeinde" in text:
        return SITE_SCOPE_SHARED
    return SITE_SCOPE_UNKNOWN


def quality_status_label(status: str) -> str:
    mapping = {
        QUALITY_STATUS_PENDING: "Noch nicht geprüft",
        QUALITY_STATUS_OK: "OK",
        QUALITY_STATUS_UNSURE: "Unsicher",
        QUALITY_STATUS_INVALID: "Falsch",
    }
    return mapping.get(status, "Unbekannt")


def site_scope_label(value: str) -> str:
    return SITE_SCOPE_LABELS.get(value or "", SITE_SCOPE_LABELS[SITE_SCOPE_UNKNOWN])


def site_scope_class(value: str) -> str:
    return SITE_SCOPE_CLASSES.get(value or "", SITE_SCOPE_CLASSES[SITE_SCOPE_UNKNOWN])


def build_data_quality_dataset() -> List[Dict[str, object]]:
    records = load_stammdaten()
    quality = load_data_quality_results()
    dataset: List[Dict[str, object]] = []
    for entry in records:
        if not isinstance(entry, dict):
            continue
        schul_id = str(entry.get("schul_id") or "").strip()
        if not schul_id:
            continue
        quality_entry = quality.get(schul_id) if isinstance(quality, dict) else None
        if not isinstance(quality_entry, dict):
            quality_entry = {}
        status = str(quality_entry.get("status") or QUALITY_STATUS_PENDING)
        formatted = dict(quality_entry)
        formatted.setdefault("status", status)
        formatted.setdefault("status_label", quality_status_label(status))
        formatted.setdefault("last_checked", None)
        formatted.setdefault("primary_reason", "")
        formatted.setdefault("confidence", None)
        formatted.setdefault("suggested_url", None)
        formatted.setdefault("suggested_confidence", None)
        formatted.setdefault("suggested_reason", "")
        formatted.setdefault("manual_note", "")
        formatted.setdefault("correction_applied", False)
        formatted.setdefault("correction_applied_at", None)
        scope_value = normalise_site_scope(formatted.get("site_scope"))
        formatted["site_scope"] = scope_value
        formatted["site_scope_label"] = site_scope_label(scope_value)
        formatted["site_scope_class"] = site_scope_class(scope_value)
        formatted.setdefault("site_scope_reason", str(formatted.get("site_scope_reason") or ""))
        formatted.setdefault("site_scope_updated_at", formatted.get("site_scope_updated_at"))
        dataset.append(
            {
                "schul_id": schul_id,
                "schulname": entry.get("schulname"),
                "ort": entry.get("ort"),
                "homepage": entry.get("homepage"),
                "status": formatted["status"],
                "status_label": formatted["status_label"],
                "correction_applied": bool(formatted.get("correction_applied")),
                "site_scope": formatted["site_scope"],
                "site_scope_label": formatted["site_scope_label"],
                "site_scope_class": formatted["site_scope_class"],
                "site_scope_reason": formatted.get("site_scope_reason"),
                "quality": formatted,
            }
        )
    return dataset


def update_stammdaten_url(school_id: str, new_url: str) -> bool:
    sanitized_id = str(school_id or "").strip()
    new_value = str(new_url or "").strip()
    if not sanitized_id or not new_value:
        return False
    records = load_stammdaten()
    updated = False
    for record in records:
        if not isinstance(record, dict):
            continue
        candidate = str(record.get("schul_id") or "").strip()
        if candidate == sanitized_id:
            record["homepage"] = new_value
            updated = True
            break
    if updated:
        save_stammdaten(records)
    return updated


def _coerce_str(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return text


def _coerce_int(value: object) -> Optional[int]:
    if value is None:
        return None
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, (int,)):
        return int(value)
    if isinstance(value, float):
        if math.isnan(value):
            return None
        return int(value)
    text = str(value).strip()
    if not text:
        return None
    cleaned = text.replace(".", "").replace(" ", "")
    try:
        return int(cleaned)
    except ValueError:
        try:
            return int(float(cleaned.replace(",", ".")))
        except ValueError:
            return None


def _coerce_bool(value: object, *, default: bool = True) -> bool:
    if value is None:
        return default
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        if isinstance(value, float) and math.isnan(value):
            return default
        return bool(int(value))
    text = str(value).strip().lower()
    if not text:
        return default
    if text in {"1", "true", "wahr", "ja", "j", "y", "aktiv", "active", "x"}:
        return True
    if text in {"0", "false", "falsch", "nein", "n", "inactive", "inaktiv"}:
        return False
    return default


def load_stammdaten() -> List[Dict[str, object]]:
    with stammdaten_lock:
        raw_data = storage.get_json("stammdaten", "records", [])
        if not isinstance(raw_data, list):
            return []

        records: List[Dict[str, object]] = []
        for entry in raw_data:
            if not isinstance(entry, dict):
                continue

            record: Dict[str, object] = {}
            for field in STAMMDATEN_FIELDS:
                key = field["key"]
                field_type = field.get("type", "text")
                value = entry.get(key)
                if value is None:
                    for alias in LEGACY_STAMMDATEN_ALIASES.get(key, []):
                        if alias in entry:
                            value = entry.get(alias)
                            if value is not None:
                                break
                if field_type == "numeric":
                    record[key] = _coerce_int(value)
                elif field_type == "bool":
                    record[key] = _coerce_bool(value, default=False)
                else:
                    record[key] = _coerce_str(value)

            homepage = _coerce_str(
                record.get("homepage")
                or entry.get("homepage")
                or entry.get("url")
            )
            if not homepage:
                continue
            record["homepage"] = homepage

            if not record.get("schulname"):
                record["schulname"] = _coerce_str(entry.get("name"))

            records.append(record)
        return records


def save_stammdaten(records: List[Dict[str, object]]) -> None:
    with stammdaten_lock:
        storage.set_json("stammdaten", "records", records)


def parse_stammdaten_excel(file_storage) -> List[Dict[str, object]]:
    try:
        file_storage.stream.seek(0)
    except (AttributeError, OSError):  # pragma: no cover - defensive
        pass

    try:
        workbook = load_workbook(file_storage, data_only=True)
    except Exception as exc:  # pragma: no cover - convert to user-facing message
        raise ValueError("Die Excel-Datei konnte nicht gelesen werden.") from exc

    worksheet = workbook.active
    rows = list(worksheet.iter_rows(values_only=True))
    if not rows:
        raise ValueError("Die Excel-Datei enthält keine Daten.")

    header_row = rows[0]
    header_map = {
        _normalize_header(str(cell) if cell is not None else ""): index
        for index, cell in enumerate(header_row)
    }

    expected_headers: Dict[str, Set[str]] = {}
    for field in STAMMDATEN_FIELDS:
        key = field["key"]
        variants = expected_headers.setdefault(key, set())
        variants.add(_normalize_header(field["label"]))
        for alias in LEGACY_STAMMDATEN_ALIASES.get(key, []):
            variants.add(_normalize_header(str(alias)))

    column_map: Dict[str, int] = {}
    for key, options in expected_headers.items():
        for normalized in options:
            if normalized in header_map:
                column_map[key] = header_map[normalized]
                break

    if len(column_map) < len(STAMMDATEN_FIELDS):
        for field in STAMMDATEN_FIELDS:
            column_map.setdefault(field["key"], None)

    records: List[Dict[str, object]] = []
    for raw_row in rows[1:]:
        if raw_row is None:
            continue
        values = list(raw_row)
        if not any(cell not in (None, "") for cell in values):
            continue

        record: Dict[str, object] = {}
        for field in STAMMDATEN_FIELDS:
            column_index = column_map.get(field["key"])
            cell_value = (
                values[column_index]
                if column_index is not None and column_index < len(values)
                else None
            )
            field_type = field.get("type", "text")
            if field_type == "numeric":
                record[field["key"]] = _coerce_int(cell_value)
            elif field_type == "bool":
                record[field["key"]] = _coerce_bool(cell_value, default=False)
            else:
                record[field["key"]] = _coerce_str(cell_value)

        if record.get("homepage"):
            records.append(record)

    if not records:
        raise ValueError("Es konnten keine gültigen Stammdaten gefunden werden.")

    return records


def collect_backup_snapshot() -> Dict[str, object]:
    return {
        "timestamp": datetime.utcnow().isoformat() + "Z",
        "settings": load_settings_data(),
        "stammdaten": load_stammdaten(),
        "saved_searches": load_saved_searches(),
        "synonym_cache": load_synonym_cache(),
        "keyword_finder_cache": load_keyword_finder_cache(),
        "data_quality_results": load_data_quality_results(),
        "searches": collect_all_search_data(),
    }


def perform_backup() -> None:
    snapshot = collect_backup_snapshot()
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    try:
        with BACKUP_FILE.open("w", encoding="utf-8") as handle:
            json.dump(snapshot, handle, ensure_ascii=False, indent=2)
    except OSError:
        pass


def restore_from_backup() -> None:
    if not BACKUP_FILE.exists():
        return
    try:
        with BACKUP_FILE.open("r", encoding="utf-8") as handle:
            snapshot = json.load(handle)
    except (json.JSONDecodeError, OSError):
        return
    if not isinstance(snapshot, dict):
        return

    settings_data = snapshot.get("settings")
    if settings_data and not load_settings_data():
        if isinstance(settings_data, dict):
            save_settings_data(settings_data)

    stammdaten_data = snapshot.get("stammdaten")
    if stammdaten_data and not load_stammdaten():
        if isinstance(stammdaten_data, list):
            save_stammdaten(stammdaten_data)

    saved_searches_data = snapshot.get("saved_searches")
    if saved_searches_data and not load_saved_searches():
        if isinstance(saved_searches_data, list):
            replace_saved_searches(saved_searches_data)

    synonym_cache_data = snapshot.get("synonym_cache")
    if synonym_cache_data and not load_synonym_cache():
        if isinstance(synonym_cache_data, dict):
            replace_synonym_cache(synonym_cache_data)

    keyword_finder_cache_data = snapshot.get("keyword_finder_cache")
    if keyword_finder_cache_data and not load_keyword_finder_cache():
        if isinstance(keyword_finder_cache_data, dict):
            replace_keyword_finder_cache(keyword_finder_cache_data)

    data_quality_results_data = snapshot.get("data_quality_results")
    if data_quality_results_data and not load_data_quality_results():
        if isinstance(data_quality_results_data, dict):
            replace_data_quality_results(data_quality_results_data)

    search_snapshot = snapshot.get("searches")
    if isinstance(search_snapshot, dict):
        if not load_search_definitions() and not list_search_runs():
            apply_searches_import(search_snapshot, "replace", dry_run=False)


def _backup_worker() -> None:
    while True:
        try:
            perform_backup()
        except Exception:
            pass
        time.sleep(BACKUP_INTERVAL_SECONDS)


def ensure_backup_thread() -> None:
    global _backup_thread_started
    if _backup_thread_started:
        return
    _backup_thread_started = True
    thread = threading.Thread(target=_backup_worker, daemon=True)
    thread.start()


restore_from_backup()
ensure_backup_thread()


@dataclass
class CrawlJob:
    """Container for the state of a running crawl job."""

    id: str
    start_urls: List[str]
    keywords: List[str]
    max_pages: int
    concurrency: int
    total_start_urls: int
    original_keywords: List[str] = field(default_factory=list)
    expanded_keywords: List[str] = field(default_factory=list)
    keyword_groups: List[Dict[str, object]] = field(default_factory=list)
    keyword_evaluations: List[Dict[str, object]] = field(default_factory=list)
    synonym_messages: List[Dict[str, str]] = field(default_factory=list)
    synonyms_enabled: bool = True
    synonyms_expanded: bool = False
    synonyms_manual: bool = False
    synonym_status: str = "noch nicht gestartet"
    start_date: Optional[date] = None
    end_date: Optional[date] = None
    respect_robots: bool = DEFAULT_RESPECT_ROBOTS
    processed_start_urls: int = 0
    current_start_url: Optional[str] = None
    current_page: Optional[str] = None
    visited_in_current: int = 0
    queue_length: int = 0
    recent_pages: List[str] = field(default_factory=list)
    status: str = "pending"  # pending, running, cancelling, cancelled, finished, error
    error: Optional[str] = None
    results: List[CrawlResult] = field(default_factory=list)
    found_count: int = 0
    progress_percent: int = 0
    started_at: float = field(default_factory=time.time)
    completed: bool = False
    result_pairs: Set[Tuple[str, str, str]] = field(default_factory=set, repr=False, compare=False)
    cancel_event: threading.Event = field(default_factory=threading.Event, repr=False, compare=False)
    cancelled: bool = False
    _lock: threading.Lock = field(default_factory=threading.Lock, repr=False, compare=False)

    def as_dict(self) -> Dict[str, object]:
        with self._lock:
            return {
                "job_id": self.id,
                "status": self.status,
                "error": self.error,
                "total_start_urls": self.total_start_urls,
                "processed_start_urls": self.processed_start_urls,
                "current_start_url": self.current_start_url,
                "current_page": self.current_page,
                "visited_in_current": self.visited_in_current,
                "queue_length": self.queue_length,
                "found_count": self.found_count,
                "progress_percent": self.progress_percent,
                "completed": self.completed,
                "recent_pages": list(self.recent_pages),
                "start_date": self.start_date.isoformat() if self.start_date else None,
                "end_date": self.end_date.isoformat() if self.end_date else None,
                "cancelled": self.cancelled,
                "concurrency": self.concurrency,
                "respect_robots": self.respect_robots,
                "original_keywords": list(self.original_keywords),
                "expanded_keywords": list(self.expanded_keywords),
                "keyword_groups": [
                    {
                        "original": group.get("original"),
                        "additional": list(group.get("additional", [])),
                        "synonyms": list(group.get("synonyms", [])),
                        "related": list(group.get("related", [])),
                        "related_details": [
                            {
                                "term": detail.get("term"),
                                "cluster": detail.get("cluster"),
                                "search_volume": detail.get("search_volume"),
                                "competition": detail.get("competition"),
                                "relevance": detail.get("relevance"),
                            }
                            for detail in group.get("related_details", [])
                            if isinstance(detail, dict)
                        ],
                    }
                    for group in self.keyword_groups
                ],
                "keyword_evaluations": [
                    {
                        "keyword": entry.get("keyword"),
                        "quality": entry.get("quality"),
                        "better_keywords": list(entry.get("better_keywords", [])),
                    }
                    for entry in self.keyword_evaluations
                ],
                "synonym_messages": list(self.synonym_messages),
                "synonyms_enabled": self.synonyms_enabled,
                "synonyms_expanded": self.synonyms_expanded,
                "synonyms_manual": self.synonyms_manual,
                "synonym_status": self.synonym_status,
                "results": [
                    {
                        "source_url": result.source_url,
                        "target_url": result.target_url,
                        "matched_keywords": list(result.matched_keywords),
                        "context": result.context,
                    }
                    for result in self.results
                ],
            }

    def update_progress(self, progress: CrawlProgress) -> None:
        with self._lock:
            if progress.queue_length is not None:
                self.queue_length = progress.queue_length

            if progress.current_url:
                self.current_page = progress.current_url
                if progress.event in {"page", "visited"}:
                    if not self.recent_pages or self.recent_pages[-1] != progress.current_url:
                        self.recent_pages.append(progress.current_url)
                        if len(self.recent_pages) > 10:
                            self.recent_pages = self.recent_pages[-10:]

            if progress.event == "visited" and progress.visited is not None:
                self.visited_in_current = progress.visited

            if progress.event == "match" and progress.result:
                keywords_tuple = tuple(progress.result.matched_keywords)
                lowered = keywords_tuple[0].lower() if keywords_tuple else ""
                context_signature = " ".join((progress.result.context or "").split()).lower()
                dedupe_signature = context_signature or lowered
                key = (
                    progress.result.source_url,
                    lowered,
                    dedupe_signature,
                )
                if key not in self.result_pairs:
                    self.results.append(progress.result)
                    self.result_pairs.add(key)
                self.found_count = len(self.results)

            pages_total = max(1, self.total_start_urls * self.max_pages)
            overall_pages = self.processed_start_urls * self.max_pages + self.visited_in_current
            self.progress_percent = min(100, int((overall_pages / pages_total) * 100))

    def mark_start_url_finished(self) -> None:
        with self._lock:
            self.processed_start_urls += 1
            self.visited_in_current = 0
            self.current_page = None
            self.queue_length = 0
            pages_total = max(1, self.total_start_urls * self.max_pages)
            overall_pages = self.processed_start_urls * self.max_pages
            self.progress_percent = min(100, int((overall_pages / pages_total) * 100))

    def mark_completed(self, *, error: Optional[str] = None) -> None:
        with self._lock:
            if error:
                self.status = "error"
                self.error = error
            else:
                self.status = "finished"
                self.progress_percent = 100
            self.completed = True

    def mark_cancelled(self) -> None:
        with self._lock:
            self.status = "cancelled"
            self.cancelled = True
            self.completed = True
            self.queue_length = 0
            self.current_page = None
            self.visited_in_current = 0

    def request_cancel(self) -> None:
        with self._lock:
            if not self.completed:
                self.status = "cancelling"
        self.cancel_event.set()


jobs: Dict[str, CrawlJob] = {}


@dataclass
class DataQualityJob:
    id: str
    school_ids: List[str]
    total: int
    mode: str = "full"
    processed: int = 0
    status: str = "pending"  # pending, running, cancelled, finished, error
    current_school: Optional[str] = None
    progress_percent: int = 0
    error: Optional[str] = None
    completed: bool = False
    results: List[Dict[str, object]] = field(default_factory=list)
    started_at: float = field(default_factory=time.time)
    messages: List[str] = field(default_factory=list)
    cancel_event: threading.Event = field(default_factory=threading.Event, repr=False, compare=False)
    _lock: threading.Lock = field(default_factory=threading.Lock, repr=False, compare=False)

    def as_dict(self) -> Dict[str, object]:
        with self._lock:
            return {
                "job_id": self.id,
                "status": self.status,
                "current_school": self.current_school,
                "processed": self.processed,
                "total": self.total,
                "progress_percent": self.progress_percent,
                "error": self.error,
                "completed": self.completed,
                "results": list(self.results),
                "messages": list(self.messages),
            }

    def update_progress(
        self,
        *,
        current_school: Optional[str] = None,
        processed_increment: int = 0,
        result: Optional[Dict[str, object]] = None,
        message: Optional[str] = None,
    ) -> None:
        with self._lock:
            if current_school is not None:
                self.current_school = current_school
            if processed_increment:
                self.processed += processed_increment
            if self.total:
                self.progress_percent = int((self.processed / self.total) * 100)
            if result:
                self.results.append(result)
                if len(self.results) > 50:
                    self.results = self.results[-50:]
            if message:
                self.messages.append(message)
                if len(self.messages) > 50:
                    self.messages = self.messages[-50:]

    def mark_completed(self, *, error: Optional[str] = None) -> None:
        with self._lock:
            if error:
                self.error = error
                self.status = "error"
            else:
                self.status = "finished"
            self.completed = True
            self.current_school = None
            if not self.progress_percent and self.total:
                self.progress_percent = int((self.processed / self.total) * 100)

    def mark_cancelled(self) -> None:
        with self._lock:
            self.status = "cancelled"
            self.completed = True
            self.current_school = None

    def request_cancel(self) -> None:
        self.cancel_event.set()


data_quality_jobs: Dict[str, DataQualityJob] = {}


@dataclass
class SearchJob:
    id: str
    definition_id: str
    definition_name: str
    keywords: List[str]
    school_ids: List[str]
    prompt: str
    model: str
    temperature: float
    max_results: int
    concurrency: int
    status: str = "pending"
    total_tasks: int = 0
    processed_tasks: int = 0
    results_count: int = 0
    error: Optional[str] = None
    messages: List[str] = field(default_factory=list)
    current_school: Optional[str] = None
    current_keyword: Optional[str] = None
    current_url: Optional[str] = None
    started_at: float = field(default_factory=time.time)
    completed_at: Optional[float] = None
    completed: bool = False
    cancel_event: threading.Event = field(default_factory=threading.Event, repr=False, compare=False)
    _lock: threading.Lock = field(default_factory=threading.Lock, repr=False, compare=False)
    school_snapshot: Dict[str, Dict[str, Any]] = field(default_factory=dict, repr=False, compare=False)
    schools_with_hits: Set[str] = field(default_factory=set, repr=False, compare=False)

    def as_dict(self) -> Dict[str, Any]:
        with self._lock:
            progress = 0
            if self.total_tasks:
                progress = min(100, int((self.processed_tasks / self.total_tasks) * 100))
            return {
                "job_id": self.id,
                "definition_id": self.definition_id,
                "definition_name": self.definition_name,
                "status": self.status,
                "total_tasks": self.total_tasks,
                "processed_tasks": self.processed_tasks,
                "results": self.results_count,
                "error": self.error,
                "messages": list(self.messages),
                "current_school": self.current_school,
                "current_keyword": self.current_keyword,
                "current_url": self.current_url,
                "progress_percent": progress,
                "started_at": self.started_at,
                "completed_at": self.completed_at,
                "completed": self.completed,
                "concurrency": self.concurrency,
            }


search_jobs: Dict[str, SearchJob] = {}


def register_search_job(job: SearchJob) -> None:
    with search_jobs_lock:
        search_jobs[job.id] = job


def get_search_job(job_id: str) -> Optional[SearchJob]:
    with search_jobs_lock:
        return search_jobs.get(job_id)


def persist_search_job_state(job: SearchJob) -> None:
    with job._lock:
        payload = {
            "id": job.id,
            "definition_id": job.definition_id,
            "definition_name": job.definition_name,
            "status": job.status,
            "keywords": list(job.keywords),
            "school_ids": list(job.school_ids),
            "results": job.results_count,
            "processed_tasks": job.processed_tasks,
            "total_tasks": job.total_tasks,
            "messages": list(job.messages),
            "prompt": job.prompt,
            "model": job.model,
            "temperature": job.temperature,
            "max_results": job.max_results,
            "concurrency": job.concurrency,
            "started_at": datetime.utcfromtimestamp(job.started_at).isoformat() + "Z",
            "error": job.error,
            "school_snapshot": job.school_snapshot,
            "schools_with_hits": sorted(job.schools_with_hits),
        }
        if job.completed_at:
            payload["completed_at"] = datetime.utcfromtimestamp(job.completed_at).isoformat() + "Z"
        save_search_run(payload)


def run_crawl_job(job: CrawlJob) -> None:
    try:
        for index, start_url in enumerate(job.start_urls, start=1):
            if job.cancel_event.is_set():
                job.mark_cancelled()
                return
            with job._lock:
                job.status = "running"
                job.current_start_url = start_url
                job.current_page = start_url
                job.visited_in_current = 0

            def progress_callback(progress: CrawlProgress) -> None:
                job.update_progress(progress)

            crawl_site(
                start_url,
                job.keywords,
                max_pages=job.max_pages,
                progress_callback=progress_callback,
                start_date=job.start_date,
                end_date=job.end_date,
                cancel_event=job.cancel_event,
                max_workers=job.concurrency,
                respect_robots=job.respect_robots,
            )
            if job.cancel_event.is_set():
                job.mark_cancelled()
                return
            job.mark_start_url_finished()

        if job.cancel_event.is_set():
            job.mark_cancelled()
        else:
            job.mark_completed()
    except Exception as exc:  # pragma: no cover - defensive safety net
        job.mark_completed(error=str(exc))
        log_error_event(
            "crawl-job-error",
            str(exc),
            details={
                "job_id": job.id,
                "keywords": job.keywords,
                "current_start_url": getattr(job, "current_start_url", None),
            },
            stack="".join(traceback.format_exception(type(exc), exc, exc.__traceback__)),
        )



def run_data_quality_job(job: DataQualityJob) -> None:
    try:
        openai_key = get_api_key()
        if not openai_key:
            job.mark_completed(error="OpenAI-Schlüssel erforderlich. Bitte unter Einstellungen speichern.")
            return

        all_records = load_stammdaten()
        records_by_id: Dict[str, Dict[str, object]] = {}
        for entry in all_records:
            schul_id = str(entry.get("schul_id") or "").strip()
            if schul_id:
                records_by_id[schul_id] = entry

        settings = get_data_quality_settings()
        mode = getattr(job, "mode", "full") or "full"
        force_google = mode == "google_search"
        site_scope_only = mode == "site_scope"
        if site_scope_only:
            force_google = False
        google_credentials = get_google_search_credentials()
        google_key = google_credentials.get("api_key", "").strip()
        google_cx = google_credentials.get("cx", "").strip()
        if force_google and (not google_key or not google_cx):
            job.mark_completed(error="Google Search API-Konfiguration erforderlich.")
            return
        threshold = float(settings.get("confidence_threshold", DEFAULT_DATA_QUALITY_SETTINGS["confidence_threshold"]))
        allowed_domains = {
            _normalise_domain(item)
            for item in settings.get("allowed_domains", [])
            if isinstance(item, str) and item.strip()
        }
        blocked_domains = {
            _normalise_domain(item)
            for item in settings.get("blocked_domains", [])
            if isinstance(item, str) and item.strip()
        }
        try:
            parallel_checks = int(
                settings.get("parallel_checks", DEFAULT_DATA_QUALITY_SETTINGS["parallel_checks"])
            )
        except (TypeError, ValueError):
            parallel_checks = DEFAULT_DATA_QUALITY_SETTINGS["parallel_checks"]
        if site_scope_only:
            parallel_checks = max(parallel_checks, 40)
        parallel_checks = max(1, min(parallel_checks, 50))

        with job._lock:
            job.status = "running"

        executor = ThreadPoolExecutor(max_workers=parallel_checks)
        futures = []

        def process_school(raw_school_id: object) -> None:
            school_id = str(raw_school_id)
            if job.cancel_event.is_set():
                return

            record = records_by_id.get(school_id)
            if not record:
                job.update_progress(
                    processed_increment=1,
                    result={"schul_id": school_id, "status": "not-found"},
                    message=f"Keine Stammdaten für Schul-ID {school_id} gefunden.",
                )
                append_data_quality_history(
                    school_id,
                    {
                        "stage": "info",
                        "status": "fehlend",
                        "detail": "Keine Stammdaten gefunden.",
                    },
                )
                return

            homepage = str(record.get("homepage", "")).strip()
            schulname = str(record.get("schulname", "")).strip() or str(record.get("Schulname", "")).strip()
            ort = str(record.get("ort", "")).strip()

            if not homepage:
                job.update_progress(
                    processed_increment=1,
                    result={"schul_id": school_id, "status": "no-url"},
                    message=f"Keine URL in den Stammdaten für {schulname or school_id}.",
                )
                append_data_quality_history(
                    school_id,
                    {
                        "stage": "info",
                        "status": "fehlend",
                        "detail": "Keine URL in den Stammdaten hinterlegt.",
                    },
                )
                return

            job.update_progress(current_school=school_id)

            snapshot = collect_site_snapshot(homepage)
            location_match, location_context, mismatch_hint = _check_impressum_location(snapshot, ort)
            structured = snapshot.get("structured_signals")
            if not isinstance(structured, dict):
                structured = {}
            structured["impressum_location_match"] = location_match
            if location_context:
                structured["impressum_location_context"] = location_context
            if mismatch_hint:
                structured["impressum_location_mismatch"] = mismatch_hint
            snapshot["structured_signals"] = structured

            if not location_match:
                expected = ort or "(kein Ort hinterlegt)"
                hint = mismatch_hint or "Ort nicht gefunden"
                job.update_progress(
                    message=f"Ort-Abgleich fehlgeschlagen für {schulname or school_id}: erwartet {expected}, gefunden {hint}.",
                )

            llm_result = assess_school_website_with_llm(
                schulname,
                ort,
                homepage,
                snapshot,
                api_key=openai_key,
                settings=settings,
            )

            status_primary = QUALITY_STATUS_UNSURE
            primary_confidence: Optional[float] = None
            primary_reason = "Keine Bewertung verfügbar."
            primary_signals: List[str] = []
            site_scope = SITE_SCOPE_UNKNOWN
            site_scope_reason = ""

            if llm_result:
                status_primary = normalise_quality_decision(llm_result.get("bewertung", ""))
                try:
                    primary_confidence = float(llm_result.get("confidence"))
                except (TypeError, ValueError):
                    primary_confidence = None
                primary_reason = str(llm_result.get("begruendung", "")).strip() or primary_reason
                signals_raw = llm_result.get("gefundene_signale")
                if isinstance(signals_raw, list):
                    primary_signals = [str(item) for item in signals_raw if str(item).strip()]
                site_scope = normalise_site_scope(llm_result.get("site_scope"))
                site_scope_reason = str(llm_result.get("site_scope_reason") or "").strip()

            meets_threshold = primary_confidence is not None and primary_confidence >= threshold
            if status_primary == QUALITY_STATUS_OK and not meets_threshold:
                status_after_primary = QUALITY_STATUS_UNSURE
            elif status_primary == QUALITY_STATUS_INVALID and not meets_threshold:
                status_after_primary = QUALITY_STATUS_UNSURE
            else:
                status_after_primary = status_primary

            if location_match and location_context:
                primary_signals.append(f"Ort bestätigt: {location_context}")
            elif not location_match:
                primary_signals.append("Ort im Impressum stimmt nicht überein")
                if primary_reason:
                    primary_reason = primary_reason.rstrip('.') + ". Ort im Impressum stimmt nicht mit den Stammdaten überein."
                else:
                    primary_reason = "Ort im Impressum stimmt nicht mit den Stammdaten überein."
                status_after_primary = QUALITY_STATUS_INVALID
                if primary_confidence is not None and primary_confidence > threshold:
                    primary_confidence = threshold - 0.01

            history_entry_site = {
                "stage": "site",
                "status": status_primary,
                "confidence": primary_confidence,
                "reason": primary_reason,
                "signals": primary_signals,
                "location_match": location_match,
                "location_context": location_context,
                "location_mismatch": mismatch_hint,
                "snapshot": {
                    "title": snapshot.get("title"),
                    "description": snapshot.get("description"),
                    "main_excerpt": snapshot.get("main_excerpt"),
                    "impressum_excerpt": snapshot.get("impressum_excerpt"),
                    "structured_signals": snapshot.get("structured_signals"),
                    "errors": snapshot.get("errors"),
                    "fetched_urls": snapshot.get("fetched_urls"),
                },
                "site_scope": site_scope,
                "site_scope_reason": site_scope_reason,
            }
            append_data_quality_history(school_id, history_entry_site)

            suggestion_url: Optional[str] = None
            suggestion_confidence: Optional[float] = None
            suggestion_reason = ""
            suggestion_signals: List[str] = []
            stage_source = "site"

            should_run_google = force_google or status_after_primary != QUALITY_STATUS_OK
            if site_scope_only:
                should_run_google = False
            if should_run_google:
                if not google_key or not google_cx:
                    job.update_progress(
                        message="Google Search API-Konfiguration fehlt. Zweite Prüfung übersprungen.",
                    )
                else:
                    try:
                        max_setting = int(settings.get("max_search_results", 10))
                    except (TypeError, ValueError):
                        max_setting = 10
                    max_results = 100 if force_google else max_setting
                    max_results = max(1, min(max_results, 100))
                    query = f"{schulname} {ort}".strip()
                    if force_google:
                        job.update_progress(
                            message=f"Google-Suche mit {max_results} Treffern für {schulname or school_id} …",
                        )
                    search_results = perform_google_search(
                        query,
                        google_key,
                        google_cx,
                        max_results=max_results,
                        language=str(settings.get("search_language", "de")),
                        region=str(settings.get("search_region", "de")),
                    )
                    filtered_results = [
                        result
                        for result in search_results
                        if isinstance(result, dict)
                        and _domain_allowed(result.get("link", ""), allowed_domains, blocked_domains)
                    ]
                    llm_search = analyse_search_results_with_llm(
                        schulname,
                        ort,
                        homepage,
                        filtered_results,
                        api_key=openai_key,
                        settings=settings,
                    )
                    if llm_search:
                        stage_source = "google" if force_google else "site+google"
                        suggestion_url = str(llm_search.get("empfehlung") or "").strip() or None
                        try:
                            suggestion_confidence = float(llm_search.get("confidence"))
                        except (TypeError, ValueError):
                            suggestion_confidence = None
                        suggestion_reason = str(llm_search.get("begruendung", "")).strip()
                        signals_raw = llm_search.get("signale")
                        if isinstance(signals_raw, list):
                            suggestion_signals = [str(item) for item in signals_raw if str(item).strip()]
                        append_data_quality_history(
                            school_id,
                            {
                                "stage": "search",
                                "status": status_after_primary,
                                "suggestion": suggestion_url,
                                "confidence": suggestion_confidence,
                                "reason": suggestion_reason,
                                "signals": suggestion_signals,
                                "results": filtered_results,
                            },
                        )
                        if (
                            suggestion_url
                            and _domain_allowed(suggestion_url, allowed_domains, blocked_domains)
                            and suggestion_confidence is not None
                            and suggestion_confidence >= threshold
                        ):
                            status_after_primary = QUALITY_STATUS_INVALID
                    else:
                        append_data_quality_history(
                            school_id,
                            {
                                "stage": "search",
                                "status": "fehlgeschlagen",
                                "reason": "Keine verwertbare Antwort von der Zweitprüfung.",
                                "results": search_results,
                            },
                        )

            timestamp = datetime.utcnow().isoformat() + "Z"
            record_update = {
                "school_name": schulname,
                "school_location": ort,
                "original_url": homepage,
                "impressum_location_match": location_match,
                "impressum_location_context": location_context,
                "impressum_location_mismatch": mismatch_hint,
                "snapshot": {
                    "title": snapshot.get("title"),
                    "description": snapshot.get("description"),
                    "main_excerpt": snapshot.get("main_excerpt"),
                    "impressum_excerpt": snapshot.get("impressum_excerpt"),
                    "structured_signals": snapshot.get("structured_signals"),
                    "errors": snapshot.get("errors"),
                    "fetched_urls": snapshot.get("fetched_urls"),
                },
                "site_scope": site_scope,
                "site_scope_reason": site_scope_reason,
                "site_scope_updated_at": timestamp,
                "site_scope_source": "llm",
            }
            if site_scope_only:
                record_update["last_checked"] = timestamp
            else:
                record_update.update(
                    {
                        "status": status_after_primary,
                        "last_checked": timestamp,
                        "confidence": primary_confidence,
                        "primary_decision": status_primary,
                        "primary_confidence": primary_confidence,
                        "primary_reason": primary_reason,
                        "primary_signals": primary_signals,
                        "suggested_url": suggestion_url,
                        "suggested_confidence": suggestion_confidence,
                        "suggested_reason": suggestion_reason,
                        "suggested_signals": suggestion_signals,
                        "last_source": stage_source,
                    }
                )
            update_data_quality_record(school_id, record_update)

            progress_message: Optional[str] = None
            if site_scope_only:
                progress_message = (
                    f"Auftrittstyp {site_scope_label(site_scope)} für {schulname or school_id}"
                )
            job.update_progress(
                processed_increment=1,
                result={
                    "schul_id": school_id,
                    "school_name": schulname or school_id,
                    "status": status_after_primary,
                    "status_label": quality_status_label(status_after_primary),
                    "confidence": primary_confidence,
                    "suggested_url": suggestion_url,
                    "suggested_confidence": suggestion_confidence,
                    "suggested_reason": suggestion_reason,
                    "original_url": homepage,
                    "location_match": location_match,
                    "site_scope": site_scope,
                    "site_scope_label": site_scope_label(site_scope),
                },
                message=progress_message,
            )

        try:
            for school_id in job.school_ids:
                if job.cancel_event.is_set():
                    break
                futures.append(executor.submit(process_school, school_id))

            for future in as_completed(futures):
                if job.cancel_event.is_set():
                    break
                try:
                    future.result()
                except Exception as exc:
                    job.update_progress(message=f"Fehler bei der Datenqualitätsprüfung: {exc}")
                    log_error_event(
                        "data-quality-worker-error",
                        str(exc),
                        details={
                            "job_id": job.id,
                            "school": getattr(job, "current_school", None),
                        },
                        stack="".join(traceback.format_exception(type(exc), exc, exc.__traceback__)),
                    )
        finally:
            executor.shutdown(wait=True, cancel_futures=True)

        if job.cancel_event.is_set():
            job.mark_cancelled()
        else:
            job.mark_completed()

    except Exception as exc:  # pragma: no cover - defensive safety net
        job.mark_completed(error=str(exc))
        log_error_event(
            "data-quality-job-error",
            str(exc),
            details={"job_id": job.id},
            stack="".join(traceback.format_exception(type(exc), exc, exc.__traceback__)),
        )


def run_search_job(job: SearchJob) -> None:
    try:
        openai_key = get_api_key()
        if not openai_key:
            with job._lock:
                job.status = "error"
                job.error = "OpenAI-Schlüssel erforderlich. Bitte in den Einstellungen speichern."
                job.completed = True
                job.completed_at = time.time()
            persist_search_job_state(job)
            update_search_definition_meta(
                job.definition_id,
                status=SEARCH_STATUS_NEW,
                last_error=job.error,
            )
            return

        google_credentials = get_google_search_credentials()
        google_key = str(google_credentials.get("api_key") or "").strip()
        google_cx = str(google_credentials.get("cx") or "").strip()
        if not google_key or not google_cx:
            with job._lock:
                job.status = "error"
                job.error = "Google Search API-Schlüssel und CX erforderlich. Bitte in den Einstellungen speichern."
                job.completed = True
                job.completed_at = time.time()
            persist_search_job_state(job)
            update_search_definition_meta(
                job.definition_id,
                status=SEARCH_STATUS_NEW,
                last_error=job.error,
            )
            return

        stammdaten = load_stammdaten()
        records_by_id: Dict[str, Dict[str, Any]] = {}
        for entry in stammdaten:
            if not isinstance(entry, dict):
                continue
            schul_id = str(entry.get("schul_id") or "").strip()
            if schul_id:
                records_by_id[schul_id] = entry

        school_ids = job.school_ids or list(records_by_id.keys())
        keywords = job.keywords or []

        combined_keywords = " OR ".join(keywords)
        keyword_display = combined_keywords or (keywords[0] if keywords else "")
        tasks: List[Tuple[str, Dict[str, Any], Optional[str]]] = []
        for school_id in school_ids:
            record = records_by_id.get(school_id)
            if not record:
                with job._lock:
                    job.messages.append(f"Keine Stammdaten für Schul-ID {school_id} gefunden.")
                continue
            homepage = str(record.get("homepage") or record.get("url") or "").strip()
            if not homepage:
                with job._lock:
                    job.messages.append(
                        f"{record.get('schulname', school_id)} hat keine Homepage in den Stammdaten."
                    )
                continue
            domain = _extract_domain_from_url(homepage)
            job.school_snapshot[school_id] = {
                "name": record.get("schulname"),
                "city": record.get("ort"),
                "homepage": homepage,
            }
            tasks.append((school_id, record, domain))

        with job._lock:
            job.total_tasks = len(tasks)
            job.status = "running"
            job.messages.append(f"Starte Suchlauf mit {len(tasks)} Schulen.")
            job.current_keyword = keyword_display or None
        persist_search_job_state(job)
        update_search_definition_meta(job.definition_id, status=SEARCH_STATUS_RUNNING, last_error=None)

        if not tasks:
            with job._lock:
                job.status = "finished"
                job.completed = True
                job.completed_at = time.time()
            persist_search_job_state(job)
            update_search_definition_meta(
                job.definition_id,
                status=SEARCH_STATUS_FINISHED,
                last_run_id=job.id,
                last_run_at=utcnow_iso(),
                results=0,
            )
            log_developer_event(
                "search-run",
                "Keine Schulen für Suchlauf gefunden",
                details={"job_id": job.id, "definition_id": job.definition_id},
            )
            return

        log_developer_event(
            "search-run",
            "Suchlauf gestartet",
            details={
                "job_id": job.id,
                "definition_id": job.definition_id,
                "schools": len(tasks),
                "keywords": len(keywords),
                "concurrency": job.concurrency,
            },
        )

        def build_query(domain_value: Optional[str], school_record: Dict[str, Any]) -> str:
            if domain_value:
                if combined_keywords:
                    return f"site:{domain_value} ({combined_keywords})"
                return f"site:{domain_value}"
            base_terms = [
                str(school_record.get("schulname") or "").strip(),
                str(school_record.get("ort") or "").strip(),
            ]
            query_parts = [part for part in base_terms if part]
            if combined_keywords:
                query_parts.append(f"({combined_keywords})")
            return " ".join(query_parts) or combined_keywords or school_record.get("schul_id", "")

        def process_school(task: Tuple[str, Dict[str, Any], Optional[str]]) -> Dict[str, Any]:
            school_id, record, domain_value = task
            school_display = str(record.get("schulname") or school_id)
            if job.cancel_event.is_set():
                return {
                    "school_id": school_id,
                    "school_display": school_display,
                    "entries": [],
                    "message": f"Suche für {school_display} abgebrochen.",
                }
            query = build_query(domain_value, record)
            with job._lock:
                job.current_school = school_display
                job.current_keyword = keyword_display or query
                job.current_url = None
            log_developer_event(
                "search-run",
                "Google-Suche für Schule",
                details={"job_id": job.id, "school_id": school_id, "query": query},
            )
            search_results = perform_google_search(
                query,
                google_key,
                google_cx,
                max_results=job.max_results,
                language="de",
                region="de",
            )
            if job.cancel_event.is_set():
                return {
                    "school_id": school_id,
                    "school_display": school_display,
                    "entries": [],
                    "message": f"Suche für {school_display} abgebrochen.",
                }
            google_items = list(search_results or [])
            page_entries: List[Dict[str, Any]] = []
            google_log_entries: List[Dict[str, Any]] = []
            for index, item in enumerate(google_items, start=1):
                link = str(item.get("link") or "").strip()
                google_log_entries.append(
                    {
                        "position": index,
                        "title": item.get("title"),
                        "link": link,
                        "snippet": item.get("snippet"),
                    }
                )
                if job.cancel_event.is_set():
                    break
                if not link:
                    continue
                text = fetch_page_text(link)
                if not text:
                    continue
                evaluation = evaluate_page_with_llm(
                    text,
                    job.prompt,
                    openai_key,
                    model=job.model,
                    temperature=job.temperature,
                )
                if not evaluation:
                    continue
                avg_value = sum(entry.get("value", 0) for entry in evaluation) / max(len(evaluation), 1)
                avg_confidence = sum(entry.get("confidence", 0.0) for entry in evaluation) / max(len(evaluation), 1)
                page_entries.append(
                    {
                        "school_id": school_id,
                        "school_name": record.get("schulname"),
                        "school_city": record.get("ort"),
                        "school_homepage": record.get("homepage"),
                        "keyword": keyword_display or query,
                        "query": query,
                        "target_url": link,
                        "google_title": item.get("title"),
                        "google_snippet": item.get("snippet"),
                        "dimensions": evaluation,
                        "average_value": avg_value,
                        "average_confidence": avg_confidence,
                        "evaluated_at": utcnow_iso(),
                        "model": job.model,
                        "prompt": job.prompt,
                        "content_excerpt": text[:600],
                    }
                )
            last_url = page_entries[-1]["target_url"] if page_entries else None
            message_text: Optional[str] = None
            if not google_items:
                message_text = f"Keine Google-Ergebnisse für {school_display} gefunden."
            elif not page_entries:
                message_text = f"Keine verwertbaren Treffer für {school_display}."
            elif job.cancel_event.is_set():
                message_text = f"Suche für {school_display} abgebrochen."
            log_developer_event(
                "search-run",
                "Auswertung abgeschlossen",
                details={
                    "job_id": job.id,
                    "school_id": school_id,
                    "results": len(page_entries),
                    "google_results": len(google_items),
                    "last_url": last_url,
                },
            )
            log_entry = {
                "school_id": school_id,
                "school_name": record.get("schulname"),
                "school_city": record.get("ort"),
                "school_homepage": record.get("homepage"),
                "keyword": keyword_display or query,
                "query": query,
                "google_results": google_log_entries,
                "llm_outputs": [
                    {
                        "target_url": entry.get("target_url"),
                        "google_title": entry.get("google_title"),
                        "google_snippet": entry.get("google_snippet"),
                        "dimensions": entry.get("dimensions"),
                        "average_value": entry.get("average_value"),
                        "average_confidence": entry.get("average_confidence"),
                    }
                    for entry in page_entries
                ],
                "message": message_text,
                "created_at": utcnow_iso(),
            }
            return {
                "school_id": school_id,
                "school_display": school_display,
                "entries": page_entries,
                "last_url": last_url,
                "message": message_text,
                "log_entry": log_entry,
            }

        futures: Dict[Any, Tuple[str, Dict[str, Any], Optional[str]]] = {}
        executor = ThreadPoolExecutor(max_workers=job.concurrency)
        try:
            for task in tasks:
                futures[executor.submit(process_school, task)] = task
            for future in as_completed(futures):
                school_id, record, _ = futures[future]
                if job.cancel_event.is_set():
                    break
                try:
                    result_data = future.result()
                except Exception as exc:  # pragma: no cover - defensive
                    message = f"Fehler bei {record.get('schulname', school_id)}: {exc}"[:200]
                    with job._lock:
                        job.processed_tasks += 1
                        job.messages.append(message)
                    persist_search_job_state(job)
                    log_error_event(
                        "search-run",
                        "Fehler bei der Auswertung",
                        details={"job_id": job.id, "school_id": school_id},
                        stack="".join(traceback.format_exception(type(exc), exc, exc.__traceback__)),
                    )
                    continue
                log_entry = result_data.get("log_entry")
                if log_entry:
                    append_search_logs(job.id, [log_entry])
                entries = result_data.get("entries", [])
                message_text = result_data.get("message")
                if entries:
                    append_search_results(job.id, entries)
                    job.schools_with_hits.add(school_id)
                with job._lock:
                    job.processed_tasks += 1
                    job.current_school = result_data.get("school_display") or job.current_school
                    job.current_keyword = keyword_display or job.current_keyword
                    job.current_url = result_data.get("last_url")
                    if entries:
                        job.results_count += len(entries)
                    if message_text:
                        job.messages.append(message_text)
                    elif entries:
                        job.messages.append(
                            f"{len(entries)} Treffer für {result_data.get('school_display') or school_id}."
                        )
                persist_search_job_state(job)
                if job.cancel_event.is_set():
                    break
        finally:
            executor.shutdown(wait=False, cancel_futures=True)

        if job.cancel_event.is_set():
            with job._lock:
                job.status = "cancelled"
                job.completed = True
                job.completed_at = time.time()
            persist_search_job_state(job)
            update_search_definition_meta(job.definition_id, status=SEARCH_STATUS_NEW, last_error="Abgebrochen")
            log_developer_event(
                "search-run",
                "Suchlauf abgebrochen",
                details={"job_id": job.id, "processed": job.processed_tasks},
            )
            return

        with job._lock:
            job.status = "finished"
            job.completed = True
            job.completed_at = time.time()
        persist_search_job_state(job)
        update_search_definition_meta(
            job.definition_id,
            status=SEARCH_STATUS_FINISHED,
            last_run_id=job.id,
            last_run_at=utcnow_iso(),
            last_error=None,
            results=job.results_count,
        )
        log_developer_event(
            "search-run",
            "Suchlauf abgeschlossen",
            details={
                "job_id": job.id,
                "definition_id": job.definition_id,
                "results": job.results_count,
                "processed": job.processed_tasks,
            },
        )
    except Exception as exc:  # pragma: no cover - defensive safety net
        message = str(exc)
        with job._lock:
            job.status = "error"
            job.error = message
            job.completed = True
            job.completed_at = time.time()
        persist_search_job_state(job)
        update_search_definition_meta(
            job.definition_id,
            status=SEARCH_STATUS_NEW,
            last_error=message,
        )
        log_error_event(
            "search-job-error",
            message,
            details={
                "job_id": job.id,
                "definition_id": job.definition_id,
                "keywords": job.keywords,
                "current_school": getattr(job, "current_school", None),
            },
            stack="".join(traceback.format_exception(type(exc), exc, exc.__traceback__)),
        )
    finally:
        with search_jobs_lock:
            search_jobs.pop(job.id, None)


@app.route("/")
def index() -> ResponseReturnValue:
    records_all = load_stammdaten()
    active_records = [record for record in records_all if record.get("aktiv", True)]
    categories = get_search_categories()
    defaults = get_search_defaults()
    google_credentials = get_google_search_credentials()
    google_configured = bool(google_credentials.get("api_key") and google_credentials.get("cx"))
    prefill_id = request.args.get("definition")
    prefill_definition: Optional[Dict[str, Any]] = None
    if prefill_id:
        existing = get_search_definition(prefill_id)
        if existing:
            prefill_definition = {
                "id": existing.get("id"),
                "name": existing.get("name"),
                "description": existing.get("description"),
                "keywords": existing.get("keywords", []),
                "school_ids": existing.get("school_ids", []),
                "prompt": existing.get("prompt"),
                "max_results": existing.get("max_results", defaults.get("max_results")),
                "model": existing.get("model", defaults.get("model")),
                "temperature": existing.get("temperature", defaults.get("temperature")),
                "category": existing.get("category"),
                "concurrency": existing.get("concurrency", defaults.get("concurrency")),
            }
    return render_template(
        "index.html",
        active_page="search_create",
        categories=categories,
        default_category=categories[0] if categories else "",
        search_defaults=defaults,
        available_models=AVAILABLE_OPENAI_MODELS,
        stammdaten_records=active_records,
        stammdaten_fields=STAMMDATEN_FIELDS,
        stammdaten_primary_fields=STAMMDATEN_PRIMARY_FIELDS,
        stammdaten_boolean_fields=STAMMDATEN_BOOLEAN_FIELDS,
        has_api_key=has_api_key(),
        google_configured=google_configured,
        search_dimensions=SEARCH_DIMENSIONS,
        prefill_definition=prefill_definition,
    )


def _extract_keywords(payload: Dict[str, Any]) -> List[str]:
    raw = payload.get("keywords")
    if isinstance(raw, str):
        values = [line.strip() for line in raw.splitlines() if line.strip()]
    elif isinstance(raw, list):
        values = [str(item or "").strip() for item in raw if str(item or "").strip()]
    else:
        values = []
    return normalise_keywords(values)


def _extract_school_ids(payload: Dict[str, Any]) -> List[str]:
    raw = payload.get("school_ids")
    if isinstance(raw, list):
        return normalise_school_ids(raw)
    return []


@app.post("/searches/test")
def search_test() -> ResponseReturnValue:
    payload = request.get_json(silent=True) or {}
    if not isinstance(payload, dict):
        return jsonify({"error": "Ungültige Anfrage."}), 400
    keywords = _extract_keywords(payload)
    school_ids = _extract_school_ids(payload)
    if not keywords:
        return jsonify({"error": "Bitte geben Sie mindestens ein Keyword an."}), 400
    if not school_ids:
        return jsonify({"error": "Bitte wählen Sie mindestens eine Schule aus."}), 400
    stammdaten = load_stammdaten()
    records_by_id = {
        str(record.get("schul_id") or "").strip(): record
        for record in stammdaten
        if isinstance(record, dict) and str(record.get("schul_id") or "").strip()
    }
    queries: List[str] = []
    for school_id in school_ids[:10]:
        record = records_by_id.get(school_id)
        if not record:
            continue
        homepage = str(record.get("homepage") or "").strip()
        domain = _extract_domain_from_url(homepage) or homepage
        if not domain:
            domain = str(record.get("url") or "").strip()
        if not domain:
            domain = school_id
        query = f"site:{domain} ({' OR '.join(keywords)})"
        queries.append(query)
    if not queries:
        return jsonify({"error": "Für die ausgewählten Schulen konnte keine Domain ermittelt werden."}), 404
    return jsonify({
        "message": "Suchanfragen erstellt. Verwenden Sie die Probelauf-Funktion für eine detaillierte Prüfung.",
        "queries": queries,
    })


@app.post("/searches/probe")
def search_probe() -> ResponseReturnValue:
    payload = request.get_json(silent=True) or {}
    if not isinstance(payload, dict):
        return jsonify({"error": "Ungültige Anfrage."}), 400
    keywords = _extract_keywords(payload)
    school_ids = _extract_school_ids(payload)
    if not keywords:
        return jsonify({"error": "Bitte geben Sie mindestens ein Keyword an."}), 400
    if not school_ids:
        return jsonify({"error": "Bitte wählen Sie mindestens eine Schule aus."}), 400

    openai_key = get_api_key()
    if not openai_key:
        return jsonify({"error": "OpenAI-Schlüssel erforderlich. Bitte in den Einstellungen hinterlegen."}), 400

    defaults = get_search_defaults()
    prompt = str(payload.get("prompt") or defaults.get("prompt") or DEFAULT_SEARCH_EVALUATION_PROMPT)
    try:
        max_results = int(payload.get("max_results", defaults.get("max_results", 5)))
    except (TypeError, ValueError):
        max_results = defaults.get("max_results", 5)
    max_results = max(1, min(max_results, 20))
    model_value = str(payload.get("model") or defaults.get("model") or DEFAULT_SEARCH_MODEL)
    try:
        temperature_value = float(payload.get("temperature", defaults.get("temperature", DEFAULT_SEARCH_TEMPERATURE)))
    except (TypeError, ValueError):
        temperature_value = defaults.get("temperature", DEFAULT_SEARCH_TEMPERATURE)

    stammdaten = load_stammdaten()
    records_by_id = {
        str(record.get("schul_id") or "").strip(): record
        for record in stammdaten
        if isinstance(record, dict) and str(record.get("schul_id") or "").strip()
    }
    first_id = school_ids[0]
    record = records_by_id.get(first_id)
    if not record:
        return jsonify({"error": "Für die erste ausgewählte Schule wurden keine Stammdaten gefunden."}), 404

    homepage = str(record.get("homepage") or "").strip()
    domain = _extract_domain_from_url(homepage) or homepage
    if not domain:
        domain = str(record.get("url") or "").strip()
    if not domain:
        domain = record.get("schul_id") or ""
    keyword_query = ' OR '.join(keywords)
    query = f"site:{domain} ({keyword_query})" if domain else keyword_query

    google_credentials = get_google_search_credentials()
    google_key = str(google_credentials.get("api_key") or "").strip()
    google_cx = str(google_credentials.get("cx") or "").strip()
    search_results: List[Dict[str, Any]] = []
    if google_key and google_cx:
        search_results = perform_google_search(
            query,
            google_key,
            google_cx,
            max_results=max_results,
            language="de",
            region="de",
        )
    if not search_results and homepage:
        search_results = [{"link": homepage, "title": record.get("schulname"), "snippet": ""}]
    if not search_results:
        return jsonify({"message": "Keine Ergebnisse für diese Suche gefunden.", "preview": None, "status": "empty"})

    candidate = search_results[0]
    link = str(candidate.get("link") or "").strip()
    if not link:
        return jsonify({"error": "Das erste Suchergebnis enthält keine URL."}), 404
    content = fetch_page_text(link)
    if not content:
        return jsonify({"error": "Die gefundene Seite konnte nicht gelesen werden."}), 502
    evaluation = evaluate_page_with_llm(
        content,
        prompt,
        openai_key,
        model=model_value,
        temperature=temperature_value,
    )
    if not evaluation:
        return jsonify({"error": "Die KI konnte keine Bewertung erstellen."}), 502

    preview = {
        "school_id": first_id,
        "school_name": record.get("schulname"),
        "school_city": record.get("ort"),
        "target_url": link,
        "query": query,
        "dimensions": evaluation,
        "google_title": candidate.get("title"),
        "google_snippet": candidate.get("snippet"),
    }
    return jsonify({"message": "Probelauf abgeschlossen.", "preview": preview})


@app.get("/searches/categories")
def list_search_categories_route() -> ResponseReturnValue:
    return jsonify({"categories": get_search_categories()})


@app.post("/searches/categories")
def create_search_category_route() -> ResponseReturnValue:
    payload = request.get_json(silent=True) or {}
    name = str(payload.get("name") or "").strip()
    if not name:
        return jsonify({"error": "Bitte geben Sie einen Kategorienamen an."}), 400
    try:
        categories = add_search_category(name)
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    return jsonify({"categories": categories})


@app.delete("/searches/categories/<path:category_name>")
def delete_search_category_route(category_name: str) -> ResponseReturnValue:
    name = str(category_name or "").strip()
    if not name:
        return jsonify({"error": "Ungültiger Kategoriename."}), 400
    try:
        categories = delete_search_category(name)
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    return jsonify({"categories": categories})


@app.route("/searches")
def searches() -> ResponseReturnValue:
    records = [record for record in load_stammdaten() if record.get("aktiv", True)]
    definitions = load_search_definitions()
    runs = list_search_runs()
    defaults = get_search_defaults()
    google_credentials = get_google_search_credentials()
    google_configured = bool(google_credentials.get("api_key") and google_credentials.get("cx"))
    with search_jobs_lock:
        active_jobs = {job_id: job.as_dict() for job_id, job in search_jobs.items()}
    return render_template(
        "searches.html",
        active_page="search_results",
        schools=records,
        definitions=definitions,
        runs=runs,
        search_defaults=defaults,
        available_models=AVAILABLE_OPENAI_MODELS,
        has_api_key=has_api_key(),
        google_configured=google_configured,
        active_jobs=active_jobs,
        search_dimensions=SEARCH_DIMENSIONS,
    )


@app.get("/searches/data")
def searches_data() -> ResponseReturnValue:
    definitions = load_search_definitions()
    runs = list_search_runs()
    with search_jobs_lock:
        active_jobs = {job_id: job.as_dict() for job_id, job in search_jobs.items()}
    return jsonify({"definitions": definitions, "runs": runs, "jobs": active_jobs})


@app.post("/searches/definitions")
def create_or_update_search_definition() -> ResponseReturnValue:
    payload = request.get_json(silent=True) or {}
    if not isinstance(payload, dict):
        return jsonify({"error": "Ungültige Anfrage."}), 400
    try:
        definition = store_search_definition(payload)
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    return jsonify({"definition": definition})


@app.delete("/searches/definitions/<definition_id>")
def delete_search_definition_route(definition_id: str) -> ResponseReturnValue:
    identifier = str(definition_id or "").strip()
    if not identifier:
        return jsonify({"error": "Ungültige ID."}), 400
    with search_jobs_lock:
        for job in search_jobs.values():
            if job.definition_id == identifier and not job.completed:
                return jsonify({"error": "Suche läuft noch. Bitte warten Sie bis zum Ende oder brechen Sie ab."}), 409
    delete_search_definitions([identifier])
    return jsonify({"status": "ok"})


@app.post("/searches/definitions/<definition_id>/duplicate")
def duplicate_search_definition(definition_id: str) -> ResponseReturnValue:
    identifier = str(definition_id or "").strip()
    if not identifier:
        return jsonify({"error": "Ungültige ID."}), 400
    original = get_search_definition(identifier)
    if not original:
        return jsonify({"error": "Suche wurde nicht gefunden."}), 404
    duplicate = dict(original)
    duplicate.pop("id", None)
    duplicate["id"] = str(uuid.uuid4())
    duplicate["name"] = f"{original.get('name')} (Kopie)"
    duplicate["status"] = SEARCH_STATUS_NEW
    duplicate.pop("last_run_id", None)
    duplicate.pop("last_run_at", None)
    duplicate.pop("last_error", None)
    duplicate["created_at"] = utcnow_iso()
    duplicate["updated_at"] = duplicate["created_at"]
    save_search_definition(duplicate)
    return jsonify({"definition": duplicate})


@app.post("/searches/<definition_id>/run")
def start_search_run(definition_id: str) -> ResponseReturnValue:
    identifier = str(definition_id or "").strip()
    if not identifier:
        return jsonify({"error": "Ungültige ID."}), 400
    definition = get_search_definition(identifier)
    if not definition:
        return jsonify({"error": "Suche wurde nicht gefunden."}), 404

    payload = request.get_json(silent=True) or {}
    overrides = payload if isinstance(payload, dict) else {}
    try:
        definition_override = store_search_definition({**definition, **overrides, "id": identifier})
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400

    job_id = str(uuid.uuid4())
    keywords = list(definition_override.get("keywords", []))
    school_ids = list(definition_override.get("school_ids", []))
    prompt_text = str(definition_override.get("prompt") or "").strip()
    model_value = str(definition_override.get("model") or DEFAULT_SEARCH_MODEL).strip()
    temperature_value = float(definition_override.get("temperature", DEFAULT_SEARCH_TEMPERATURE))
    max_results_value = int(definition_override.get("max_results", DEFAULT_SEARCH_MAX_RESULTS))
    try:
        concurrency_value = int(definition_override.get("concurrency", DEFAULT_SEARCH_CONCURRENCY))
    except (TypeError, ValueError):
        concurrency_value = DEFAULT_SEARCH_CONCURRENCY
    if concurrency_value < 1:
        concurrency_value = 1

    job = SearchJob(
        id=job_id,
        definition_id=identifier,
        definition_name=str(definition_override.get("name") or identifier),
        keywords=keywords,
        school_ids=school_ids,
        prompt=prompt_text,
        model=model_value,
        temperature=temperature_value,
        max_results=max_results_value,
        concurrency=concurrency_value,
    )
    register_search_job(job)
    persist_search_job_state(job)

    thread = threading.Thread(target=run_search_job, args=(job,), daemon=True)
    thread.start()

    return jsonify({"job_id": job_id, "definition": definition_override})


@app.get("/searches/jobs/<job_id>")
def search_job_status(job_id: str) -> ResponseReturnValue:
    identifier = str(job_id or "").strip()
    if not identifier:
        return jsonify({"error": "Ungültige Job-ID."}), 400
    job = get_search_job(identifier)
    if job:
        return jsonify({"job": job.as_dict()})
    run = get_search_run(identifier)
    if run:
        return jsonify({"job": run})
    return jsonify({"error": "Job nicht gefunden."}), 404


@app.post("/searches/jobs/<job_id>/cancel")
def cancel_search_job(job_id: str) -> ResponseReturnValue:
    identifier = str(job_id or "").strip()
    if not identifier:
        return jsonify({"error": "Ungültige Job-ID."}), 400
    job = get_search_job(identifier)
    if not job:
        return jsonify({"error": "Job nicht gefunden."}), 404
    job.request_cancel()
    return jsonify({"status": "cancelling"})


@app.get("/searches/runs/<run_id>")
def search_run_info(run_id: str) -> ResponseReturnValue:
    identifier = str(run_id or "").strip()
    if not identifier:
        return jsonify({"error": "Ungültige Run-ID."}), 400
    run = get_search_run(identifier)
    if not run:
        return jsonify({"error": "Run nicht gefunden."}), 404
    results = load_search_results(identifier)
    logs = load_search_logs(identifier)
    return jsonify({"run": run, "result_count": len(results)})


@app.get("/searches/runs/<run_id>/results")
def search_run_results(run_id: str) -> ResponseReturnValue:
    identifier = str(run_id or "").strip()
    if not identifier:
        return jsonify({"error": "Ungültige Run-ID."}), 400
    run = get_search_run(identifier)
    if not run:
        return jsonify({"error": "Run nicht gefunden."}), 404
    results = load_search_results(identifier)

    school_filter = request.args.get("school_id", "").strip()
    keyword_filter = request.args.get("keyword", "").strip()
    dimension_filter = request.args.get("dimension", "").strip()
    min_value_raw = request.args.get("min_value")
    min_confidence_raw = request.args.get("min_confidence")

    try:
        min_value = float(min_value_raw) if min_value_raw is not None else None
    except (TypeError, ValueError):
        min_value = None
    try:
        min_confidence = float(min_confidence_raw) if min_confidence_raw is not None else None
    except (TypeError, ValueError):
        min_confidence = None

    filtered: List[Dict[str, Any]] = []
    for entry in results:
        if school_filter and str(entry.get("school_id")) != school_filter:
            continue
        if keyword_filter and str(entry.get("keyword")) != keyword_filter:
            continue
        if dimension_filter:
            dimensions = entry.get("dimensions") if isinstance(entry.get("dimensions"), list) else []
            matches_dimension = False
            for dim in dimensions:
                if not isinstance(dim, dict):
                    continue
                if str(dim.get("dimension")) != dimension_filter:
                    continue
                value = dim.get("value")
                confidence = dim.get("confidence")
                if min_value is not None and (value is None or float(value) < min_value):
                    continue
                if min_confidence is not None and (confidence is None or float(confidence) < min_confidence):
                    continue
                matches_dimension = True
                break
            if not matches_dimension:
                continue
        filtered.append(entry)

    dimension_summary: Dict[str, Dict[str, float]] = {}
    school_summary: Dict[str, Dict[str, Any]] = {}
    keyword_summary: Dict[str, Dict[str, Any]] = {}
    snapshot_raw = run.get("school_snapshot") if isinstance(run, dict) else {}
    school_snapshot: Dict[str, Dict[str, Any]] = (
        snapshot_raw if isinstance(snapshot_raw, dict) else {}
    )
    aggregated_schools: Dict[str, Dict[str, Any]] = {}

    for entry in filtered:
        school_id = str(entry.get("school_id") or "")
        school_summary.setdefault(
            school_id,
            {
                "school_id": school_id,
                "school_name": entry.get("school_name"),
                "school_city": entry.get("school_city"),
                "count": 0,
            },
        )["count"] += 1
        keyword = str(entry.get("keyword") or "")
        keyword_summary.setdefault(
            keyword,
            {
                "keyword": keyword,
                "count": 0,
            },
        )["count"] += 1
        aggregated = aggregated_schools.setdefault(
            school_id,
            {
                "school_id": school_id,
                "school_name": entry.get("school_name")
                or school_snapshot.get(school_id, {}).get("name"),
                "school_city": entry.get("school_city")
                or school_snapshot.get(school_id, {}).get("city"),
                "school_homepage": entry.get("school_homepage")
                or school_snapshot.get(school_id, {}).get("homepage"),
                "total_hits": 0,
                "dimensions": {},
                "result_ids": [],
            },
        )
        aggregated["total_hits"] += 1
        if entry.get("id") is not None:
            aggregated["result_ids"].append(entry.get("id"))
        for dim in entry.get("dimensions", []):
            if not isinstance(dim, dict):
                continue
            label = str(dim.get("dimension"))
            value = float(dim.get("value", 0))
            confidence = float(dim.get("confidence", 0.0))
            stats = dimension_summary.setdefault(
                label,
                {
                    "dimension": label,
                    "count": 0,
                    "value_sum": 0.0,
                    "confidence_sum": 0.0,
                },
            )
            stats["count"] += 1
            stats["value_sum"] += value
            stats["confidence_sum"] += confidence
            agg_dim = aggregated["dimensions"].setdefault(
                label,
                {
                    "dimension": label,
                    "value_sum": 0.0,
                    "confidence_sum": 0.0,
                    "count": 0,
                },
            )
            agg_dim["value_sum"] += value
            agg_dim["confidence_sum"] += confidence
            agg_dim["count"] += 1

    for stats in dimension_summary.values():
        count = max(stats.get("count", 0), 1)
        stats["average_value"] = stats["value_sum"] / count
        stats["average_confidence"] = stats["confidence_sum"] / count
        stats.pop("value_sum", None)
        stats.pop("confidence_sum", None)

    aggregated_list: List[Dict[str, Any]] = []
    for school_id, summary in aggregated_schools.items():
        dimensions_compiled: List[Dict[str, Any]] = []
        for label, dim_stats in summary["dimensions"].items():
            count = max(dim_stats.get("count", 0), 1)
            dimensions_compiled.append(
                {
                    "dimension": label,
                    "average_value": dim_stats.get("value_sum", 0.0) / count,
                    "average_confidence": dim_stats.get("confidence_sum", 0.0) / count,
                    "count": dim_stats.get("count", 0),
                }
            )
        dimensions_compiled.sort(key=lambda item: item["dimension"])
        aggregated_list.append(
            {
                "school_id": school_id,
                "school_name": summary.get("school_name") or school_snapshot.get(school_id, {}).get("name"),
                "school_city": summary.get("school_city") or school_snapshot.get(school_id, {}).get("city"),
                "school_homepage": summary.get("school_homepage")
                or school_snapshot.get(school_id, {}).get("homepage"),
                "total_hits": summary.get("total_hits", 0),
                "result_ids": summary.get("result_ids", []),
                "dimensions": dimensions_compiled,
            }
        )
    aggregated_list.sort(key=lambda item: (item["school_name"] or "", item["school_id"]))

    dimension_matrix: List[Dict[str, Any]] = []
    for dimension_label, stats in dimension_summary.items():
        buckets: Dict[int, Dict[str, Any]] = {}
        for entry in filtered:
            for dim in entry.get("dimensions", []):
                if not isinstance(dim, dict):
                    continue
                if str(dim.get("dimension")) != dimension_label:
                    continue
                try:
                    value = int(dim.get("value", 0))
                except (TypeError, ValueError):
                    value = 0
                bucket = buckets.setdefault(
                    value,
                    {
                        "value": value,
                        "count": 0,
                        "entries": [],
                    },
                )
                bucket["count"] += 1
                bucket["entries"].append(
                    {
                        "result_id": entry.get("id"),
                        "school_id": entry.get("school_id"),
                        "school_name": entry.get("school_name"),
                        "school_city": entry.get("school_city"),
                        "target_url": entry.get("target_url"),
                        "keyword": entry.get("keyword"),
                        "confidence": dim.get("confidence"),
                    }
                )
        total = sum(bucket["count"] for bucket in buckets.values()) or 1
        values_sorted = sorted(buckets.values(), key=lambda item: item["value"], reverse=True)
        for bucket in values_sorted:
            bucket["percentage"] = bucket["count"] / total
        dimension_matrix.append(
            {
                "dimension": dimension_label,
                "total": total,
                "values": values_sorted,
            }
        )

    total_results = len(results)
    filtered_results = len(filtered)
    message: Optional[str] = None
    if total_results == 0:
        message = "Keine Ergebnisse für diese Suche gefunden."
    elif filtered_results == 0:
        message = "Keine Ergebnisse entsprechen den aktuellen Filtern."

    selected_school_ids = {
        str(item)
        for item in (run.get("school_ids") if isinstance(run, dict) else [])
        if str(item)
    }
    if not selected_school_ids and school_snapshot:
        selected_school_ids = set(school_snapshot.keys())
    hit_school_ids = {item["school_id"] for item in aggregated_list if item.get("school_id")}
    missing_school_ids = sorted(selected_school_ids - hit_school_ids)
    schools_without_hits: List[Dict[str, Any]] = []
    for school_id in missing_school_ids:
        info = school_snapshot.get(school_id, {})
        schools_without_hits.append(
            {
                "school_id": school_id,
                "school_name": info.get("name"),
                "school_city": info.get("city"),
                "school_homepage": info.get("homepage"),
            }
        )

    return jsonify(
        {
            "run": run,
            "total_results": total_results,
            "filtered_results": filtered_results,
            "results": filtered,
            "dimension_summary": list(dimension_summary.values()),
            "dimension_matrix": dimension_matrix,
            "school_summary": list(school_summary.values()),
            "keyword_summary": list(keyword_summary.values()),
            "aggregated_schools": aggregated_list,
            "schools_without_hits": schools_without_hits,
            "total_selected_schools": len(selected_school_ids),
            "message": message,
            "logs": logs,
        }
    )


@app.patch("/searches/results/<int:result_id>")
def update_search_result_route(result_id: int) -> ResponseReturnValue:
    existing = get_search_result_entry(result_id)
    if not existing:
        return jsonify({"error": "Suchergebnis nicht gefunden."}), 404
    payload = request.get_json(silent=True) or {}
    if not isinstance(payload, dict):
        return jsonify({"error": "Ungültige Anfrage."}), 400
    dimensions_payload = payload.get("dimensions")
    if dimensions_payload is not None:
        if not isinstance(dimensions_payload, list):
            return jsonify({"error": "Die Dimensionen müssen als Liste übermittelt werden."}), 400
        cleaned_dimensions: List[Dict[str, Any]] = []
        for entry in dimensions_payload:
            if not isinstance(entry, dict):
                continue
            label = str(entry.get("dimension") or "").strip()
            if not label:
                continue
            try:
                value = int(entry.get("value", 0))
            except (TypeError, ValueError):
                try:
                    value = int(float(entry.get("value", 0)))
                except (TypeError, ValueError):
                    value = 0
            try:
                confidence = float(entry.get("confidence", 0.0))
            except (TypeError, ValueError):
                confidence = 0.0
            cleaned_dimensions.append(
                {
                    "dimension": label,
                    "value": max(0, min(value, 5)),
                    "confidence": max(0.0, min(confidence, 1.0)),
                }
            )
        existing["dimensions"] = cleaned_dimensions
    if "keyword" in payload:
        existing["keyword"] = str(payload.get("keyword") or existing.get("keyword") or "").strip()
    if "target_url" in payload:
        existing["target_url"] = str(payload.get("target_url") or existing.get("target_url") or "").strip()
    if "manual_note" in payload:
        existing["manual_note"] = str(payload.get("manual_note") or "").strip()
    with search_result_lock:
        storage.update_search_result(result_id, existing)
    updated = get_search_result_entry(result_id)
    return jsonify({"result": updated})


@app.delete("/searches/results/<int:result_id>")
def delete_search_result_route(result_id: int) -> ResponseReturnValue:
    with search_result_lock:
        storage.delete_search_results([result_id])
    return jsonify({"status": "deleted"})


@app.route("/stammdaten", methods=["GET", "POST"])
def stammdaten():
    message: Optional[str] = None
    message_category: Optional[str] = None

    if request.method == "POST":
        uploaded = request.files.get("excel_file")
        if uploaded is None or not uploaded.filename:
            message = "Bitte wählen Sie eine Excel-Datei aus."
            message_category = "danger"
        else:
            try:
                records = parse_stammdaten_excel(uploaded)
            except ValueError as exc:
                message = str(exc)
                message_category = "danger"
            else:
                save_stammdaten(records)
                message = f"{len(records)} Stammdatensätze wurden übernommen."
                message_category = "success"

    records = load_stammdaten()
    return render_template(
        "stammdaten.html",
        records=records,
        message=message,
        message_category=message_category,
        stammdaten_fields=STAMMDATEN_FIELDS,
        stammdaten_primary_fields=STAMMDATEN_PRIMARY_FIELDS,
        stammdaten_boolean_fields=STAMMDATEN_BOOLEAN_FIELDS,
        storage_file=storage.path.name,
        active_page="stammdaten",
    )


@app.route("/data-quality")
def data_quality() -> ResponseReturnValue:
    dataset = build_data_quality_dataset()
    dq_settings = get_data_quality_settings()
    google_credentials = get_google_search_credentials()
    google_configured = bool(google_credentials.get("api_key") and google_credentials.get("cx"))
    return render_template(
        "data_quality.html",
        records=dataset,
        data_quality_settings=dq_settings,
        google_credentials=google_credentials,
        openai_present=has_api_key(),
        google_configured=google_configured,
        status_labels={
            QUALITY_STATUS_PENDING: quality_status_label(QUALITY_STATUS_PENDING),
            QUALITY_STATUS_OK: quality_status_label(QUALITY_STATUS_OK),
            QUALITY_STATUS_UNSURE: quality_status_label(QUALITY_STATUS_UNSURE),
            QUALITY_STATUS_INVALID: quality_status_label(QUALITY_STATUS_INVALID),
        },
        site_scope_labels=SITE_SCOPE_LABELS,
        site_scope_classes=SITE_SCOPE_CLASSES,
        site_scope_unknown=SITE_SCOPE_UNKNOWN,
        active_page="data_quality",
    )


@app.route("/debug")
def debug_page() -> ResponseReturnValue:
    logs = list_error_log_entries()
    return render_template(
        "debug.html",
        active_page="debug",
        initial_logs=logs,
    )


@app.get("/debug/logs")
def debug_logs() -> ResponseReturnValue:
    limit = request.args.get("limit", default=200, type=int)
    if limit is None:
        limit = 200
    limit = max(1, min(limit, 2000))
    return jsonify({"logs": list_error_log_entries(limit)})


@app.get("/debug/export")
def debug_export() -> Response:
    logs = list_error_log_entries(limit=None)
    payload = {
        "generated_at": datetime.utcnow().isoformat() + "Z",
        "logs": logs,
    }
    raw = json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8")
    filename = f"crawler-debug-{datetime.utcnow().strftime('%Y%m%d-%H%M%S')}.json"
    response = Response(raw, mimetype="application/json")
    response.headers["Content-Disposition"] = f"attachment; filename={filename}"
    return response


@app.route("/data-quality/data")
def data_quality_data() -> ResponseReturnValue:
    return jsonify({"records": build_data_quality_dataset()})


@app.route("/data-quality/details/<school_id>")
def data_quality_details(school_id: str) -> ResponseReturnValue:
    results = load_data_quality_results()
    record = results.get(str(school_id)) if isinstance(results, dict) else None
    if not isinstance(record, dict):
        return jsonify({"error": "Keine Daten vorhanden."}), 404
    return jsonify({"school_id": school_id, "quality": record, "history": record.get("history", [])})


@app.route("/data-quality/start", methods=["POST"])
def start_data_quality_job() -> ResponseReturnValue:
    payload = request.get_json(silent=True) or {}
    school_ids_raw = payload.get("school_ids")
    dataset = build_data_quality_dataset()
    if isinstance(school_ids_raw, list) and school_ids_raw:
        school_ids = [str(item).strip() for item in school_ids_raw if str(item).strip()]
    else:
        school_ids = [record["schul_id"] for record in dataset]
    school_ids = [sid for sid in school_ids if sid]
    if not school_ids:
        return jsonify({"error": "Keine Schulen ausgewählt."}), 400
    mode_raw = payload.get("mode")
    if isinstance(mode_raw, str):
        candidate = mode_raw.strip().lower()
        if candidate in {"google", "google_search", "google-only"}:
            mode = "google_search"
        elif candidate in {"site_scope", "site-scope", "scope", "auftritt", "auftrittstyp"}:
            mode = "site_scope"
        else:
            mode = "full"
    else:
        mode = "full"
    settings = get_data_quality_settings()
    try:
        batch_size = int(settings.get("batch_size", len(school_ids)))
    except (TypeError, ValueError):
        batch_size = len(school_ids)
    if batch_size > 0 and mode != "site_scope":
        school_ids = school_ids[:batch_size]

    job_id = str(uuid.uuid4())
    job = DataQualityJob(id=job_id, school_ids=school_ids, total=len(school_ids), mode=mode)
    data_quality_jobs[job_id] = job
    thread = threading.Thread(target=run_data_quality_job, args=(job,), daemon=True)
    thread.start()
    return jsonify({"job_id": job_id})


@app.route("/data-quality/status/<job_id>")
def data_quality_status(job_id: str) -> ResponseReturnValue:
    job = data_quality_jobs.get(job_id)
    if not job:
        return jsonify({"error": "Unbekannte Job-ID"}), 404
    return jsonify(job.as_dict())


@app.route("/data-quality/cancel/<job_id>", methods=["POST"])
def cancel_data_quality_job(job_id: str) -> ResponseReturnValue:
    job = data_quality_jobs.get(job_id)
    if not job:
        return jsonify({"error": "Unbekannte Job-ID"}), 404
    job.request_cancel()
    return jsonify({"status": "cancelling"})


@app.route("/data-quality/export")
def export_data_quality() -> ResponseReturnValue:
    data = load_data_quality_results()
    response = app.response_class(
        json.dumps(data, ensure_ascii=False, indent=2),
        mimetype="application/json",
    )
    response.headers["Content-Disposition"] = "attachment; filename=data_quality_results.json"
    return response


@app.route("/data-quality/record/<school_id>", methods=["POST"])
def manage_data_quality_record(school_id: str) -> ResponseReturnValue:
    payload = request.get_json(silent=True) or {}
    action = (payload.get("action") or "").strip().lower()
    if not action:
        return jsonify({"error": "Keine Aktion übermittelt."}), 400

    school_id = str(school_id)
    dry_run = bool(get_data_quality_settings().get("dry_run", False))
    results = load_data_quality_results()
    record = results.get(school_id) if isinstance(results, dict) else None
    if not isinstance(record, dict):
        record = {"status": QUALITY_STATUS_PENDING}

    now_iso = datetime.utcnow().isoformat() + "Z"

    if action == "accept_suggestion":
        suggestion = str(record.get("suggested_url") or "").strip()
        if not suggestion:
            return jsonify({"error": "Kein URL-Vorschlag vorhanden."}), 400
        if _apply_data_quality_url(
            school_id,
            suggestion,
            dry_run=dry_run,
            history_stage="manual",
            history_status="accepted",
            history_detail="URL-Vorschlag übernommen",
        ):
            return jsonify({"status": "ok", "records": build_data_quality_dataset()})
        return jsonify({"error": "URL konnte nicht aktualisiert werden."}), 500

    if action == "update_url":
        new_url = str(payload.get("url") or "").strip()
        if not new_url:
            return jsonify({"error": "Bitte geben Sie eine gültige URL an."}), 400
        if _apply_data_quality_url(
            school_id,
            new_url,
            dry_run=dry_run,
            manual_note_suffix="manuell",
            history_stage="manual",
            history_status="updated",
            history_detail="URL manuell angepasst",
            history_extra={"url": new_url},
        ):
            return jsonify({"status": "ok", "records": build_data_quality_dataset()})
        return jsonify({"error": "URL konnte nicht aktualisiert werden."}), 500

    if action == "update_site_scope":
        scope_value = normalise_site_scope(payload.get("site_scope"))
        reason_value = str(payload.get("reason") or "").strip()
        update_data_quality_record(
            school_id,
            {
                "site_scope": scope_value,
                "site_scope_reason": reason_value,
                "site_scope_manual": True,
                "site_scope_updated_at": now_iso,
                "site_scope_source": "manual",
                "_force_site_scope": True,
            },
        )
        append_data_quality_history(
            school_id,
            {
                "stage": "manual",
                "status": "site_scope",
                "detail": f"Auftrittstyp gesetzt auf {site_scope_label(scope_value)}.",
                "reason": reason_value,
            },
        )
        return jsonify({"status": "ok", "records": build_data_quality_dataset()})

    if action == "mark_ok":
        update_data_quality_record(
            school_id,
            {
                "status": QUALITY_STATUS_OK,
                "last_checked": now_iso,
                "confidence": 1.0,
                "manual_note": "Manuell als korrekt markiert",
                "correction_applied": False,
                "correction_applied_at": None,
            },
        )
        append_data_quality_history(
            school_id,
            {
                "stage": "manual",
                "status": "ok",
                "detail": "Manuell als korrekt markiert.",
            },
        )
        return jsonify({"status": "ok", "records": build_data_quality_dataset()})

    if action == "ignore":
        update_data_quality_record(
            school_id,
            {
                "ignored": True,
                "last_checked": now_iso,
                "manual_note": "Als ignoriert markiert",
                "correction_applied": False,
                "correction_applied_at": None,
            },
        )
        append_data_quality_history(
            school_id,
            {
                "stage": "manual",
                "status": "ignored",
                "detail": "Als ignoriert markiert.",
            },
        )
        return jsonify({"status": "ok", "records": build_data_quality_dataset()})

    if action == "recheck":
        job_id = str(uuid.uuid4())
        job = DataQualityJob(id=job_id, school_ids=[school_id], total=1)
        data_quality_jobs[job_id] = job
        thread = threading.Thread(target=run_data_quality_job, args=(job,), daemon=True)
        thread.start()
        return jsonify({"status": "started", "job_id": job_id})

    return jsonify({"error": f"Unbekannte Aktion: {action}"}), 400


@app.route("/data-quality/bulk-accept", methods=["POST"])
def bulk_accept_data_quality() -> ResponseReturnValue:
    payload = request.get_json(silent=True) or {}
    requested_ids_raw = payload.get("school_ids")

    dataset = build_data_quality_dataset()
    dataset_map = {str(record.get("schul_id")): record for record in dataset if record.get("schul_id")}

    if isinstance(requested_ids_raw, list) and requested_ids_raw:
        target_ids = []
        seen: Set[str] = set()
        for raw_id in requested_ids_raw:
            sid = str(raw_id).strip()
            if not sid or sid in seen or sid not in dataset_map:
                continue
            seen.add(sid)
            target_ids.append(sid)
    else:
        target_ids = [
            record_id
            for record_id, record in dataset_map.items()
            if str(record.get("status")).lower() == QUALITY_STATUS_INVALID
        ]

    if not target_ids:
        return jsonify({"error": "Keine als falsch markierten Vorschläge gefunden."}), 400

    dry_run = bool(get_data_quality_settings().get("dry_run", False))
    applied: List[str] = []
    skipped: List[Dict[str, object]] = []
    failed: List[Dict[str, object]] = []

    for school_id in target_ids:
        record = dataset_map.get(school_id) or {}
        quality = record.get("quality") if isinstance(record.get("quality"), dict) else {}
        suggestion = str(quality.get("suggested_url") or "").strip()
        if not suggestion:
            skipped.append({"schul_id": school_id, "reason": "Kein Vorschlag vorhanden"})
            continue
        success = _apply_data_quality_url(
            school_id,
            suggestion,
            dry_run=dry_run,
            manual_note_suffix="automatisch",
            history_stage="bulk",
            history_status="accepted",
            history_detail="URL-Vorschlag automatisch übernommen",
            history_extra={"source": "bulk"},
        )
        if success:
            applied.append(school_id)
        else:
            failed.append({"schul_id": school_id, "reason": "URL konnte nicht aktualisiert werden"})

    summary = {
        "requested": len(target_ids),
        "applied": len(applied),
        "skipped": len(skipped),
        "failed": len(failed),
        "dry_run": dry_run,
    }

    response: Dict[str, object] = {
        "status": "ok",
        "summary": summary,
        "applied": applied,
        "skipped": skipped,
        "failed": failed,
        "records": build_data_quality_dataset(),
    }
    if dry_run:
        response["message"] = "Dry-Run aktiv: Stammdaten wurden nicht geändert."
    if failed and not applied:
        response["status"] = "error"
        response["error"] = "Keine URL konnte aktualisiert werden."
        return jsonify(response), 500

    return jsonify(response)


@app.route("/settings", methods=["GET", "POST"])
def settings():
    message: Optional[str] = None
    message_category: Optional[str] = None
    key_present = has_api_key()
    planner_key_present = has_keyword_planner_key()
    synonym_defaults = get_synonym_defaults()
    keyword_finder_defaults = get_keyword_finder_defaults()
    crawl_defaults = get_crawl_defaults()
    google_credentials = get_google_search_credentials()
    data_quality_settings = get_data_quality_settings()
    developer_mode_enabled = is_developer_mode_enabled()

    if request.method == "POST":
        form_id = request.form.get("form_id", "api")
        if form_id == "synonym-defaults":
            values = {
                "prompt": request.form.get("synonym_prompt", ""),
                "max_results": request.form.get("synonym_max_results"),
                "temperature": request.form.get("synonym_temperature"),
                "model": request.form.get("synonym_model", ""),
            }
            update_synonym_defaults(values)
            synonym_defaults = get_synonym_defaults()
            message = "Die Standardwerte für die Synonymsuche wurden gespeichert."
            message_category = "success"
        elif form_id == "crawl-defaults":
            values = {
                "max_pages": request.form.get("crawl_max_pages"),
                "concurrency": request.form.get("crawl_concurrency"),
            }
            crawl_defaults = update_crawl_defaults(values)
            message = "Die Standardwerte für den Crawl wurden gespeichert."
            message_category = "success"
        elif form_id == "keyword-finder-defaults":
            values = {
                "prompt": request.form.get("finder_prompt", ""),
                "max_results": request.form.get("finder_max_results"),
                "temperature": request.form.get("finder_temperature"),
                "model": request.form.get("finder_model", ""),
            }
            keyword_finder_defaults = update_keyword_finder_defaults(values)
            message = "Die Standardwerte für den Keyword-Finder wurden gespeichert."
            message_category = "success"
        elif form_id == "search-defaults":
            values = {
                "prompt": request.form.get("search_prompt", ""),
                "max_results": request.form.get("search_max_results"),
                "temperature": request.form.get("search_temperature"),
                "model": request.form.get("search_model", ""),
                "concurrency": request.form.get("search_concurrency"),
            }
            search_defaults = update_search_defaults(values)
            message = "Die Bewertungs-Standardwerte wurden gespeichert."
            message_category = "success"
        elif form_id == "google-search":
            action = request.form.get("action", "save")
            if action == "remove":
                clear_google_search_credentials()
                google_credentials = get_google_search_credentials()
                message = "Die Google-Search-Zugangsdaten wurden entfernt."
                message_category = "success"
            elif action == "test":
                candidate_key = request.form.get("google_search_api_key", "").strip()
                candidate_cx = request.form.get("google_search_cx", "").strip()
                api_key = candidate_key or google_credentials.get("api_key")
                cx = candidate_cx or google_credentials.get("cx")
                if not api_key or not cx:
                    message = "Bitte geben Sie API-Key und Search-Engine-ID an oder speichern Sie sie zuerst."
                    message_category = "danger"
                else:
                    results = perform_google_search(
                        "Beispielschule Stuttgart",
                        api_key,
                        cx,
                        max_results=1,
                    )
                    if results:
                        if candidate_key or candidate_cx:
                            google_credentials = update_google_search_credentials(api_key, cx)
                        message = "Die Google-Suche war erfolgreich."
                        message_category = "success"
                    else:
                        message = "Testsuche fehlgeschlagen. Bitte Zugangsdaten prüfen."
                        message_category = "danger"
            else:
                api_key = request.form.get("google_search_api_key", "").strip()
                cx = request.form.get("google_search_cx", "").strip()
                google_credentials = update_google_search_credentials(api_key, cx)
                message = "Die Google-Search-Zugangsdaten wurden gespeichert."
                message_category = "success"
        elif form_id == "data-quality-defaults":
            values = {
                "model": request.form.get("dq_model", ""),
                "temperature": request.form.get("dq_temperature"),
                "confidence_threshold": request.form.get("dq_confidence_threshold"),
                "max_search_results": request.form.get("dq_max_search_results"),
                "batch_size": request.form.get("dq_batch_size"),
                "dry_run": request.form.get("dq_dry_run"),
                "search_language": request.form.get("dq_search_language", ""),
                "search_region": request.form.get("dq_search_region", ""),
                "allowed_domains": request.form.get("dq_allowed_domains", ""),
                "blocked_domains": request.form.get("dq_blocked_domains", ""),
            }
            data_quality_settings = update_data_quality_settings(values)
            message = "Die Einstellungen für die Datenqualität wurden gespeichert."
            message_category = "success"
        elif form_id == "google-key":
            action = request.form.get("action", "save")
            if action == "remove":
                if planner_key_present:
                    clear_keyword_planner_key()
                    planner_key_present = False
                    message = "Der Keyword-Planner-Schlüssel wurde entfernt."
                    message_category = "success"
                else:
                    message = "Es ist kein Keyword-Planner-Schlüssel hinterlegt."
                    message_category = "info"
            else:
                key_value = request.form.get("google_api_key", "").strip()
                if not key_value:
                    message = "Bitte geben Sie einen gültigen Keyword-Planner-Schlüssel ein."
                    message_category = "danger"
                else:
                    set_keyword_planner_key(key_value)
                    planner_key_present = True
                    message = "Der Keyword-Planner-Schlüssel wurde gespeichert."
                    message_category = "success"
        elif form_id == "developer-mode":
            enabled = request.form.get("developer_mode") in {"on", "1", "true", "True"}
            developer_mode_enabled = set_developer_mode_enabled(enabled)
            if developer_mode_enabled:
                message = "Der Entwicklungsmodus ist nun aktiviert."
                message_category = "success"
            else:
                message = "Der Entwicklungsmodus wurde deaktiviert."
                message_category = "info"
        else:
            action = request.form.get("action", "save")
            if action == "remove":
                if key_present:
                    clear_api_key()
                    key_present = False
                    message = "Der API-Schlüssel wurde entfernt."
                    message_category = "success"
                else:
                    message = "Es ist kein API-Schlüssel hinterlegt."
                    message_category = "info"
            elif action == "test":
                candidate = request.form.get("api_key", "").strip()
                key_to_test = candidate or get_api_key()
                if not key_to_test:
                    message = "Bitte geben Sie einen API-Schlüssel ein oder speichern Sie ihn zuerst."
                    message_category = "danger"
                else:
                    try:
                        test_openai_api_key(key_to_test)
                    except OpenAIIntegrationError as exc:
                        message = str(exc)
                        message_category = "danger"
                    else:
                        if candidate:
                            set_api_key(candidate)
                            message = "Der API-Schlüssel wurde gespeichert und erfolgreich geprüft."
                        else:
                            message = "Der hinterlegte API-Schlüssel wurde erfolgreich geprüft."
                        message_category = "success"
                    key_present = has_api_key()
            else:
                key_value = request.form.get("api_key", "").strip()
                if not key_value:
                    message = "Bitte geben Sie einen gültigen API-Schlüssel ein."
                    message_category = "danger"
                else:
                    set_api_key(key_value)
                    key_present = True
                    message = "Der API-Schlüssel wurde gespeichert."
                    message_category = "success"

    key_present = has_api_key()
    planner_key_present = has_keyword_planner_key()
    keyword_finder_defaults = get_keyword_finder_defaults()
    crawl_defaults = get_crawl_defaults()
    search_defaults = get_search_defaults()
    google_credentials = get_google_search_credentials()
    data_quality_settings = get_data_quality_settings()
    developer_mode_enabled = is_developer_mode_enabled()

    return render_template(
        "settings.html",
        message=message,
        message_category=message_category,
        key_present=key_present,
        keyword_planner_key_present=planner_key_present,
        synonym_defaults=synonym_defaults,
        available_models=AVAILABLE_OPENAI_MODELS,
        crawl_defaults=crawl_defaults,
        max_concurrency=MAX_CONCURRENCY,
        max_max_pages=MAX_MAX_PAGES,
        keyword_finder_defaults=keyword_finder_defaults,
        keyword_finder_max_results=DEFAULT_KEYWORD_FINDER_MAX_RESULTS,
        search_defaults=search_defaults,
        google_credentials=google_credentials,
        data_quality_settings=data_quality_settings,
        export_sections=EXPORT_SECTION_ORDER,
        export_definitions=EXPORT_SECTION_DEFINITIONS,
        developer_mode_enabled=developer_mode_enabled,
        active_page="settings",
    )


@app.post("/settings/export")
def export_data_bundle() -> ResponseReturnValue:
    request_data = request.get_json(silent=True) or {}
    sections_value = request_data.get("sections", [])
    if not isinstance(sections_value, list):
        sections_value = []
    selected_sections = {str(value) for value in sections_value}
    payload, summary, contains_sensitive = build_export_payload(selected_sections)
    raw = json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8")
    encoded = base64.b64encode(raw).decode("ascii")
    filename = f"crawler-export-{datetime.utcnow().strftime('%Y%m%d-%H%M%S')}.json"
    record_audit_event(
        "export",
        {
            "sections": summary.get("requested_sections", []),
            "bytes": len(raw),
            "contains_sensitive": contains_sensitive,
        },
    )
    return jsonify(
        {
            "status": "ok",
            "filename": filename,
            "filesize": len(raw),
            "summary": summary,
            "download": encoded,
            "contains_sensitive": contains_sensitive,
        }
    )


@app.post("/settings/import/preview")
def import_preview() -> ResponseReturnValue:
    if "file" not in request.files:
        return jsonify({"status": "error", "message": "Bitte wählen Sie eine Exportdatei aus."}), 400
    file_storage = request.files["file"]
    try:
        payload = json.load(file_storage)
    except json.JSONDecodeError:
        return jsonify({"status": "error", "message": "Die Datei konnte nicht gelesen werden."}), 400
    try:
        summary, contains_sensitive = analyze_import_payload(payload)
    except ValueError as exc:
        return jsonify({"status": "error", "message": str(exc)}), 400

    session_id = create_import_session(payload, summary, contains_sensitive)
    record_audit_event(
        "import-preview",
        {
            "session": session_id,
            "sections": summary.get("available_sections", []),
            "contains_sensitive": contains_sensitive,
        },
    )
    return jsonify(
        {
            "status": "ok",
            "session_id": session_id,
            "summary": summary,
            "contains_sensitive": contains_sensitive,
        }
    )


@app.post("/settings/import/apply")
def import_apply_route() -> ResponseReturnValue:
    request_data = request.get_json(silent=True) or {}
    session_id = str(request_data.get("session_id") or "").strip()
    if not session_id:
        return jsonify({"status": "error", "message": "Die Vorschau ist abgelaufen. Bitte erneut laden."}), 400
    session = get_import_session(session_id)
    if not session:
        return jsonify({"status": "error", "message": "Die Vorschau ist abgelaufen. Bitte erneut laden."}), 410

    actions = request_data.get("actions", {})
    if not isinstance(actions, dict):
        actions = {}
    dry_run = bool(request_data.get("dry_run", False))
    confirm_sensitive = bool(request_data.get("confirm_sensitive", False))

    contains_sensitive = bool(session.get("contains_sensitive"))
    if contains_sensitive and not confirm_sensitive:
        return (
            jsonify(
                {
                    "status": "error",
                    "message": "Bitte bestätigen Sie, dass Sie sensible Daten importieren möchten.",
                }
            ),
            400,
        )

    payload = session.get("payload")
    try:
        result = apply_import(payload, {str(k): str(v) for k, v in actions.items()}, dry_run)
    except ValueError as exc:
        return jsonify({"status": "error", "message": str(exc)}), 400

    record_audit_event(
        "import-apply",
        {
            "session": session_id,
            "dry_run": dry_run,
            "actions": {k: actions.get(k, "ignore") for k in EXPORT_SECTION_DEFINITIONS.keys()},
            "contains_sensitive": contains_sensitive,
        },
    )

    if not dry_run:
        consume_import_session(session_id)

    return jsonify(
        {
            "status": "ok",
            "dry_run": dry_run,
            "result": result,
            "summary": session.get("summary"),
        }
    )


@app.route("/status/<job_id>")
def job_status(job_id: str):
    job = jobs.get(job_id)
    if not job:
        return jsonify({"error": "Unbekannte Job-ID"}), 404
    return jsonify(job.as_dict())


@app.route("/cancel/<job_id>", methods=["POST"])
def cancel_job(job_id: str):
    job = jobs.get(job_id)
    if not job:
        return jsonify({"error": "Unbekannte Job-ID"}), 404
    job.request_cancel()
    return jsonify({"status": job.status})


@app.route("/save-search/<job_id>", methods=["POST"])
def save_search(job_id: str):
    job = jobs.get(job_id)
    if not job:
        return jsonify({"error": "Unbekannte Job-ID"}), 404
    with job._lock:
        if not job.completed:
            return jsonify({"error": "Der Crawl läuft noch. Bitte warten Sie, bis er beendet ist."}), 400
        entry = {
            "saved_at": datetime.utcnow().isoformat() + "Z",
            "job_id": job.id,
            "start_urls": list(job.start_urls),
            "keywords": list(job.keywords),
            "original_keywords": list(job.original_keywords),
            "expanded_keywords": list(job.expanded_keywords),
            "keyword_groups": [
                {
                    "original": group.get("original"),
                    "additional": list(group.get("additional", [])),
                    "synonyms": list(group.get("synonyms", [])),
                    "related": list(group.get("related", [])),
                    "related_details": [
                        {
                            "term": detail.get("term"),
                            "cluster": detail.get("cluster"),
                            "search_volume": detail.get("search_volume"),
                            "competition": detail.get("competition"),
                            "relevance": detail.get("relevance"),
                        }
                        for detail in group.get("related_details", [])
                        if isinstance(detail, dict)
                    ],
                }
                for group in job.keyword_groups
            ],
            "keyword_evaluations": [
                {
                    "keyword": entry.get("keyword"),
                    "quality": entry.get("quality"),
                    "better_keywords": list(entry.get("better_keywords", [])),
                }
                for entry in job.keyword_evaluations
            ],
            "synonym_messages": list(job.synonym_messages),
            "synonyms_enabled": job.synonyms_enabled,
            "synonyms_expanded": job.synonyms_expanded,
            "synonyms_manual": job.synonyms_manual,
            "synonym_status": job.synonym_status,
            "start_date": job.start_date.isoformat() if job.start_date else None,
            "end_date": job.end_date.isoformat() if job.end_date else None,
            "concurrency": job.concurrency,
            "respect_robots": job.respect_robots,
            "status": job.status,
            "result_count": len(job.results),
            "results": [
                {
                    "source_url": result.source_url,
                    "target_url": result.target_url,
                    "matched_keywords": list(result.matched_keywords),
                }
                for result in job.results
            ],
        }
    append_saved_search(entry)
    return jsonify({"status": "saved"})


@app.route("/saved-searches")
def get_saved_searches():
    return jsonify({"searches": load_saved_searches()})


if __name__ == "__main__":
    app.run(debug=True)
